# ═══════════════════════════════════════════════════════════════
#  HÜCRE 9 — BACKTEST  (v2 — 15.10.2026)
#  ⚠️ Günlük sıraya GİRMEZ. Ayda bir ya da kurallar değişince çalıştır.
#     Yapay zekâ kullanmaz (API kredisi harcamaz).
#     İlk çalıştırma ~15-20 dk (veri indirme); aynı oturumda tekrar ~3-4 dk.
#
#  Ne yapar:
#   1) Son ~2,5 yılda her ayın ilk işlem gününde hücre 2'nin (v2.4) ve hücre 4'ün
#      puanlamasını — kodları birebir kopyalandı — SADECE O GÜN BİLİNEN verilerle hesaplar
#   2) Her bileşeni AYRI AYRI test eder: momentum, değerleme, büyüme, borç, FCF, RSI,
#      MACD, hareketli ortalamalar... + dünyada bilinen ölçütler (12 ay momentum, PD/DD,
#      kâr getirisi, düşük oynaklık). Hangisi BIST'te gerçekten işe yarıyor?
#   3) Geçmişe uydurmamak için: ilk yarıdaki aylarda işe yarayan ölçütleri seçer,
#      bunlardan bir "bileşik skor" kurar ve SADECE ikinci yarıda (seçimde hiç
#      kullanılmamış aylarda) test eder
#   4) Çıkış kuralları: 60 gün tut, −%15, −%20, ATR 2,5× / 3,5× / 4× stop, stop+hedef
#   5) Hücre 8'in makro skoru: makro kötüyken beklemek işe yarıyor mu?
#   6) Hücre 7'nin günlük kaydı (sistem_kayit.csv) varsa: yapay zekânın haber kararları
#      ve Öneri/Giriş kararları gerçekte ne getirdi?
#
#  Sınırlar (sonuç sayfasında da yazıyor):
#   - Haber/AI, KAP, DCF, İsyatırım puanları geçmişe dönük test edilemez (6. madde ileriye dönük)
#   - Çeyrek bilançolar yok: sadece yıllık (yıl sonundan 80 gün sonra biliniyor sayılır)
#   - Bugün borsada olan hisseler test edilir (sonradan borsadan çıkanlar yok)
# ═══════════════════════════════════════════════════════════════
import yfinance as yf
import pandas as pd
import numpy as np
import os
import time
import warnings
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
warnings.filterwarnings('ignore')
try:
    from tqdm.notebook import tqdm
except Exception:
    def tqdm(x, **k): return x

DRIVE      = '/content/drive/MyDrive/isyatirim'
EXCEL_YOLU = f'{DRIVE}/BIST_Tam_KAP_v9_11mart.xlsx'
KAYIT_YOLU = f'{DRIVE}/sistem_kayit.csv'        # hücre 7 v4.8+ yazıyor
HAM_CSV    = f'{DRIVE}/backtest_ham.csv'        # tüm gözlemler (Claude'a ayrıntılı analiz için)

# ── AYARLAR (hücre 2 ile aynı olmalı) ────────────────────────
YILLIK_TUFE = {2020: 14.60, 2021: 36.08, 2022: 64.27, 2023: 64.77,
               2024: 44.38, 2025: 30.89}
VARSAYILAN_TUFE    = 30.0
BEKLENEN_ENFLASYON = 0.25
ENFLASYON_DUZELT   = True
KAP_PUANA_KAT      = False
ILERI_REEL_TAVAN   = 10
FK_MIN_TAVAN       = 15
FK_ORT_TAVAN       = 20
FK_IYI_TAVAN       = 25
SEKTOR_MIN_HISSE   = 5
FK_SUPHELI         = 2

# ── BACKTEST AYARLARI ────────────────────────────────────────
AL_ESIK           = 55            # AL listesi eşiği (hücre 3-7)
KISA_GUN          = 20            # faktör testi: sonraki 20 işlem günü (~1 ay, aylar üst üste binmez)
TUTMA_GUN         = 60            # tutma testi: sonraki 60 işlem günü (~3 ay)
RAPOR_GECIKME_GUN = 80            # yıllık bilanço yıl sonundan kaç gün sonra biliniyor sayılsın
ILK_TEST          = '2024-04-01'  # ilk test tarihi (öncesinde en az 2 yıllık bilanço yok)
MAKRO_ESIK        = -20           # hücre 7: bunun altı 'RİSK YÜKSEK'
MIN_TL_HACIM      = 20            # portföy testinde en az günlük işlem hacmi (mn TL) — alınabilir hisseler
PORTFOY_SAYISI    = 15            # bileşik skorla seçilen portföydeki hisse sayısı
SECIM_T           = 1.5           # ilk yarıda bir ölçütün seçilmesi için gereken en az |t|

print("=" * 65)
print("  📊 BACKTEST v2 — bileşen bileşen, bugünkü kurallar, geçmişin verisi")
print(f"  {datetime.now().strftime('%d.%m.%Y %H:%M')}")
print("=" * 65)

# ── İSYATIRIM: hisse listesi, sektör, piyasa değeri (hücre 2 ile aynı) ──
ozet = pd.read_excel(f'{DRIVE}/temelozet.xlsx')
HISSELER = list(dict.fromkeys(ozet['Kod'].dropna().astype(str).str.upper().str.strip().tolist()))
PD_ISY = {}
pd_kolon = next((c for c in ozet.columns if 'Piyasa Değeri' in str(c)), None)
if pd_kolon is not None:
    for k, v in zip(ozet['Kod'], pd.to_numeric(ozet[pd_kolon], errors='coerce')):
        if pd.notna(k) and pd.notna(v) and v > 0:
            PD_ISY[str(k).upper().strip()] = float(v) * 1_000_000
SEKTOR = {}
if 'Sektör' in ozet.columns:
    for k, s in zip(ozet['Kod'], ozet['Sektör']):
        if pd.notna(k):
            SEKTOR[str(k).upper().strip()] = str(s) if pd.notna(s) else ''

def sektor_turu(kod):
    s = SEKTOR.get(kod, '').lower()
    if 'gyo' in s or 'gayrimenkul' in s:
        return 'gyo'
    if any(k in s for k in ['banka', 'sigorta', 'kiralama', 'faktoring',
                             'aracı kurum', 'menkul', 'finansman', 'emeklilik']):
        return 'finans'
    return ''

_KUR_CACHE = {}
def kur_serisi(para):
    para = (para or 'TRY').upper()
    if para in ('TRY', 'TL'):
        return None
    if para not in _KUR_CACHE:
        k = yf.Ticker(f'{para}TRY=X').history(period='6y')['Close']
        if getattr(k.index, 'tz', None) is not None:
            k.index = k.index.tz_localize(None)
        _KUR_CACHE[para] = k
    return _KUR_CACHE[para]

def kur(seri, tarih):
    if seri is None:
        return 1.0
    o = seri[seri.index <= pd.Timestamp(tarih)]
    return float(o.iloc[-1]) if len(o) else float(seri.iloc[0])

# ── HÜCRE 2'DEN AYNEN: enflasyon yardımcıları ───────────────
def tufe(yil):
    return YILLIK_TUFE.get(int(yil), VARSAYILAN_TUFE) / 100

def enflasyon_carpani(yil_bas, yil_son):
    """yil_bas sonundan yil_son sonuna kadar fiyatlar kaç kat arttı."""
    c = 1.0
    for y in range(int(yil_bas) + 1, int(yil_son) + 1):
        c *= 1 + tufe(y)
    return c

def yillik_buyume(degerler, yillar):
    """degerler/yillar: en yeni başta. (nominal, reel) yıllık bileşik büyüme, %."""
    if len(degerler) < 2 or degerler[0] <= 0 or degerler[-1] <= 0:
        return None, None
    n = max(int(yillar[0]) - int(yillar[-1]), 1)
    oran = degerler[0] / degerler[-1]
    nominal = (oran ** (1 / n) - 1) * 100
    reel = ((oran / enflasyon_carpani(yillar[-1], yillar[0])) ** (1 / n) - 1) * 100
    return nominal, reel

# ── HÜCRE 4'TEN AYNEN: RSI, MACD, teknik skor ───────────────
def rsi_hesapla(close, period=14):
    delta = close.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.ewm(com=period-1, min_periods=period).mean()
    avg_loss = loss.ewm(com=period-1, min_periods=period).mean()
    rs = avg_gain / avg_loss.replace(0, np.nan)
    return 100 - (100 / (1 + rs))

def macd_hesapla(close, fast=12, slow=26, signal=9):
    ema_fast = close.ewm(span=fast, adjust=False).mean()
    ema_slow = close.ewm(span=slow, adjust=False).mean()
    macd_line = ema_fast - ema_slow
    signal_line = macd_line.ewm(span=signal, adjust=False).mean()
    histogram = macd_line - signal_line
    return macd_line, signal_line, histogram

def teknik_skor_hesapla(veri):
    if veri is None:
        return 0, 'Veri yok'
    skor = 0
    ac = []
    rsi = veri.get('RSI')
    macd = veri.get('MACD')
    sig = veri.get('MACD_Signal')
    kes = veri.get('MACD_Kesisim')
    fiy = veri.get('Fiyat')
    ma20 = veri.get('MA20')
    ma50 = veri.get('MA50')
    bb = veri.get('BB_Yuzde')
    hvol = veri.get('Hacim_Orani')
    if rsi is not None:
        if rsi < 30:   skor += 12; ac.append(f'RSI {rsi:.0f} aşırı satım 🟢')
        elif rsi < 45: skor += 8;  ac.append(f'RSI {rsi:.0f} alım bölgesi 🟢')
        elif rsi < 60: skor += 5;  ac.append(f'RSI {rsi:.0f} nötr 🟡')
        elif rsi < 70: skor += 0;  ac.append(f'RSI {rsi:.0f} yüksek 🟡')
        else:          skor -= 10; ac.append(f'RSI {rsi:.0f} aşırı alım 🔴')
    if macd is not None and sig is not None:
        if kes:              skor += 10; ac.append('MACD kesişim ⚡')
        elif macd>sig>0:     skor += 8;  ac.append('MACD pozitif & yukarı 🟢')
        elif macd>sig:       skor += 4;  ac.append('MACD yukarı (negatif bölge) 🟡')
        elif macd<sig and macd>0: skor -= 3; ac.append('MACD satış 🟠')
        else:                skor -= 8;  ac.append('MACD negatif & aşağı 🔴')
    if fiy and ma20:
        if fiy > ma20: skor += 8; ac.append('MA20 üstünde 🟢')
        else:          skor -= 5; ac.append('MA20 altında 🔴')
    if fiy and ma50:
        if fiy > ma50: skor += 5; ac.append('MA50 üstünde 🟢')
        else:          skor -= 3; ac.append('MA50 altında 🟠')
    if bb is not None:
        if bb < 20:   skor += 5; ac.append(f'Bollinger alt ({bb:.0f}%) 🟢')
        elif bb > 80: skor -= 5; ac.append(f'Bollinger üst ({bb:.0f}%) 🔴')
        else:         skor += 2; ac.append(f'Bollinger orta ({bb:.0f}%) 🟡')
    if hvol is not None:
        if hvol >= 2.5:   skor += 8; ac.append(f'Hacim patlaması {hvol:.1f}x 🔥')
        elif hvol >= 1.5: skor += 5; ac.append(f'Hacim artışı {hvol:.1f}x 🟢')
        elif hvol >= 1.0: skor += 2; ac.append(f'Hacim normal {hvol:.1f}x 🟡')
        else:             skor -= 3; ac.append(f'Hacim düşük {hvol:.1f}x 🔴')
    return skor, ' | '.join(ac)

def teknik_hesapla(df):
    """Hücre 4'teki teknik_veri_cek'in hesap kısmı, aynen (df = son 3 ayın fiyatı)."""
    close = df['Close']
    volume = df['Volume']
    rsi_son = rsi_hesapla(close).iloc[-1]
    macd_line, signal_line, histogram = macd_hesapla(close)
    macd_son = macd_line.iloc[-1]
    signal_son = signal_line.iloc[-1]
    histogram_son = histogram.iloc[-1]
    histogram_onc = histogram.iloc[-2] if len(histogram) > 1 else 0
    macd_kesisim = histogram_onc < 0 and histogram_son > 0
    ma20 = close.rolling(20).mean().iloc[-1]
    ma50 = close.rolling(50).mean().iloc[-1] if len(close) >= 50 else None
    fiyat = close.iloc[-1]
    bb_orta = close.rolling(20).mean()
    bb_std = close.rolling(20).std()
    bb_ust = (bb_orta + 2 * bb_std).iloc[-1]
    bb_alt = (bb_orta - 2 * bb_std).iloc[-1]
    bb_yuzde = (fiyat - bb_alt) / (bb_ust - bb_alt) * 100
    vol_ort_20 = volume.rolling(20).mean().iloc[-1]
    vol_son3 = volume.iloc[-3:].mean()
    hacim_orani = vol_son3 / vol_ort_20 if vol_ort_20 > 0 else 1
    trend_5 = (fiyat - close.iloc[-5]) / close.iloc[-5] * 100 if len(close) >= 5 else 0
    trend_20 = (fiyat - close.iloc[-20]) / close.iloc[-20] * 100 if len(close) >= 20 else 0
    teknik = {
        'RSI': round(rsi_son, 1), 'MACD': round(macd_son, 4),
        'MACD_Signal': round(signal_son, 4), 'MACD_Histogram': round(histogram_son, 4),
        'MACD_Kesisim': macd_kesisim, 'MA20': round(ma20, 2),
        'MA50': round(ma50, 2) if ma50 else None, 'BB_Yuzde': round(bb_yuzde, 1),
        'Hacim_Orani': round(hacim_orani, 2), 'Trend_5G': round(trend_5, 1),
        'Trend_20G': round(trend_20, 1), 'Fiyat': round(fiyat, 2),
    }
    return teknik

# ── HÜCRE 8'DEN AYNEN: makro skor ───────────────────────────
def gosterge_hesapla(seriler, tarih=None):
    """Seri sözlüğünden göstergeleri hesaplar. tarih verilirse o güne kadarki veriyle."""
    gostergeler = {}
    for ad, k in seriler.items():
        if tarih is not None:
            k = k[k.index <= tarih]
        if len(k) < 5:
            continue
        son     = float(k.iloc[-1])
        once_1g = float(k.iloc[-2]) if len(k) >= 2 else son
        once_1h = float(k.iloc[-6]) if len(k) >= 6 else son
        once_1a = float(k.iloc[-22]) if len(k) >= 22 else float(k.iloc[0])
        gostergeler[ad] = {
            'Son Değer':    round(son, 2),
            'Günlük (%)':   round((son / once_1g - 1) * 100, 2),
            'Haftalık (%)': round((son / once_1h - 1) * 100, 2),
            'Aylık (%)':    round((son / once_1a - 1) * 100, 2),
        }

    # Gram altın (TL) ≈ ons ($) × USD/TL ÷ 31,1035 (yaklaşık; değişimler bileşik)
    try:
        g, u = gostergeler['Altın (ons $)'], gostergeler['USD/TL']
        def bilesik(a, b): return round(((1 + a / 100) * (1 + b / 100) - 1) * 100, 2)
        gostergeler['Gram Altın (TL)'] = {
            'Son Değer':    round(g['Son Değer'] * u['Son Değer'] / 31.1035, 0),
            'Günlük (%)':   bilesik(g['Günlük (%)'],   u['Günlük (%)']),
            'Haftalık (%)': bilesik(g['Haftalık (%)'], u['Haftalık (%)']),
            'Aylık (%)':    bilesik(g['Aylık (%)'],    u['Aylık (%)']),
        }
    except: pass
    return gostergeler

def karar_ver(skor):
    if skor >= 30:    return '🟢 UYGUN ORTAM', '00FF88', '001A10'
    elif skor >= 5:   return '🟡 DİKKATLİ OL', 'FFB300', '1A1400'
    elif skor >= -20: return '🟠 RİSKLİ ORTAM', 'FF8C00', '1A0A00'
    return '🔴 RİSK YÜKSEK', 'FF4757', '2D0000'

def makro_skor_hesapla(gostergeler):
    """
    Göstergelerden makro risk skoru üretir
    Negatif = riskli ortam, Pozitif = iyi ortam
    -100 ile +100 arası
    """
    skor = 0
    sinyaller = []
    risk_seviyesi = []

    # ── DOLAR/TL ──
    try:
        usd = gostergeler['USD/TL']          # veri yoksa KeyError → puan verilmez
        gunluk = usd.get('Günlük (%)', 0)
        aylik  = usd.get('Aylık (%)', 0)
        if gunluk >= 2:      skor -= 15; risk_seviyesi.append('🔴'); sinyaller.append(f'USD/TL günlük +%{gunluk:.1f} — TL değer kaybı')
        elif gunluk >= 1:    skor -= 8;  risk_seviyesi.append('🟠'); sinyaller.append(f'USD/TL günlük +%{gunluk:.1f} — TL baskı altında')
        elif gunluk <= -1:   skor += 8;  risk_seviyesi.append('🟢'); sinyaller.append(f'USD/TL günlük %{gunluk:.1f} — TL güçleniyor')
        else:                skor += 2;  risk_seviyesi.append('🟡'); sinyaller.append(f'USD/TL stabil ({usd.get("Son Değer",0):.2f})')

        if aylik >= 5:       skor -= 10; sinyaller.append(f'USD/TL aylık +%{aylik:.1f} — TL trend kaybı')
        elif aylik <= -3:    skor += 5;  sinyaller.append(f'USD/TL aylık %{aylik:.1f} — TL güçlü trend')
    except: pass

    # ── BRENT PETROL ──
    try:
        brent = gostergeler['Brent Petrol']          # veri yoksa KeyError → puan verilmez
        gunluk = brent.get('Günlük (%)', 0)
        aylik  = brent.get('Aylık (%)', 0)
        son    = brent.get('Son Değer', 0)

        if aylik >= 30:      skor -= 20; risk_seviyesi.append('🔴'); sinyaller.append(f'Brent petrol aylık +%{aylik:.1f} — jeopolitik kriz!')
        elif aylik >= 15:    skor -= 12; risk_seviyesi.append('🔴'); sinyaller.append(f'Brent petrol aylık +%{aylik:.1f} — ciddi yükseliş')
        elif aylik >= 5:     skor -= 5;  risk_seviyesi.append('🟠'); sinyaller.append(f'Brent petrol aylık +%{aylik:.1f} — yükseliş baskısı')
        elif aylik <= -10:   skor += 8;  risk_seviyesi.append('🟢'); sinyaller.append(f'Brent petrol aylık %{aylik:.1f} — düşüş, enflasyon azalır')
        else:                skor += 1;  risk_seviyesi.append('🟡'); sinyaller.append(f'Brent petrol aylık %{aylik:+.1f} — aylık değişim sınırlı')

        if son >= 100:       skor -= 10; sinyaller.append(f'Brent {son:.1f}$ — 100$ eşiğinin üstünde!')
        elif son >= 85:      skor -= 5;  sinyaller.append(f'Brent {son:.1f}$ — yüksek seviye')
    except: pass

    # ── VIX (KORKU ENDEKSİ) ──
    try:
        vix = gostergeler['VIX']          # veri yoksa KeyError → puan verilmez
        son = vix.get('Son Değer', 20)
        gunluk = vix.get('Günlük (%)', 0)

        if son >= 35:        skor -= 20; risk_seviyesi.append('🔴'); sinyaller.append(f'VIX {son:.1f} — PANIK! Global korku çok yüksek')
        elif son >= 25:      skor -= 12; risk_seviyesi.append('🔴'); sinyaller.append(f'VIX {son:.1f} — Global belirsizlik yüksek')
        elif son >= 20:      skor -= 5;  risk_seviyesi.append('🟠'); sinyaller.append(f'VIX {son:.1f} — Orta belirsizlik')
        elif son <= 15:      skor += 10; risk_seviyesi.append('🟢'); sinyaller.append(f'VIX {son:.1f} — Global piyasalar sakin')
        else:                skor += 3;  risk_seviyesi.append('🟡'); sinyaller.append(f'VIX {son:.1f} — Normal seviye')

        if gunluk >= 20:     skor -= 10; sinyaller.append(f'VIX günlük +%{gunluk:.1f} — ani korku artışı!')
    except: pass

    # ── S&P500 ──
    try:
        sp = gostergeler['S&P500']          # veri yoksa KeyError → puan verilmez
        gunluk = sp.get('Günlük (%)', 0)
        haftalik = sp.get('Haftalık (%)', 0)

        if gunluk <= -2:     skor -= 15; risk_seviyesi.append('🔴'); sinyaller.append(f'S&P500 günlük %{gunluk:.1f} — sert düşüş')
        elif gunluk <= -1:   skor -= 8;  risk_seviyesi.append('🟠'); sinyaller.append(f'S&P500 günlük %{gunluk:.1f} — düşüş')
        elif gunluk >= 1:    skor += 8;  risk_seviyesi.append('🟢'); sinyaller.append(f'S&P500 günlük +%{gunluk:.1f} — yükseliş')
        else:                skor += 2;  risk_seviyesi.append('🟡'); sinyaller.append(f'S&P500 stabil (%{gunluk:.1f})')

        if haftalik <= -3:   skor -= 8;  sinyaller.append(f'S&P500 haftalık %{haftalik:.1f} — haftalık trend kötü')
        elif haftalik >= 2:  skor += 5;  sinyaller.append(f'S&P500 haftalık +%{haftalik:.1f} — haftalık trend iyi')
    except: pass

    # ── ALTIN ──
    try:
        altin = gostergeler['Altın (ons $)']          # veri yoksa KeyError → puan verilmez
        gunluk = altin.get('Günlük (%)', 0)
        aylik  = altin.get('Aylık (%)', 0)

        if gunluk >= 2:      skor -= 8;  sinyaller.append(f'Altın günlük +%{gunluk:.1f} — güvenli liman talebi artıyor')
        elif gunluk <= -2:   skor += 5;  sinyaller.append(f'Altın günlük %{gunluk:.1f} — risk iştahı artıyor')
        else:                sinyaller.append(f'Altın stabil (%{gunluk:.1f})')

        if aylik >= 10:      skor -= 10; sinyaller.append(f'Altın aylık +%{aylik:.1f} — uzun süreli kaçış!')
        elif aylik <= -8:    skor += 8;  sinyaller.append(f'Altın aylık %{aylik:.1f} — risk iştahı güçlü')
    except: pass

    # ── BIST100 ──
    try:
        bist = gostergeler['BIST100']          # veri yoksa KeyError → puan verilmez
        gunluk = bist.get('Günlük (%)', 0)
        haftalik = bist.get('Haftalık (%)', 0)
        aylik = bist.get('Aylık (%)', 0)

        if gunluk <= -2:     skor -= 12; risk_seviyesi.append('🔴'); sinyaller.append(f'BIST100 günlük %{gunluk:.1f} — sert düşüş')
        elif gunluk <= -1:   skor -= 6;  risk_seviyesi.append('🟠'); sinyaller.append(f'BIST100 günlük %{gunluk:.1f} — düşüş')
        elif gunluk >= 1:    skor += 6;  risk_seviyesi.append('🟢'); sinyaller.append(f'BIST100 günlük +%{gunluk:.1f} — yükseliş')
        else:                skor += 1;  risk_seviyesi.append('🟡'); sinyaller.append(f'BIST100 stabil (%{gunluk:.1f})')

        if aylik <= -5:      skor -= 10; sinyaller.append(f'BIST100 aylık %{aylik:.1f} — trend kötü')
        elif aylik >= 5:     skor += 8;  sinyaller.append(f'BIST100 aylık +%{aylik:.1f} — trend iyi')
    except: pass

    # Verisi gelmeyen göstergeler puana katılmadı
    for ad in ['USD/TL', 'Brent Petrol', 'VIX', 'S&P500', 'Altın (ons $)', 'BIST100']:
        if ad not in gostergeler:
            sinyaller.append(f'⚠️ {ad} verisi alınamadı — puana katılmadı')

    # Skoru -100 ile +100 arasında tut
    skor = max(min(skor, 100), -100)
    karar, karar_renk, karar_bg = karar_ver(skor)
    return skor, karar, karar_renk, karar_bg, sinyaller

# ════════════════════════════════════════════════════════════════
#  1) VERİ İNDİR (aynı Colab oturumunda tekrar çalıştırınca indirmez)
# ════════════════════════════════════════════════════════════════
def _naive(df):
    df = df.copy()
    if getattr(df.index, 'tz', None) is not None:
        df.index = df.index.tz_localize(None)
    df.index = pd.to_datetime(df.index).normalize()
    return df[~df.index.duplicated(keep='last')]

try:
    BT_VERI
except NameError:
    BT_VERI = {}
eksik = [h for h in HISSELER if h not in BT_VERI]
print(f"\n📥 {len(HISSELER)} hisse | indirilecek: {len(eksik)} (öncekiler bellekte)")
for h in tqdm(eksik, desc='Veri'):
    try:
        t = yf.Ticker(f'{h}.IS')
        hist = t.history(period='5y')
        if hist is None or len(hist) < 300:
            BT_VERI[h] = None
            continue
        info = t.info or {}
        BT_VERI[h] = {'hist': _naive(hist[['Open', 'High', 'Low', 'Close', 'Volume']]),
                      'inc': t.income_stmt, 'bal': t.balance_sheet, 'cf': t.cashflow,
                      'shares': info.get('sharesOutstanding', 0) or 0,
                      'para': (info.get('financialCurrency') or 'TRY').upper()}
    except Exception:
        pass                       # kaydetme: tekrar çalıştırınca yeniden denenir
    time.sleep(0.3)

for h, v in BT_VERI.items():      # hisse sayısı: hücre 2'deki gibi İsyatırım ile düzelt
    if not v:
        continue
    sh = v['shares']; fiyat_simdi = float(v['hist']['Close'].iloc[-1])
    if h in PD_ISY and fiyat_simdi > 0:
        isy = PD_ISY[h] / fiyat_simdi
        if sh <= 0:
            sh = isy
        else:
            oran = isy / sh
            if 0.05 <= oran <= 20 and abs(oran - 1) > 0.25:
                sh = isy
    v['shares_duz'] = sh

def seri_indir(semboller, deneme=3):
    for s in semboller:
        for _ in range(deneme):
            try:
                hh = yf.Ticker(s).history(period='5y')
                if hh is not None and len(hh) > 100:
                    return _naive(hh)
            except Exception:
                pass
            time.sleep(1)
    return None

try:
    BT_MAKRO
except NameError:
    BT_MAKRO = {}
for ad, sem in {'BIST100': ['XU100.IS'], 'USD/TL': ['USDTRY=X', 'TRY=X'], 'Brent Petrol': ['BZ=F'],
                'VIX': ['^VIX'], 'S&P500': ['^GSPC'], 'Altın (ons $)': ['GC=F']}.items():
    if BT_MAKRO.get(ad) is None:
        BT_MAKRO[ad] = seri_indir(sem)
XU = BT_MAKRO['BIST100']
if XU is None:
    raise SystemExit('❌ BIST100 verisi alınamadı, biraz sonra tekrar çalıştır.')
print(f"✅ Veri hazır: {sum(1 for v in BT_VERI.values() if v)} hisse | "
      f"makro: {', '.join(k for k, v in BT_MAKRO.items() if v is not None)}")

# ════════════════════════════════════════════════════════════════
#  2) GEÇMİŞ TARİHTE PUAN (hücre 2 v2.4 + hücre 4 kodu birebir) + ölçütler
# ════════════════════════════════════════════════════════════════
class _CeyrekYok:
    quarterly_income_stmt = None       # geçmiş tarihte çeyrek verisi yok → yıllık kullanılır
_T_YOK = _CeyrekYok()
_TUTARLILIK = {'🟢 Sürekli Büyüme': 3, '🟢 Güçlü Büyüme': 2, '🟡 Kararsız': 1, '🔴 Tutarsız': 0}

def mevcut(df, tarih):
    """tarih itibarıyla açıklanmış yıllık sütunlar (en yeni başta)."""
    if df is None or getattr(df, 'empty', True):
        return df
    kol = [c for c in df.columns if pd.Timestamp(c) + pd.Timedelta(days=RAPOR_GECIKME_GUN) <= tarih]
    kol = sorted(kol, key=lambda c: pd.Timestamp(c), reverse=True)
    return df[kol]

def _f(x):
    try:
        x = float(x)
        return x if np.isfinite(x) else None
    except Exception:
        return None

def puan_hesapla(h, v, tarih):
    hist = v['hist'][v['hist'].index <= tarih]
    if len(hist) < 252 or (tarih - hist.index[-1]).days > 7:
        return None
    inc = mevcut(v['inc'], tarih); bal = mevcut(v['bal'], tarih); cf = mevcut(v['cf'], tarih)
    if inc is None or inc.empty:
        return None
    shares = v['shares_duz']
    if not shares or shares < 1000:
        return None
    kur_s = kur_serisi(v['para'])
    tur = sektor_turu(h); info = {}; t = _T_YOK; kap_alimlar = {}
    # ↓↓↓ hücre 2'den aynen ↓↓↓
    kapanis = hist['Close'].copy()
    if getattr(kapanis.index, 'tz', None) is not None:
        kapanis.index = kapanis.index.tz_localize(None)

    son_fiyat  = hist['Close'].iloc[-1]
    yuksek_52h = hist['High'].iloc[-252:].max()
    dusuk_52h  = hist['Low'].iloc[-252:].min()
    zirveden   = ((son_fiyat - yuksek_52h) / yuksek_52h) * 100
    dipten     = ((son_fiyat - dusuk_52h)  / dusuk_52h)  * 100
    yil_ici    = ((son_fiyat / hist['Close'].iloc[-252]) - 1) * 100 if len(hist) >= 252 else 0
    aylik      = ((son_fiyat / hist['Close'].iloc[-21])  - 1) * 100 if len(hist) >= 21  else 0
    haftalik   = ((son_fiyat / hist['Close'].iloc[-5])   - 1) * 100 if len(hist) >= 5   else 0
    gunluk     = ((son_fiyat / hist['Close'].iloc[-2])   - 1) * 100 if len(hist) >= 2   else 0
    son3       = ((son_fiyat / hist['Close'].iloc[-4])   - 1) * 100 if len(hist) >= 4   else 0

    son_1ay_adet = hist['Volume'].iloc[-20:].mean()
    onceki_adet  = hist['Volume'].iloc[-60:-20].mean()
    hacim_artis  = ((son_1ay_adet - onceki_adet) / onceki_adet * 100) if onceki_adet > 0 else 0
    tl_hacim_ort = (hist['Close'] * hist['Volume']).iloc[-20:].mean() / 1_000_000
    if tl_hacim_ort >= 200:  likidite = '🟢 Yüksek'
    elif tl_hacim_ort >= 50: likidite = '🟢 Normal'
    elif tl_hacim_ort >= 10: likidite = '🟡 Düşük'
    else:                    likidite = '🔴 Çok Düşük'

    # ── Yıllık net kârlar (en yeni başta) ──
    net_karlar = []; kar_tarihleri = []
    if 'Net Income' in inc.index:
        for col in inc.columns[:4]:
            v = inc.loc['Net Income', col]
            if pd.notna(v) and float(v) != 0:
                net_karlar.append(float(v) * kur(kur_s, col)); kar_tarihleri.append(pd.Timestamp(col))
    if len(net_karlar) < 2:
        return None
    kar_yillari = [d.year for d in kar_tarihleri]
    hbk_gecmis  = [nk / shares for nk in net_karlar]

    # ── Güncel HBK: son 4 çeyrek, olmazsa son yıllık ──
    hbk_guncel = hbk_gecmis[0]; hbk_kaynak = 'Yıllık'
    q_son_tarih = None
    try:
        q = t.quarterly_income_stmt
        if q is not None and not q.empty and 'Net Income' in q.index:
            qs = q.loc['Net Income'].iloc[:4]
            qv = [float(v) * kur(kur_s, d) for d, v in qs.items() if pd.notna(v)]
            if len(qv) == 4:
                ttm = sum(qv) / shares
                q_son_tarih = pd.Timestamp(qs.index[0])
                # Akıl kontrolü: yıllıktan çok saparsa çeyrek verisi bozuk say
                if hbk_gecmis[0] > 0 and 0.2 <= ttm / hbk_gecmis[0] <= 2.5:
                    hbk_guncel = ttm; hbk_kaynak = 'Son 4 çeyrek'
                elif hbk_gecmis[0] <= 0:
                    hbk_guncel = ttm; hbk_kaynak = 'Son 4 çeyrek'
    except: pass

    if hbk_guncel <= 0:
        return None
    guncel_fk = son_fiyat / hbk_guncel

    # ── Gerçek tarihi F/K: o yılın kapanış fiyatı ÷ o yılın HBK ──
    fk_tarihi = []
    for tarih, hbk in zip(kar_tarihleri, hbk_gecmis):
        if hbk <= 0:
            continue
        o_gun = kapanis[kapanis.index <= tarih]
        if len(o_gun) == 0 or (tarih - o_gun.index[-1]).days > 10:
            continue          # o tarihte fiyat verisi yok
        fk_tarihi.append(o_gun.iloc[-1] / hbk)

    if len(fk_tarihi) >= 2:
        fk_min    = min(min(fk_tarihi), FK_MIN_TAVAN)
        fk_medyan = min(float(np.median(fk_tarihi)), FK_ORT_TAVAN)
        fk_med_ham = float(np.median(fk_tarihi))          # sektör istatistiği için (tavansız)
        fk_max    = min(max(fk_tarihi), 35)
    else:
        fk_min = fk_medyan = fk_max = fk_med_ham = None

    # ── Büyüme (nominal ve reel) ──
    buyume, reel_buyume = yillik_buyume(net_karlar, kar_yillari)
    buyume_yok = buyume is None          # baş ya da son yıl zararda
    if buyume_yok:
        buyume, reel_buyume = 0.0, 0.0
    # puanlamada kullanılan büyüme (hesaplanamıyorsa ceza alsın)
    b_puan = -100.0 if buyume_yok else (reel_buyume if ENFLASYON_DUZELT else buyume)

    buyume_tutarli_skor = 'Veri Yok'
    if len(net_karlar) >= 3:
        b_yil = 0
        for i in range(len(net_karlar) - 1):
            esik = net_karlar[i + 1] * (1 + tufe(kar_yillari[i])) if ENFLASYON_DUZELT and net_karlar[i + 1] > 0 \
                   else net_karlar[i + 1]
            if net_karlar[i] > esik:
                b_yil += 1
        oran = b_yil / (len(net_karlar) - 1)
        if oran == 1.0:    buyume_tutarli_skor = '🟢 Sürekli Büyüme'
        elif oran >= 0.75: buyume_tutarli_skor = '🟢 Güçlü Büyüme'
        elif oran >= 0.5:  buyume_tutarli_skor = '🟡 Kararsız'
        else:              buyume_tutarli_skor = '🔴 Tutarsız'

    roe_tutarli_skor = 'Veri Yok'; roe_ort = None
    try:
        if bal is not None and not bal.empty and 'Net Income' in inc.index:
            roe_l = []
            for i, col in enumerate(inc.columns[:4]):
                ni = inc.loc['Net Income', col]
                if i < len(bal.columns) and 'Common Stock Equity' in bal.index:
                    eq = bal.loc['Common Stock Equity', bal.columns[i]]
                    if pd.notna(ni) and pd.notna(eq) and float(eq) > 0:
                        roe_l.append(float(ni) / float(eq) * 100)
            if len(roe_l) >= 2:
                roe_ort = np.mean(roe_l)
                roe_cv  = np.std(roe_l) / abs(roe_ort) if roe_ort != 0 else 999
                if roe_ort >= 15 and roe_cv < 0.3:   roe_tutarli_skor = '🟢 Yüksek & Tutarlı'
                elif roe_ort >= 10 and roe_cv < 0.5: roe_tutarli_skor = '🟢 İyi ROE'
                elif roe_ort >= 5:                   roe_tutarli_skor = '🟡 Orta ROE'
                elif roe_ort > 0:                    roe_tutarli_skor = '🟠 Düşük ROE'
                else:                                roe_tutarli_skor = '🔴 Negatif ROE'
    except: pass

    peg = None; peg_skor = 'Veri Yok'
    if guncel_fk > 0 and b_puan > 5:
        peg = round(guncel_fk / b_puan, 2)
        if peg < 0.5:    peg_skor = '🟢 Çok Ucuz'
        elif peg < 1.0:  peg_skor = '🟢 Ucuz'
        elif peg < 1.5:  peg_skor = '🟡 Makul'
        elif peg < 2.0:  peg_skor = '🟠 Pahalı'
        else:            peg_skor = '🔴 Çok Pahalı'

    kap_islem_sayisi = kap_alimlar.get(h, {}).get('islem_sayisi', 0)
    if kap_islem_sayisi >= 3:   kap_sinyal = '🟢 Güçlü Alım'
    elif kap_islem_sayisi >= 1: kap_sinyal = '🟡 Alım Var'
    else:                       kap_sinyal = '-'

    # ── HEDEF FİYAT: 12 ay sonraki HBK × tarihi F/K ──
    # Son dönem: son 4 çeyrek kârı, son yıllık kârla reel olarak karşılaştırılır
    reel_son = None
    if hbk_kaynak == 'Son 4 çeyrek' and q_son_tarih is not None and hbk_gecmis[0] > 0:
        ara_yil = (q_son_tarih - kar_tarihleri[0]).days / 365.25
        if ara_yil >= 0.2:
            enf = (1 + tufe(kar_tarihleri[0].year + 1)) ** ara_yil
            reel_son = ((hbk_guncel / hbk_gecmis[0]) / enf - 1) * 100
    # Gelecek yıl reel büyüme: uzun dönem ortalaması ile son dönemden düşük olanı
    reel_ileri = reel_buyume if reel_son is None else min(reel_buyume, reel_son)
    reel_ileri = max(min(reel_ileri, ILERI_REEL_TAVAN), -20)
    if ENFLASYON_DUZELT:
        g_ileri = (1 + reel_ileri / 100) * (1 + BEKLENEN_ENFLASYON) - 1
    else:
        g_ileri = max(min(buyume, 25), -20) / 100
    tahmini_hbk = hbk_guncel * (1 + g_ileri)
    veri_uyari = ''
    if guncel_fk < FK_SUPHELI:
        veri_uyari = f'F/K {guncel_fk:.1f}: veri şüpheli (tek seferlik kâr ya da hisse sayısı hatası)'
    # Hedef fiyat döngüden sonra (sektör F/K'ları belli olunca) hesaplanıyor

    pd_dd = info.get('priceToBook', None)
    roe   = (info.get('returnOnEquity', 0) or 0) * 100

    ciro_buyume = None; reel_ciro = None
    try:
        if 'Total Revenue' in inc.index:
            cirolar = []; ciro_yil = []
            for col in inc.columns[:4]:
                v = inc.loc['Total Revenue', col]
                if pd.notna(v) and float(v) > 0:
                    cirolar.append(float(v) * kur(kur_s, col)); ciro_yil.append(pd.Timestamp(col).year)
            nom, reel = yillik_buyume(cirolar, ciro_yil)
            if nom is not None:
                ciro_buyume = round(nom, 1); reel_ciro = round(reel, 1)
    except: pass

    faiz_karsilama = None; faiz_kategori = 'Veri Yok'
    if tur:
        faiz_kategori = 'Uygulanmaz'
    else:
        try:
            favok = None; faiz_gider = None
            for k in ['EBITDA', 'Operating Income']:
                if k in inc.index:
                    v = inc.loc[k].iloc[0]
                    if pd.notna(v): favok = float(v); break
            for k in ['Interest Expense', 'Interest Expense Non Operating']:
                if k in inc.index:
                    v = inc.loc[k].iloc[0]
                    if pd.notna(v) and float(v) != 0:
                        faiz_gider = abs(float(v)); break
            if favok and faiz_gider and faiz_gider > 0:
                faiz_karsilama = round(favok / faiz_gider, 1)
                if faiz_karsilama >= 5:   faiz_kategori = '🟢 Güçlü'
                elif faiz_karsilama >= 3: faiz_kategori = '🟡 Yeterli'
                elif faiz_karsilama >= 1: faiz_kategori = '🟠 Riskli'
                else:                    faiz_kategori = '🔴 Kritik'
        except: pass

    fcf = None; fcf_verim = None; fcf_kategori = 'Veri Yok'
    if tur:
        fcf_kategori = 'Uygulanmaz'
    else:
        try:
            if cf is not None and not cf.empty:
                op_cf = None; capex = None
                for k in ['Operating Cash Flow', 'Cash From Operations']:
                    if k in cf.index:
                        v = cf.loc[k].iloc[0]
                        if pd.notna(v): op_cf = float(v); break
                for k in ['Capital Expenditure', 'Purchase Of PPE', 'Capital Expenditures']:
                    if k in cf.index:
                        v = cf.loc[k].iloc[0]
                        if pd.notna(v): capex = abs(float(v)); break
                if op_cf is not None and capex is not None:
                    fcf_raw       = (op_cf - capex) * kur(kur_s, cf.columns[0])
                    piyasa_degeri = son_fiyat * shares
                    if piyasa_degeri > 0:
                        fcf_verim = round((fcf_raw / piyasa_degeri) * 100, 1)
                    fcf = round(fcf_raw / 1_000_000, 1)
                    if fcf_verim is not None:
                        if fcf_verim >= 8:   fcf_kategori = '🟢 Güçlü FCF'
                        elif fcf_verim >= 3: fcf_kategori = '🟢 Pozitif FCF'
                        elif fcf_verim >= 0: fcf_kategori = '🟡 Zayıf FCF'
                        else:                fcf_kategori = '🔴 Negatif FCF'
        except: pass

    # ── Borç: sadece finansal borç (Total Debt) ──
    borc_ozsermaye = None; net_borc_ozserm = None; borc_riski = 'Veri Yok'
    if tur:
        borc_riski = 'Uygulanmaz'
    else:
        try:
            if bal is not None and not bal.empty:
                tb = None; nakit = 0.0; oz = None
                if 'Total Debt' in bal.index and pd.notna(bal.loc['Total Debt'].iloc[0]):
                    tb = float(bal.loc['Total Debt'].iloc[0])
                if 'Common Stock Equity' in bal.index and pd.notna(bal.loc['Common Stock Equity'].iloc[0]):
                    oz = float(bal.loc['Common Stock Equity'].iloc[0])
                for k in ['Cash Cash Equivalents And Short Term Investments', 'Cash And Cash Equivalents']:
                    if k in bal.index and pd.notna(bal.loc[k].iloc[0]):
                        nakit = float(bal.loc[k].iloc[0]); break
                if tb is not None and oz and oz > 0:
                    borc_ozsermaye  = round(tb / oz, 2)
                    net_borc_ozserm = round((tb - nakit) / oz, 2)
                if net_borc_ozserm is not None:
                    if net_borc_ozserm < 0:      borc_riski = '🟢 Net Nakit'
                    elif net_borc_ozserm < 0.5:  borc_riski = '🟢 Düşük Borç'
                    elif net_borc_ozserm < 1.0:  borc_riski = '🟡 Orta Borç'
                    elif net_borc_ozserm < 2.0:  borc_riski = '🟠 Yüksek Borç'
                    else:                         borc_riski = '🔴 Tehlikeli Borç'
        except: pass

    # ── SKOR ───────────────────────────────────────────
    skor = 0      # değerleme puanı (+30 / −15) döngüden sonra eklenecek
    if zirveden >= -5:    skor += 12
    elif zirveden >= -15: skor += 8
    elif zirveden >= -25: skor += 4
    elif zirveden >= -40: skor += 1
    else:                 skor -= 6
    if aylik >= 15:     skor += 6
    elif aylik >= 8:    skor += 4
    elif aylik >= 3:    skor += 2
    elif aylik < 0:     skor -= 3
    if haftalik >= 5:   skor += 3
    elif haftalik >= 2: skor += 2
    elif haftalik < 0:  skor -= 1
    if hacim_artis >= 60:   skor += 6
    elif hacim_artis >= 25: skor += 4
    elif hacim_artis >= 0:  skor += 1
    else:                   skor -= 2
    # Büyüme: reel ise 20/10/0 eşikleri, nominal ise eski 40/20/5
    e1, e2, e3 = (20, 10, 0) if ENFLASYON_DUZELT else (40, 20, 5)
    if b_puan >= e1:   skor += 15
    elif b_puan >= e2: skor += 10
    elif b_puan >= e3: skor += 5
    else:              skor -= 8
    if buyume_tutarli_skor == '🟢 Sürekli Büyüme':  skor += 15
    elif buyume_tutarli_skor == '🟢 Güçlü Büyüme':  skor += 10
    elif buyume_tutarli_skor == '🔴 Tutarsız':       skor -= 12
    if peg is not None:
        if peg < 0.5:    skor += 8
        elif peg < 1.0:  skor += 5
        elif peg < 1.5:  skor += 2
        elif peg < 2.0:  skor -= 3
        else:            skor -= 8
    if roe_tutarli_skor == '🟢 Yüksek & Tutarlı':  skor += 8
    elif roe_tutarli_skor == '🟢 İyi ROE':          skor += 6
    elif roe_tutarli_skor == '🟡 Orta ROE':         skor += 2
    elif roe_tutarli_skor == '🟠 Düşük ROE':        skor -= 4
    elif roe_tutarli_skor == '🔴 Negatif ROE':      skor -= 10
    c_puan = reel_ciro if ENFLASYON_DUZELT else ciro_buyume
    c1, c2 = (15, 5) if ENFLASYON_DUZELT else (30, 15)
    if c_puan is not None:
        if c_puan >= c1:   skor += 6
        elif c_puan >= c2: skor += 4
        elif c_puan >= 0:  skor += 1
        else:              skor -= 4
    if net_borc_ozserm is not None:
        if net_borc_ozserm < 0:      skor += 10
        elif net_borc_ozserm < 0.5:  skor += 6
        elif net_borc_ozserm < 1.0:  skor += 0
        elif net_borc_ozserm < 2.0:  skor -= 12
        else:                         skor -= 25
    if faiz_karsilama is not None:
        if faiz_karsilama >= 5:   skor += 10
        elif faiz_karsilama >= 3: skor += 6
        elif faiz_karsilama >= 1: skor += 0
        elif faiz_karsilama > 0:  skor -= 8
        else:                     skor -= 15
    if fcf_verim is not None:
        if fcf_verim >= 8:   skor += 15
        elif fcf_verim >= 3: skor += 10
        elif fcf_verim >= 0: skor += 2
        else:                skor -= 12
    # Likidite (düzeltildi: 50-200 mn TL artık ceza almıyor)
    if tl_hacim_ort >= 200:  skor += 5
    elif tl_hacim_ort >= 50: skor += 0
    elif tl_hacim_ort >= 10: skor -= 10
    else:                    skor -= 25
    if son3 <= -5:   skor -= 4
    elif son3 <= -3: skor -= 2
    if KAP_PUANA_KAT:
        if kap_islem_sayisi >= 3:   skor += 12
        elif kap_islem_sayisi >= 1: skor += 6
    tavan5 = zirveden <= -60 and aylik < 0      # sert düşen hisse: skor en fazla 5
    # ↑↑↑ hücre 2'den aynen ↑↑↑
    _kap = hist['Close']
    tr = pd.concat([hist['High'] - hist['Low'], (hist['High'] - _kap.shift()).abs(),
                    (hist['Low'] - _kap.shift()).abs()], axis=1).max(axis=1)
    sonuc_d = {'Sembol': h, 'Son Fiyat': round(son_fiyat, 2), 'Toplam Skor': skor,
               'Güncel F/K': round(guncel_fk, 1) if guncel_fk > 0 else None,
               'İSY Sektör': SEKTOR.get(h, ''),
               'Hedef (Kötü)': None, 'Hedef (Ort.)': None, 'Hedef (İyi)': None, 'Hedefe Uzaklık (%)': 0.0,
               '_tahmini_hbk': tahmini_hbk, '_fk_min_t': fk_min, '_fk_med_t': fk_medyan,
               '_fk_max_t': fk_max, '_fk_med_ham': fk_med_ham,
               '_veri_uyari': veri_uyari, '_tavan5': tavan5,
               'ATR': float(tr.rolling(14).mean().iloc[-1]), 'TL Hacim (mn)': tl_hacim_ort,
               # ── hücre 2 bileşenleri (ham değerler) ──
               'f_zirve': _f(zirveden), 'f_aylik': _f(aylik), 'f_haftalik': _f(haftalik),
               'f_hacim': _f(hacim_artis), 'f_reel_buyume': None if buyume_yok else _f(reel_buyume),
               'f_tutarlilik': _TUTARLILIK.get(buyume_tutarli_skor), 'f_peg': _f(peg),
               'f_roe_ort': _f(roe_ort), 'f_reel_ciro': _f(reel_ciro), 'f_net_borc': _f(net_borc_ozserm),
               'f_faiz': _f(faiz_karsilama), 'f_fcf': _f(fcf_verim),
               'f_likidite': _f(np.log1p(max(tl_hacim_ort, 0))), 'f_son3': _f(son3),
               # ── bilinen ölçütler ──
               'f_mom12': _f((_kap.iloc[-22] / _kap.iloc[-253] - 1) * 100) if len(_kap) >= 253 else None,
               'f_mom6': _f((_kap.iloc[-1] / _kap.iloc[-127] - 1) * 100) if len(_kap) >= 127 else None,
               'f_vol': _f(_kap.pct_change().iloc[-60:].std() * np.sqrt(252) * 100),
               'f_kar_getirisi': _f(100 / guncel_fk) if guncel_fk > 0 else None,
               'f_buyukluk': _f(np.log(son_fiyat * shares))}
    try:
        oz = float(bal.loc['Common Stock Equity'].iloc[0]) * kur(kur_s, bal.columns[0])
        sonuc_d['f_pddd'] = _f(son_fiyat * shares / oz) if oz > 0 else None
        sonuc_d['f_roe'] = _f(net_karlar[0] / oz * 100) if oz > 0 else None
    except Exception:
        sonuc_d['f_pddd'] = sonuc_d['f_roe'] = None
    # ── hücre 4 teknik (son 3 ay) ──
    try:
        tk = teknik_hesapla(hist.iloc[-63:])
        ts, _ = teknik_skor_hesapla(tk)
        sonuc_d.update({'Teknik Skor': ts, 'f_rsi': _f(tk['RSI']),
                        'f_macd': _f(tk['MACD_Histogram'] / tk['Fiyat'] * 100),
                        'f_ma20': _f((tk['Fiyat'] / tk['MA20'] - 1) * 100),
                        'f_ma50': _f((tk['Fiyat'] / tk['MA50'] - 1) * 100) if tk['MA50'] else None,
                        'f_bb': _f(tk['BB_Yuzde']), 'f_hacim_orani': _f(tk['Hacim_Orani'])})
    except Exception:
        sonuc_d['Teknik Skor'] = None
    return sonuc_d

def _sessiz(*a, **k):
    pass

def degerleme(df):
    """Hücre 2'nin 2. aşaması (sektör F/K + hedef fiyat + değerleme puanı), aynen."""
    # ── HEDEF FİYAT (2. aşama): kendi geçmiş F/K'sı + bugünkü sektör F/K'sı ──
    # Sektör F/K'sı: bugünkü F/K ile geçmiş F/K medyanlarından düşük olanı (temkinli)
    def _fk_istat(seri):
        seri = pd.to_numeric(seri, errors='coerce').dropna()
        return (float(seri.quantile(0.25)), float(seri.median()), float(seri.quantile(0.75)))
    _temiz = df['_veri_uyari'] == ''
    g_bugun  = df[_temiz & df['Güncel F/K'].between(2, 60)]
    g_gecmis = df[_temiz & pd.to_numeric(df['_fk_med_ham'], errors='coerce').between(2, 60)]
    def _sektor_istat(grup_b, grup_g, alt=SEKTOR_MIN_HISSE):
        adaylar = []
        if len(grup_b) >= alt: adaylar.append(_fk_istat(grup_b['Güncel F/K']))
        if len(grup_g) >= alt: adaylar.append(_fk_istat(grup_g['_fk_med_ham']))
        if not adaylar:
            return None
        return tuple(min(a[j] for a in adaylar) for j in range(3))
    PIYASA_FK = _sektor_istat(g_bugun, g_gecmis, 1) or (7.0, 10.0, 14.0)
    sek = {}
    for ad in set(df['İSY Sektör'].dropna()):
        if not ad:
            continue
        st = _sektor_istat(g_bugun[g_bugun['İSY Sektör'] == ad], g_gecmis[g_gecmis['İSY Sektör'] == ad])
        if st:
            sek[ad] = st
    _sessiz(f'\n📐 Piyasa F/K (bugünkü ve geçmişin düşüğü): alt çeyrek {PIYASA_FK[0]:.1f} | medyan {PIYASA_FK[1]:.1f}'
          f' | üst çeyrek {PIYASA_FK[2]:.1f} | kendi sektörüyle karşılaştırılan sektör: {len(sek)}')
    if 'Çimento' in sek:
        _sessiz(f"   örnek: Çimento sektörü F/K medyanı {sek['Çimento'][1]:.1f}")

    def _ort(*xs):
        xs = [float(x) for x in xs if x is not None and pd.notna(x) and x > 0]
        return float(np.mean(xs)) if xs else None

    for col in ['Reel Potansiyel (%)', 'Sektör F/K Medyan', 'Hedef F/K (Kötü)', 'Hedef F/K (Orta)', 'Hedef F/K (İyi)']:
        df[col] = np.nan
    for i, r in df.iterrows():
        s_p25, s_med, s_p75 = sek.get(r['İSY Sektör'], PIYASA_FK)
        df.at[i, 'Sektör F/K Medyan'] = round(s_med, 1)
        skor = r['Toplam Skor']
        thbk = r['_tahmini_hbk']; son_fiyat = r['Son Fiyat']
        if (not r['_veri_uyari'] and pd.notna(thbk) and thbk > 0
                and pd.notna(r['_fk_min_t']) and r['_fk_min_t']):
            fk_orta = min(_ort(r['_fk_med_t'], s_med), FK_ORT_TAVAN)
            fk_kotu = min(_ort(r['_fk_min_t'], s_p25), FK_MIN_TAVAN, fk_orta)
            fk_iyi  = max(min(_ort(r['_fk_max_t'], s_p75), FK_IYI_TAVAN), fk_orta)
            hedef_kotu = min(thbk * fk_kotu, son_fiyat * 3)
            hedef_ort  = min(thbk * fk_orta, son_fiyat * 3)
            hedef_iyi  = min(thbk * fk_iyi,  son_fiyat * 3)
            df.at[i, 'Hedef (Kötü)'] = round(hedef_kotu, 2)
            df.at[i, 'Hedef (Ort.)'] = round(hedef_ort, 2)
            df.at[i, 'Hedef (İyi)']  = round(hedef_iyi, 2)
            df.at[i, 'Hedefe Uzaklık (%)'] = round((hedef_ort - son_fiyat) / son_fiyat * 100, 1)
            df.at[i, 'Hedef F/K (Kötü)'] = round(fk_kotu, 1)
            df.at[i, 'Hedef F/K (Orta)'] = round(fk_orta, 1)
            df.at[i, 'Hedef F/K (İyi)']  = round(fk_iyi, 1)
            # Puan: 12 ay sonraki hedefler bugünkü paraya çevrilip fiyatla karşılaştırılıyor.
            # Sadece enflasyon kadar artış beklenen hisse "ucuz" sayılmıyor.
            ind = 1 + BEKLENEN_ENFLASYON if ENFLASYON_DUZELT else 1
            df.at[i, 'Reel Potansiyel (%)'] = round((hedef_ort / ind - son_fiyat) / son_fiyat * 100, 1)
            if son_fiyat < hedef_kotu / ind:   skor += 30
            elif son_fiyat < hedef_ort / ind:  skor += 22
            elif son_fiyat < hedef_iyi / ind:  skor += 10
            else:                              skor -= 15
        if r['_tavan5']:
            skor = min(skor, 5)
        df.at[i, 'Toplam Skor'] = int(skor)
    return df

# (sütun, ad, sistemin varsaydığı yön: +1 büyüğü iyi / −1 küçüğü iyi, nereden)
FAKTORLER = [
    ('f_toplam',       'TOPLAM SKOR (hücre 2)',                +1, 'Hücre 2'),
    ('f_zirve',        'Zirveye yakınlık (52 hafta)',          +1, 'Hücre 2'),
    ('f_aylik',        'Son 1 ay getiri',                      +1, 'Hücre 2'),
    ('f_haftalik',     'Son 1 hafta getiri',                   +1, 'Hücre 2'),
    ('f_son3',         'Son 3 gün getiri',                     +1, 'Hücre 2'),
    ('f_hacim',        'Hacim artışı (20g / önceki 40g)',      +1, 'Hücre 2'),
    ('f_reel_buyume',  'Reel kâr büyümesi',                    +1, 'Hücre 2'),
    ('f_tutarlilik',   'Büyüme tutarlılığı',                   +1, 'Hücre 2'),
    ('f_peg',          'PEG (düşüğü iyi)',                     -1, 'Hücre 2'),
    ('f_roe_ort',      'ROE ortalaması',                       +1, 'Hücre 2'),
    ('f_reel_ciro',    'Reel ciro büyümesi',                   +1, 'Hücre 2'),
    ('f_net_borc',     'Net borç / özsermaye (düşüğü iyi)',    -1, 'Hücre 2'),
    ('f_faiz',         'Faiz karşılama',                       +1, 'Hücre 2'),
    ('f_fcf',          'FCF verimi',                           +1, 'Hücre 2'),
    ('f_likidite',     'Likidite (TL hacim)',                  +1, 'Hücre 2'),
    ('f_reel_pot',     'Hedef fiyata reel potansiyel',         +1, 'Hücre 2'),
    ('f_teknik',       'TEKNİK SKOR (hücre 4)',                +1, 'Hücre 4'),
    ('f_kombine',      'KOMBİNE SKOR (hücre 2 + 4)',           +1, 'Hücre 4'),
    ('f_rsi',          'RSI (düşüğü iyi)',                     -1, 'Hücre 4'),
    ('f_macd',         'MACD histogram',                       +1, 'Hücre 4'),
    ('f_ma20',         'Fiyat / 20 günlük ortalama',           +1, 'Hücre 4'),
    ('f_ma50',         'Fiyat / 50 günlük ortalama',           +1, 'Hücre 4'),
    ('f_bb',           'Bollinger konumu (düşüğü iyi)',        -1, 'Hücre 4'),
    ('f_hacim_orani',  'Hacim oranı (3g / 20g)',               +1, 'Hücre 4'),
    ('f_mom12',        'Momentum: 12 ay (son ay hariç)',       +1, 'Bilinen'),
    ('f_mom6',         'Momentum: 6 ay',                       +1, 'Bilinen'),
    ('f_vol',          'Oynaklık 60 gün (düşüğü iyi)',         -1, 'Bilinen'),
    ('f_kar_getirisi', 'Kâr getirisi = 1 / (F/K)',             +1, 'Bilinen'),
    ('f_pddd',         'PD/DD (düşüğü iyi)',                   -1, 'Bilinen'),
    ('f_roe',          'ROE son yıl',                          +1, 'Bilinen'),
    ('f_buyukluk',     'Büyüklük (piyasa değeri)',             +1, 'Bilinen'),
]
BILESIK_DISI = {'f_toplam', 'f_teknik', 'f_kombine'}   # bunlar zaten bileşik, seçime girmez

# ════════════════════════════════════════════════════════════════
#  3) AYLIK TEST: her ayın ilk işlem günü puanla, 20 ve 60 gün sonrasına bak
# ════════════════════════════════════════════════════════════════
takvim = XU.index
ilkler = pd.Series(takvim, index=takvim).groupby([takvim.year, takvim.month]).first().tolist()
test_tarihleri = [d for d in ilkler if d >= pd.Timestamp(ILK_TEST)
                  and takvim.get_loc(d) + 1 + TUTMA_GUN < len(takvim)]
print(f"\n🗓️  {len(test_tarihleri)} test tarihi: {test_tarihleri[0]:%d.%m.%Y} → {test_tarihleri[-1]:%d.%m.%Y}")

def fiyat_at(hist, t):
    o = hist['Close'][hist.index <= t]
    if len(o) == 0 or (t - o.index[-1]).days > 7:
        return None
    return float(o.iloc[-1])

def cikis_sim(hist, giris_t, cikis_t, giris, stop, hedef):
    """Gün gün: önce stop, sonra hedef (aynı gün ikisi de olursa stop sayılır, temkinli)."""
    p = hist[(hist.index > giris_t) & (hist.index <= cikis_t)]
    if len(p) == 0:
        return 0.0, 'Süre'
    O, H, L = p['Open'].values, p['High'].values, p['Low'].values
    for i in range(len(p)):
        if stop is not None and L[i] <= stop:
            return min(O[i], stop) / giris - 1, 'Stop'
        if hedef is not None and H[i] >= hedef:
            return max(O[i], hedef) / giris - 1, 'Hedef'
    return float(p['Close'].iloc[-1]) / giris - 1, 'Süre'

CIKIS_KURALLARI = [          # (ad, stop fonksiyonu, hedef türü)
    (f'{TUTMA_GUN} gün tut (stopsuz)', None,                       None),
    ('Stop −%15',                      lambda g, a: g * 0.85,      None),
    ('Stop −%20',                      lambda g, a: g * 0.80,      None),
    ('ATR stop 2,5× (sistemdeki)',     lambda g, a: g - 2.5 * a,   None),
    ('ATR stop 3,5×',                  lambda g, a: g - 3.5 * a,   None),
    ('ATR stop 4×',                    lambda g, a: g - 4.0 * a,   None),
    ('Stop −%15 + Hedef +%30',         lambda g, a: g * 0.85,      'yuzde30'),
    ('ATR 2,5× + Hedef (Ort.)',        lambda g, a: g - 2.5 * a,   'sistem'),
]

kayitlar = []
for tarih in tqdm(test_tarihleri, desc='Test'):
    satirlar = []
    for h, v in BT_VERI.items():
        if not v:
            continue
        try:
            r = puan_hesapla(h, v, tarih)
            if r:
                satirlar.append(r)
        except Exception:
            continue
    if len(satirlar) < 20:
        continue
    df = degerleme(pd.DataFrame(satirlar))
    pos = takvim.get_loc(tarih)
    giris_t, kisa_t, cikis_t = takvim[pos + 1], takvim[pos + 1 + KISA_GUN], takvim[pos + 1 + TUTMA_GUN]
    xc = XU['Close']
    xu20 = float(xc.loc[kisa_t] / xc.loc[giris_t] - 1) * 100
    xu60 = float(xc.loc[cikis_t] / xc.loc[giris_t] - 1) * 100
    for _, r in df.iterrows():
        hist = BT_VERI[r['Sembol']]['hist']
        g, k20, c60 = fiyat_at(hist, giris_t), fiyat_at(hist, kisa_t), fiyat_at(hist, cikis_t)
        if not g or not k20 or not c60:
            continue
        kayit = {'Tarih': tarih, 'Kod': r['Sembol'], 'Skor': int(r['Toplam Skor']),
                 'Teknik Skor': r.get('Teknik Skor'), 'Giriş': round(g, 2),
                 'TL Hacim (mn)': r['TL Hacim (mn)'],
                 'G20 (%)': (k20 / g - 1) * 100, 'BIST100 20g (%)': xu20,
                 'Getiri (%)': (c60 / g - 1) * 100, 'BIST100 (%)': xu60,
                 'f_toplam': float(r['Toplam Skor']),
                 'f_teknik': _f(r.get('Teknik Skor')),
                 'f_kombine': _f(r['Toplam Skor'] + r['Teknik Skor']) if pd.notna(r.get('Teknik Skor')) else None,
                 'f_reel_pot': _f(r.get('Reel Potansiyel (%)'))}
        for kol, _, _, _ in FAKTORLER:
            if kol not in kayit:
                kayit[kol] = r.get(kol)
        kayit['Fark (%)'] = kayit['Getiri (%)'] - xu60
        if r['Toplam Skor'] >= AL_ESIK:
            hed_sis = r['Hedef (Ort.)'] if pd.notna(r['Hedef (Ort.)']) and r['Hedef (Ort.)'] > g else None
            for ad, stop_f, hed_t in CIKIS_KURALLARI:
                stop = stop_f(g, r['ATR']) if stop_f else None
                hed = g * 1.30 if hed_t == 'yuzde30' else hed_sis if hed_t == 'sistem' else None
                getiri, sonuc = cikis_sim(hist, giris_t, cikis_t, g, stop, hed)
                kayit[f'ÇK|{ad}'] = getiri * 100
                kayit[f'ÇS|{ad}'] = sonuc
        kayitlar.append(kayit)

bt = pd.DataFrame(kayitlar)
if bt.empty:
    raise SystemExit('❌ Test için yeterli veri çıkmadı.')
bt['Tüm Hisseler (%)'] = bt.groupby('Tarih')['Getiri (%)'].transform('mean')
bt['Tüme Fark (%)'] = bt['Getiri (%)'] - bt['Tüm Hisseler (%)']
bt['Tüm 20g (%)'] = bt.groupby('Tarih')['G20 (%)'].transform('mean')
aylar = sorted(bt['Tarih'].unique())
ilk_yari, son_yari = aylar[:len(aylar) // 2], aylar[len(aylar) // 2:]
print(f"✅ {len(bt)} gözlem | {len(aylar)} ay | {bt['Kod'].nunique()} hisse "
      f"| seçim dönemi {len(ilk_yari)} ay, doğrulama dönemi {len(son_yari)} ay")

# ════════════════════════════════════════════════════════════════
#  4) FAKTÖR TESTİ: her ölçüt tek tek
#     IC = o ay ölçüte göre sıralama ile sonraki 20 günün getiri sıralaması arasındaki
#     uyum (−1 … +1). Sürekli 0'ın üstündeyse ölçüt işe yarıyor demektir.
# ════════════════════════════════════════════════════════════════
def faktor_istat(df, kol, yon, aylar_, getiri='G20 (%)'):
    ics, yay20, yay60 = [], [], []
    for d in aylar_:
        g = df[df['Tarih'] == d]
        x = pd.to_numeric(g[kol], errors='coerce') * yon
        ok = x.notna() & g[getiri].notna()
        if ok.sum() < 30:
            continue
        xs, ys, y60 = x[ok], g.loc[ok, getiri], g.loc[ok, 'Getiri (%)']
        ic = xs.rank().corr(ys.rank())
        if pd.notna(ic):
            ics.append(ic)
        q = pd.qcut(xs.rank(method='first'), 5, labels=False)
        yay20.append(ys[q == 4].mean() - ys[q == 0].mean())
        yay60.append(y60[q == 4].mean() - y60[q == 0].mean())
    n = len(ics)
    if n < 3:
        return None
    ic = np.array(ics)
    t = ic.mean() / (ic.std(ddof=1) / np.sqrt(n)) if ic.std(ddof=1) > 0 else 0.0
    return {'Ay': n, 'Ort. IC': ic.mean(), 'IC>0 Ay (%)': (ic > 0).mean() * 100, 't': t,
            'Yay20': np.nanmean(yay20), 'Yay20>0 (%)': (np.array(yay20) > 0).mean() * 100,
            'Yay60': np.nanmean(yay60)}

def hukum(t, ic1, ic2):
    if ic1 is None or ic2 is None:
        return '⚪ Veri az'
    if t >= 2 and ic1 > 0 and ic2 > 0:   return '✅ İşe yarıyor'
    if t <= -2 and ic1 < 0 and ic2 < 0:  return '🔴 Ters çalışıyor'
    if t >= 1 and ic1 > 0 and ic2 > 0:   return '🟡 Zayıf olumlu'
    if t <= -1 and ic1 < 0 and ic2 < 0:  return '🟠 Zayıf ters'
    return '⚪ Etkisiz'

frows = []
for kol, ad, yon, kaynak in FAKTORLER:
    a = faktor_istat(bt, kol, yon, aylar)
    if a is None:
        continue
    b1 = faktor_istat(bt, kol, yon, ilk_yari); b2 = faktor_istat(bt, kol, yon, son_yari)
    ic1 = b1['Ort. IC'] if b1 else None; ic2 = b2['Ort. IC'] if b2 else None
    frows.append({'Nereden': kaynak, 'Ölçüt': ad, 'Ay': a['Ay'],
                  'Ort. IC': round(a['Ort. IC'], 3), 'IC>0 Ay (%)': round(a['IC>0 Ay (%)'], 0),
                  't': round(a['t'], 2),
                  'İlk Yarı IC': round(ic1, 3) if ic1 is not None else None,
                  'Son Yarı IC': round(ic2, 3) if ic2 is not None else None,
                  'En İyi %20 − En Kötü %20, 20g (%)': round(a['Yay20'], 2),
                  'Bu Fark >0 Ay (%)': round(a['Yay20>0 (%)'], 0),
                  'Aynı Fark, 60g (%)': round(a['Yay60'], 2),
                  'Sonuç': hukum(a['t'], ic1, ic2), '_kol': kol, '_yon': yon,
                  '_t1': b1['t'] if b1 else 0.0, '_ic1': ic1})
t_faktor = pd.DataFrame(frows)
_sira = {'✅ İşe yarıyor': 0, '🟡 Zayıf olumlu': 1, '⚪ Etkisiz': 2, '⚪ Veri az': 3, '🟠 Zayıf ters': 4, '🔴 Ters çalışıyor': 5}
t_faktor = t_faktor.sort_values(['Nereden', 't'], ascending=[True, False],
                                key=lambda s: s.map({'Hücre 2': 0, 'Hücre 4': 1, 'Bilinen': 2}) if s.name == 'Nereden' else s)

# ════════════════════════════════════════════════════════════════
#  5) BİLEŞİK SKOR — ilk yarıda seç, son yarıda test et (geçmişe uydurmaya karşı)
# ════════════════════════════════════════════════════════════════
secilen = []
for _, f in t_faktor.iterrows():
    if f['_kol'] in BILESIK_DISI or f['_ic1'] is None:
        continue
    if abs(f['_t1']) >= SECIM_T:
        secilen.append((f['_kol'], f['_yon'] * (1 if f['_ic1'] > 0 else -1), f['Ölçüt'], f['_ic1'] > 0))
def bilesik(g):
    parca = []
    for kol, isaret, _, _ in secilen:
        x = pd.to_numeric(g[kol], errors='coerce') * isaret
        parca.append(x.rank(pct=True))
    if not parca:
        return pd.Series(np.nan, index=g.index)
    m = pd.concat(parca, axis=1)
    return m.mean(axis=1).where(m.notna().sum(axis=1) >= max(1, len(parca) // 2))
bt['Bileşik'] = bt.groupby('Tarih', group_keys=False).apply(bilesik)

def portfoy_testi(kolon, aylar_, n=PORTFOY_SAYISI):
    s = []
    for d in aylar_:
        g = bt[(bt['Tarih'] == d) & (bt['TL Hacim (mn)'] >= MIN_TL_HACIM)].dropna(subset=[kolon])
        if len(g) < n * 2:
            continue
        sec = g.nlargest(n, kolon)
        s.append({'Tarih': d, 'P20': sec['G20 (%)'].mean(), 'U20': g['G20 (%)'].mean(),
                  'X20': g['BIST100 20g (%)'].iloc[0], 'P60': sec['Getiri (%)'].mean(),
                  'U60': g['Getiri (%)'].mean(), 'X60': g['BIST100 (%)'].iloc[0]})
    return pd.DataFrame(s)

prow = []
for kolon, ad in [('Bileşik', f'Bileşik skor — en iyi {PORTFOY_SAYISI} hisse'),
                  ('f_toplam', f'Toplam skor (hücre 2) — en iyi {PORTFOY_SAYISI} hisse'),
                  ('f_kombine', f'Kombine skor (2+4) — en iyi {PORTFOY_SAYISI} hisse')]:
    for donem, aylar_ in [('Seçim dönemi (ilk yarı)', ilk_yari), ('DOĞRULAMA (son yarı)', son_yari)]:
        p = portfoy_testi(kolon, aylar_)
        if p.empty:
            continue
        prow.append({'Portföy': ad, 'Dönem': donem, 'Ay': len(p),
                     'Portföy 20g (%)': round(p['P20'].mean(), 2), 'Tüm Hisseler 20g (%)': round(p['U20'].mean(), 2),
                     'BIST100 20g (%)': round(p['X20'].mean(), 2),
                     'Tüm Hisseleri Yendiği Ay (%)': round((p['P20'] > p['U20']).mean() * 100, 0),
                     'BIST100\'ü Yendiği Ay (%)': round((p['P20'] > p['X20']).mean() * 100, 0),
                     'Portföy 60g (%)': round(p['P60'].mean(), 2), 'BIST100 60g (%)': round(p['X60'].mean(), 2)})
t_portfoy = pd.DataFrame(prow)
t_secilen = pd.DataFrame([{'Seçilen Ölçüt': ad, 'Yön': 'sistemin varsaydığı gibi' if ayni else 'TERS yönde',
                           'Seçim Dönemi IC': round(t_faktor.set_index('_kol').loc[k, '_ic1'], 3)}
                          for k, _, ad, ayni in secilen]) if secilen else pd.DataFrame([{'Seçilen Ölçüt': 'yok (ilk yarıda hiçbiri yeterince güçlü değil)'}])

# ════════════════════════════════════════════════════════════════
#  6) SKOR GRUPLARI, ÇIKIŞ KURALLARI, MAKRO
# ════════════════════════════════════════════════════════════════
def ozet_satir(g):
    return pd.Series({
        'Gözlem': int(len(g)), 'Ay': int(g['Tarih'].nunique()),
        'Ort. Getiri 60g (%)': round(g['Getiri (%)'].mean(), 1),
        'Medyan (%)': round(g['Getiri (%)'].median(), 1),
        'Ort. BIST100 (%)': round(g['BIST100 (%)'].mean(), 1),
        'BIST100\'e Fark (%)': round(g['Fark (%)'].mean(), 1),
        'Endeksi Yenme (%)': round((g['Fark (%)'] > 0).mean() * 100, 1),
        'Tüm Hisselere Fark (%)': round(g['Tüme Fark (%)'].mean(), 1),
    })
def skor_grubu(s):
    if s >= 80: return '1) 80 ve üstü'
    if s >= AL_ESIK: return f'2) {AL_ESIK}-79 (AL)'
    if s >= 35: return f'3) 35-{AL_ESIK - 1}'
    return '4) 35 altı'
bt['Grup'] = bt['Skor'].apply(skor_grubu)
t_grup = bt.groupby('Grup').apply(ozet_satir).reset_index()
al = bt[bt['Skor'] >= AL_ESIK]
t_grup = pd.concat([t_grup, pd.DataFrame([{'Grup': f'AL listesi toplam (≥{AL_ESIK})', **ozet_satir(al).to_dict()},
                                          {'Grup': 'Tüm hisseler', **ozet_satir(bt).to_dict()}])], ignore_index=True)
t_grup[['Gözlem', 'Ay']] = t_grup[['Gözlem', 'Ay']].astype(int)

cik = []
for ad, _, _ in CIKIS_KURALLARI:
    g = pd.to_numeric(al.get(f'ÇK|{ad}'), errors='coerce').dropna() if f'ÇK|{ad}' in al else pd.Series(dtype=float)
    s_ = al[f'ÇS|{ad}'].dropna() if f'ÇS|{ad}' in al else pd.Series(dtype=object)
    if len(g) == 0:
        continue
    cik.append({'Çıkış Kuralı (AL listesi)': ad, 'İşlem': len(g), 'Ort. Getiri (%)': round(g.mean(), 1),
                'Medyan (%)': round(g.median(), 1), 'Kâr Eden (%)': round((g > 0).mean() * 100, 1),
                'En Kötü %10 Ort. (%)': round(g[g <= g.quantile(0.1)].mean(), 1),
                'Stop Olan (%)': round((s_ == 'Stop').mean() * 100, 1),
                'Hedefe Ulaşan (%)': round((s_ == 'Hedef').mean() * 100, 1)})
t_cikis = pd.DataFrame(cik)

seriler = {ad: s['Close'] for ad, s in BT_MAKRO.items() if s is not None}
mgunler = [d for d in takvim if d >= pd.Timestamp(ILK_TEST) - pd.Timedelta(days=400)
           and takvim.get_loc(d) + 1 + TUTMA_GUN < len(takvim)]
makro_g = {}
for d in tqdm(mgunler, desc='Makro'):
    try:
        makro_g[d] = makro_skor_hesapla(gosterge_hesapla(seriler, d))[0]
    except Exception:
        pass
mk = pd.DataFrame({'Makro (gün)': pd.Series(makro_g)})
mk['Makro (3 gün ort.)'] = mk['Makro (gün)'].rolling(3).mean().round()
xcl = XU['Close']
mk['BIST100 60g (%)'] = [float(xcl.iloc[takvim.get_loc(d) + 1 + TUTMA_GUN] / xcl.iloc[takvim.get_loc(d) + 1] - 1) * 100
                         for d in mk.index]
def bolge(m):
    if pd.isna(m): return None
    if m < MAKRO_ESIK: return f'1) {MAKRO_ESIK} altı (RİSK YÜKSEK)'
    if m < -5:         return f'2) {MAKRO_ESIK} … −5 (riskli)'
    if m < 0:          return '3) −5 … 0'
    return '4) 0 ve üstü'
mak = []
for kol, ad in [('Makro (gün)', 'Bugünkü skor (MAKRO_GUN=1)'), ('Makro (3 gün ort.)', '3 gün ort. (MAKRO_GUN=3)')]:
    for bl, g in mk.groupby(mk[kol].apply(bolge)):
        mak.append({'Skor Türü': ad, 'Makro Bölgesi': bl, 'Gün': len(g),
                    'Bu Bölge Kaç Ayrı Dönem': int((g.index.to_series().diff().dt.days.fillna(99) > 5).sum()),
                    'BIST100 Sonraki 60g Ort. (%)': round(g['BIST100 60g (%)'].mean(), 1),
                    'BIST100 Yükseldi (%)': round((g['BIST100 60g (%)'] > 0).mean() * 100, 1)})
t_makro = pd.DataFrame(mak)

# ════════════════════════════════════════════════════════════════
#  7) KAYIT DEĞERLENDİRMESİ (hücre 7'nin sistem_kayit.csv'si): ileriye dönük gerçek sonuç
# ════════════════════════════════════════════════════════════════
t_kayit = pd.DataFrame(); kayit_not = 'Kayıt dosyası yok (hücre 7 v4.8 çalıştıkça oluşur).'
if os.path.exists(KAYIT_YOLU):
    kd = pd.read_csv(KAYIT_YOLU)
    kd['Tarih'] = pd.to_datetime(kd['Tarih'])
    sat = []
    for _, r in kd.iterrows():
        v = BT_VERI.get(r['Kod'])
        if not v:
            continue
        sonraki = takvim[takvim > r['Tarih']]
        if len(sonraki) < KISA_GUN + 1:
            continue                               # 20 iş günü dolmadı
        g = fiyat_at(v['hist'], sonraki[0]); k = fiyat_at(v['hist'], sonraki[KISA_GUN])
        if not g or not k:
            continue
        x20 = float(xcl.loc[sonraki[KISA_GUN]] / xcl.loc[sonraki[0]] - 1) * 100
        sat.append({**r.to_dict(), 'G20 (%)': (k / g - 1) * 100, 'BIST100 20g (%)': x20})
    kayit_not = (f"Kayıtta {kd['Tarih'].nunique()} gün / {len(kd)} karar var; 20 iş günü dolan {len(sat)} karar "
                 f"değerlendirildi." + (" Henüz az — 1-2 ay sonra anlamlı olur." if len(sat) < 100 else ""))
    if sat:
        ks = pd.DataFrame(sat); ks['Fark'] = ks['G20 (%)'] - ks['BIST100 20g (%)']
        for kol in ['Haber Kararı', 'Öneri', 'Giriş Kararı', 'Süper AL']:
            if kol not in ks:
                continue
            for deger, g in ks.groupby(kol):
                t_kayit = pd.concat([t_kayit, pd.DataFrame([{
                    'Karar Türü': kol, 'Karar': deger, 'Adet': len(g), 'Gün': g['Tarih'].nunique(),
                    'Sonraki 20g Ort. (%)': round(g['G20 (%)'].mean(), 2),
                    'BIST100 20g (%)': round(g['BIST100 20g (%)'].mean(), 2),
                    'Fark (%)': round(g['Fark'].mean(), 2),
                    'Endeksi Yenme (%)': round((g['Fark'] > 0).mean() * 100, 0)}])], ignore_index=True)

# ════════════════════════════════════════════════════════════════
#  8) SONUÇ YORUMU
# ════════════════════════════════════════════════════════════════
yorum = []
f_ise = t_faktor[t_faktor['Sonuç'] == '✅ İşe yarıyor']['Ölçüt'].tolist()
f_ters = t_faktor[t_faktor['Sonuç'] == '🔴 Ters çalışıyor']['Ölçüt'].tolist()
f_zayif = t_faktor[t_faktor['Sonuç'] == '🟡 Zayıf olumlu']['Ölçüt'].tolist()
yorum.append(f"İşe yarayan ölçütler ({len(f_ise)}): " + (', '.join(f_ise) if f_ise else 'yok'))
if f_zayif:
    yorum.append(f"Zayıf olumlu ({len(f_zayif)}): " + ', '.join(f_zayif))
yorum.append(f"Ters çalışan ölçütler ({len(f_ters)}) — sistem bunlara yanlış yönde puan veriyor: "
             + (', '.join(f_ters) if f_ters else 'yok'))
tp = t_faktor.set_index('_kol')
for k in ['f_toplam', 'f_teknik', 'f_kombine']:
    if k in tp.index:
        yorum.append(f"{tp.loc[k, 'Ölçüt']}: IC {tp.loc[k, 'Ort. IC']:+.3f}, t {tp.loc[k, 't']:+.2f} → {tp.loc[k, 'Sonuç']}")
dg = t_portfoy[t_portfoy['Dönem'].str.startswith('DOĞRULAMA')] if len(t_portfoy) else pd.DataFrame()
if len(dg):
    for _, r in dg.iterrows():
        yorum.append(f"DOĞRULAMA dönemi — {r['Portföy']}: ayda ort. %{r['Portföy 20g (%)']:+.2f} "
                     f"(tüm hisseler %{r['Tüm Hisseler 20g (%)']:+.2f}, BIST100 %{r['BIST100 20g (%)']:+.2f}); "
                     f"tüm hisseleri {int(r['Ay'])} ayın %{r['Tüm Hisseleri Yendiği Ay (%)']:.0f}'inde yendi")
if len(t_cikis):
    en = t_cikis.sort_values('Ort. Getiri (%)', ascending=False).iloc[0]
    yorum.append(f"Çıkış: en yüksek ortalama '{en['Çıkış Kuralı (AL listesi)']}' (%{en['Ort. Getiri (%)']:+.1f}).")
mg = mk[mk['Makro (gün)'] < MAKRO_ESIK]['BIST100 60g (%)'].mean(); md = mk[mk['Makro (gün)'] >= MAKRO_ESIK]['BIST100 60g (%)'].mean()
yorum.append(f"Makro {MAKRO_ESIK} altındayken BIST100 sonraki 60 günde %{mg:+.1f}, diğer günlerde %{md:+.1f} → "
             + ("beklemek işe yaramış." if mg < md else "beklemek işe YARAMAMIŞ."))
yorum.append(kayit_not)

SINIRLAR = [
    "Test edilen: hücre 2'nin puanı ve bileşenleri, hücre 4'ün teknik skoru, bilinen ölçütler, hücre 8'in makro skoru.",
    "Geçmişe dönük test EDİLEMEYEN: haber/AI, KAP, DCF, İsyatırım puanları — bunlar için hücre 7'nin günlük kaydı (7. bölüm) ileriye dönük ölçüyor.",
    f"Bilançolar yıl sonundan {RAPOR_GECIKME_GUN} gün sonra biliniyor sayıldı; çeyrek verisi yok.",
    "Sadece bugün borsada olan hisseler var (sonradan çıkanlar yok) — sonuçlar biraz iyimser olabilir.",
    "IC: bir ölçüte göre sıralama ile sonraki 20 günün getiri sıralaması arasındaki uyum (−1…+1). "
    "BIST'te 0,03-0,05 bile anlamlı sayılır. t ≥ 2 ve iki yarıda da aynı yön → 'işe yarıyor'.",
    f"Bileşik skor: ilk {len(ilk_yari)} ayda |t| ≥ {SECIM_T} olan ölçütler seçildi, sadece son {len(son_yari)} ayda test edildi. "
    f"Portföy: günlük hacmi {MIN_TL_HACIM} mn TL üstü hisselerden en iyi {PORTFOY_SAYISI}, eşit ağırlık, aylık yenileme.",
    "Komisyon ve vergi düşülmedi. Dönem ~2,5 yıl; geçmiş performans geleceği garanti etmez.",
]

print(f"\n{'=' * 65}\n  💡 SONUÇ\n{'=' * 65}")
for y in yorum:
    print('  • ' + y)
gos = t_faktor.drop(columns=[c for c in t_faktor.columns if c.startswith('_')])
print('\n📋 ÖLÇÜTLER\n' + gos[['Nereden', 'Ölçüt', 'Ort. IC', 't', 'İlk Yarı IC', 'Son Yarı IC',
                             'En İyi %20 − En Kötü %20, 20g (%)', 'Sonuç']].to_string(index=False))
print('\n📋 BİLEŞİK SKORA SEÇİLENLER\n' + t_secilen.to_string(index=False))
if len(t_portfoy):
    print('\n📋 PORTFÖY TESTİ\n' + t_portfoy.to_string(index=False))
print('\n📋 ÇIKIŞ KURALLARI\n' + t_cikis.to_string(index=False))
print('\n📋 MAKRO\n' + t_makro.to_string(index=False))
if len(t_kayit):
    print('\n📋 KAYIT (ileriye dönük)\n' + t_kayit.to_string(index=False))

# ════════════════════════════════════════════════════════════════
#  9) EXCEL + CSV
# ════════════════════════════════════════════════════════════════
def tablo_yaz(ws, satir, baslik, df):
    ws.cell(satir, 1, baslik).font = Font(bold=True, size=12, color='FF0F3460')
    satir += 1
    for ci, c in enumerate(df.columns, 1):
        x = ws.cell(satir, ci, c)
        x.font = Font(bold=True, color='FFFFFFFF'); x.fill = PatternFill('solid', fgColor='FF1A1A2E')
        x.alignment = Alignment(horizontal='center', wrap_text=True)
    for _, r in df.iterrows():
        satir += 1
        for ci, c in enumerate(df.columns, 1):
            v = r[c]
            if isinstance(v, (pd.Timestamp, datetime)):
                v = v.strftime('%d.%m.%Y')
            elif isinstance(v, (float, np.floating)) and np.isnan(v):
                v = None
            elif isinstance(v, np.integer):
                v = int(v)
            elif isinstance(v, np.floating):
                v = float(v)
            elif isinstance(v, (bool, np.bool_)):
                v = 'Evet' if v else 'Hayır'
            x = ws.cell(satir, ci, v)
            if isinstance(v, (int, float)) and any(k in c for k in ['Fark', 'Getiri', 'IC', 'Portföy 20g', ' t']):
                x.font = Font(color='FF1B5E20' if v > 0 else 'FFB71C1C' if v < 0 else 'FF000000')
            if c == 'Sonuç' and isinstance(v, str):
                x.font = Font(bold=True, color='FF1B5E20' if '✅' in v else 'FFB71C1C' if '🔴' in v else 'FF8D6E00' if '🟡' in v or '🟠' in v else 'FF555555')
    return satir + 2

wb = load_workbook(EXCEL_YOLU)
for ad in ['📊 Backtest', '📊 Backtest Faktör', '📊 Backtest Ham']:
    if ad in wb.sheetnames:
        del wb[ad]
ws = wb.create_sheet('📊 Backtest')
ws['A1'] = (f"📊 BACKTEST v2 — {datetime.now():%d.%m.%Y %H:%M} | bugünkü kurallar (hücre 2 v2.4, hücre 4, hücre 8) | "
            f"{len(aylar)} ay, {len(bt)} gözlem")
ws['A1'].font = Font(bold=True, size=13, color='FFFFFFFF'); ws['A1'].fill = PatternFill('solid', fgColor='FF0F3460')
s = 3
ws.cell(s, 1, '💡 SONUÇ').font = Font(bold=True, size=12)
for y in yorum:
    s += 1; ws.cell(s, 1, '• ' + y)
s += 2
s = tablo_yaz(ws, s, 'Bileşik skor ve portföy testi (asıl ölçü: DOĞRULAMA satırları)', t_portfoy) if len(t_portfoy) else s
s = tablo_yaz(ws, s, 'Bileşik skora seçilen ölçütler (sadece ilk yarının verisiyle seçildi)', t_secilen)
s = tablo_yaz(ws, s, 'Skor gruplarına göre (sonraki 60 işlem günü)', t_grup)
s = tablo_yaz(ws, s, 'Çıkış kuralları (AL listesi, 60 işlem günü içinde)', t_cikis)
s = tablo_yaz(ws, s, 'Makro skor ve BIST100\'ün sonraki 60 günü', t_makro)
if len(t_kayit):
    s = tablo_yaz(ws, s, 'Hücre 7 kaydı: gerçek kararların sonraki 20 günü (ileriye dönük)', t_kayit)
else:
    ws.cell(s, 1, 'Hücre 7 kaydı: ' + kayit_not); s += 2
ws.cell(s, 1, '⚠️ SINIRLAR VE AÇIKLAMALAR').font = Font(bold=True, size=12)
for y in SINIRLAR:
    s += 1; ws.cell(s, 1, '• ' + y)
for i, w in enumerate([40, 22, 12, 14, 14, 14, 14, 14, 14, 14, 14], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

wsf = wb.create_sheet('📊 Backtest Faktör')
tablo_yaz(wsf, 1, 'Her ölçüt tek başına: sonraki 20 işlem günü (IC ve en iyi/en kötü %20 farkı)', gos)
for i, w in enumerate([10, 36, 6, 9, 10, 7, 10, 10, 16, 12, 12, 18], 1):
    wsf.column_dimensions[get_column_letter(i)].width = w
wsf.freeze_panes = 'C3'

ws2 = wb.create_sheet('📊 Backtest Ham')
ham = bt[['Tarih', 'Kod', 'Skor', 'Teknik Skor', 'Bileşik', 'G20 (%)', 'BIST100 20g (%)',
          'Getiri (%)', 'BIST100 (%)', 'Fark (%)']].copy()
for c in ham.columns:
    if ham[c].dtype.kind == 'f':
        ham[c] = ham[c].round(2)
tablo_yaz(ws2, 1, f'Tüm gözlemler ({len(ham)}) — ölçütlerin tamamı Drive\'da backtest_ham.csv', ham)
ws2.freeze_panes = 'A3'
wb.save(EXCEL_YOLU)
try:
    bt.drop(columns=['Grup']).to_csv(HAM_CSV, index=False)
except Exception as e:
    print(f"⚠️ CSV yazılamadı: {e}")
print(f"\n✅ Excel: '📊 Backtest' (özet), '📊 Backtest Faktör' (ölçütler), '📊 Backtest Ham' | CSV: backtest_ham.csv")
print("🎉 BACKTEST TAMAMLANDI — çıktının SONUÇ bölümünü ve backtest_ham.csv'yi Claude'a gönder")
