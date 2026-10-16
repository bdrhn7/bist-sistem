# ═══════════════════════════════════════════════════════════════
#  HÜCRE 4 — TEKNİK ANALİZ + EXCEL SAYFALARI  (v2.1 — 16.10.2026)
#  v2 değişiklikleri:
#   - Seans açıkken bugünün yarım günlük verisi kullanılmıyor
#     (RSI, hacim, MACD dünkü kapanışa göre hesaplanıyor)
#   - Finansal Sağlık ve Borç sayfalarında hiç görünmeyen sütunlar düzeltildi
#   - Sektör Analizi sayfası İsyatırım sektörüyle yeniden çalışıyor
#   - KAP sayfası hücre 2'nin verisini kullanıyor (iki ayrı kaynak yok)
#   - AL listesi ya da teknik onaylı hisse boşsa çökme düzeltildi
#   - Kalite ve Tüm Hisseler sayfalarına reel büyüme sütunları eklendi
#   - Teknik onaylı liste '✅ Teknik Onay' sayfasına yazılıyor
#     ('🎯 Kombine AL' sayfasını artık sadece hücre 7 yazıyor)
#   v2.1: Yahoo fiyatı boş (NaN) satır gönderirse o satır atılıyor (16.10 sabahı yaşandı)
# ═══════════════════════════════════════════════════════════════
import yfinance as yf
import pandas as pd
import numpy as np
from datetime import datetime
import time
import warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

# ── AYARLAR ──────────────────────────────────────────────
SISTEM_YOLU = '/content/drive/MyDrive/isyatirim/BIST_Sistem.xlsx'
EXCEL_YOLU  = '/content/drive/MyDrive/isyatirim/BIST_Tam_KAP_v9_11mart.xlsx'
DRIVE       = '/content/drive/MyDrive/isyatirim'

BORDER2 = Border(left=Side(style='thin',color='1E3A5F'), right=Side(style='thin',color='1E3A5F'),
                 top=Side(style='thin',color='1E3A5F'), bottom=Side(style='thin',color='1E3A5F'))

def rsi_hesapla(close, period=14):
    delta = close.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.ewm(com=period-1, min_periods=period).mean()
    avg_loss = loss.ewm(com=period-1, min_periods=period).mean()
    rs = avg_gain / avg_loss.replace(0, np.nan)
    return 100 - (100 / (1 + rs))

def macd_hesapla(close, fast=12, slow=26, signal=9):
    ema_fast = close.ewm(span=fast, adjust=False).mean()
    ema_slow = close.ewm(span=slow, adjust=False).mean()
    macd_line = ema_fast - ema_slow
    signal_line = macd_line.ewm(span=signal, adjust=False).mean()
    histogram = macd_line - signal_line
    return macd_line, signal_line, histogram

def tamamlanmis_gunler(df_h):
    """Fiyatı boş gelen satırları atar; seans kapanmadan (18:15'ten önce) bugünün
    satırlarını (yarım ya da boş mum) çıkarır."""
    try:
        if 'Close' in df_h.columns:
            df_h = df_h[df_h['Close'].notna() & (df_h['Close'] > 0)]
        simdi = datetime.now()
        if len(df_h) and (simdi.hour, simdi.minute) < (18, 15):
            df_h = df_h[[d.date() != simdi.date() for d in df_h.index]]
    except: pass
    return df_h

def direnc_destek_hesapla(df_hist):
    try:
        fiyat = df_hist['Close'].iloc[-1]
        yuksek_3ay = df_hist['High'].max()
        dusuk_3ay = df_hist['Low'].min()
        son_gun = df_hist.iloc[-2]
        pivot = (son_gun['High'] + son_gun['Low'] + son_gun['Close']) / 3
        r1 = 2 * pivot - son_gun['Low']
        r2 = pivot + (son_gun['High'] - son_gun['Low'])
        s1 = 2 * pivot - son_gun['High']
        s2 = pivot - (son_gun['High'] - son_gun['Low'])
        fib_range = yuksek_3ay - dusuk_3ay
        fib_382 = yuksek_3ay - fib_range * 0.382
        fib_500 = yuksek_3ay - fib_range * 0.500
        fib_618 = yuksek_3ay - fib_range * 0.618
        destekler = [d for d in [s1, s2, fib_382, fib_500, fib_618, dusuk_3ay] if d < fiyat]
        direngler = [r for r in [r1, r2, yuksek_3ay] if r > fiyat]
        en_yakin_destek = max(destekler) if destekler else dusuk_3ay
        en_yakin_direnc = min(direngler) if direngler else yuksek_3ay
        destege_uzaklik = ((fiyat - en_yakin_destek) / fiyat) * 100
        direce_uzaklik = ((en_yakin_direnc - fiyat) / fiyat) * 100
        risk_odul = round(direce_uzaklik / destege_uzaklik, 2) if destege_uzaklik > 0 else 0
        if destege_uzaklik <= 3:      pozisyon = '🟢 Desteğe yakın'
        elif direce_uzaklik <= 3:     pozisyon = '🔴 Dirençte'
        elif destege_uzaklik <= 8:    pozisyon = '🟡 Orta bölge'
        else:                         pozisyon = '⚪ Serbest bölge'
        return {
            'Destek': round(en_yakin_destek, 2),
            'Direnç': round(en_yakin_direnc, 2),
            'Desteğe %': round(destege_uzaklik, 1),
            'Dirençe %': round(direce_uzaklik, 1),
            'Risk/Ödül': risk_odul,
            'Pozisyon': pozisyon,
        }
    except:
        return None

def teknik_veri_cek(kod):
    try:
        ticker = yf.Ticker(f"{kod}.IS")
        df = ticker.history(period='3mo', interval='1d')
        if df is None or len(df) < 31:
            return None, None
        df = tamamlanmis_gunler(df)
        close = df['Close']
        volume = df['Volume']
        rsi_son = rsi_hesapla(close).iloc[-1]
        macd_line, signal_line, histogram = macd_hesapla(close)
        macd_son = macd_line.iloc[-1]
        signal_son = signal_line.iloc[-1]
        histogram_son = histogram.iloc[-1]
        histogram_onc = histogram.iloc[-2] if len(histogram) > 1 else 0
        macd_kesisim = histogram_onc < 0 and histogram_son > 0
        ma20 = close.rolling(20).mean().iloc[-1]
        ma50 = close.rolling(50).mean().iloc[-1] if len(close) >= 50 else None
        fiyat = close.iloc[-1]
        bb_orta = close.rolling(20).mean()
        bb_std = close.rolling(20).std()
        bb_ust = (bb_orta + 2 * bb_std).iloc[-1]
        bb_alt = (bb_orta - 2 * bb_std).iloc[-1]
        bb_yuzde = (fiyat - bb_alt) / (bb_ust - bb_alt) * 100
        vol_ort_20 = volume.rolling(20).mean().iloc[-1]
        vol_son3 = volume.iloc[-3:].mean()
        hacim_orani = vol_son3 / vol_ort_20 if vol_ort_20 > 0 else 1
        trend_5 = (fiyat - close.iloc[-5]) / close.iloc[-5] * 100 if len(close) >= 5 else 0
        trend_20 = (fiyat - close.iloc[-20]) / close.iloc[-20] * 100 if len(close) >= 20 else 0
        teknik = {
            'RSI': round(rsi_son, 1), 'MACD': round(macd_son, 4),
            'MACD_Signal': round(signal_son, 4), 'MACD_Histogram': round(histogram_son, 4),
            'MACD_Kesisim': macd_kesisim, 'MA20': round(ma20, 2),
            'MA50': round(ma50, 2) if ma50 else None, 'BB_Yuzde': round(bb_yuzde, 1),
            'Hacim_Orani': round(hacim_orani, 2), 'Trend_5G': round(trend_5, 1),
            'Trend_20G': round(trend_20, 1), 'Fiyat': round(fiyat, 2),
        }
        dd = direnc_destek_hesapla(df)
        return teknik, dd
    except:
        return None, None

def teknik_skor_hesapla(veri):
    if veri is None:
        return 0, 'Veri yok'
    skor = 0
    ac = []
    rsi = veri.get('RSI')
    macd = veri.get('MACD')
    sig = veri.get('MACD_Signal')
    kes = veri.get('MACD_Kesisim')
    fiy = veri.get('Fiyat')
    ma20 = veri.get('MA20')
    ma50 = veri.get('MA50')
    bb = veri.get('BB_Yuzde')
    hvol = veri.get('Hacim_Orani')
    if rsi is not None:
        if rsi < 30:   skor += 12; ac.append(f'RSI {rsi:.0f} aşırı satım 🟢')
        elif rsi < 45: skor += 8;  ac.append(f'RSI {rsi:.0f} alım bölgesi 🟢')
        elif rsi < 60: skor += 5;  ac.append(f'RSI {rsi:.0f} nötr 🟡')
        elif rsi < 70: skor += 0;  ac.append(f'RSI {rsi:.0f} yüksek 🟡')
        else:          skor -= 10; ac.append(f'RSI {rsi:.0f} aşırı alım 🔴')
    if macd is not None and sig is not None:
        if kes:              skor += 10; ac.append('MACD kesişim ⚡')
        elif macd>sig>0:     skor += 8;  ac.append('MACD pozitif & yukarı 🟢')
        elif macd>sig:       skor += 4;  ac.append('MACD yukarı (negatif bölge) 🟡')
        elif macd<sig and macd>0: skor -= 3; ac.append('MACD satış 🟠')
        else:                skor -= 8;  ac.append('MACD negatif & aşağı 🔴')
    if fiy and ma20:
        if fiy > ma20: skor += 8; ac.append('MA20 üstünde 🟢')
        else:          skor -= 5; ac.append('MA20 altında 🔴')
    if fiy and ma50:
        if fiy > ma50: skor += 5; ac.append('MA50 üstünde 🟢')
        else:          skor -= 3; ac.append('MA50 altında 🟠')
    if bb is not None:
        if bb < 20:   skor += 5; ac.append(f'Bollinger alt ({bb:.0f}%) 🟢')
        elif bb > 80: skor -= 5; ac.append(f'Bollinger üst ({bb:.0f}%) 🔴')
        else:         skor += 2; ac.append(f'Bollinger orta ({bb:.0f}%) 🟡')
    if hvol is not None:
        if hvol >= 2.5:   skor += 8; ac.append(f'Hacim patlaması {hvol:.1f}x 🔥')
        elif hvol >= 1.5: skor += 5; ac.append(f'Hacim artışı {hvol:.1f}x 🟢')
        elif hvol >= 1.0: skor += 2; ac.append(f'Hacim normal {hvol:.1f}x 🟡')
        else:             skor -= 3; ac.append(f'Hacim düşük {hvol:.1f}x 🔴')
    return skor, ' | '.join(ac)

def piyasa_yonu():
    try:
        df = yf.Ticker('XU100.IS').history(period='3mo', interval='1d')
        if df is None or len(df) < 31: return None, 'DİKKATLİ'
        df = tamamlanmis_gunler(df)
        close = df['Close']
        fiyat = close.iloc[-1]
        ma20 = close.rolling(20).mean().iloc[-1]
        ma50 = close.rolling(50).mean().iloc[-1] if len(close) >= 50 else None
        rsi = rsi_hesapla(close).iloc[-1]
        macd_line, signal_line, _ = macd_hesapla(close)
        ps = 0; dl = []
        if fiyat > ma20: ps += 2; dl.append('MA20 üstü')
        else:            ps -= 2; dl.append('MA20 altı ⚠️')
        if ma50 and fiyat > ma50: ps += 2; dl.append('MA50 üstü')
        elif ma50:                ps -= 2; dl.append('MA50 altı ⚠️')
        if macd_line.iloc[-1] > signal_line.iloc[-1]: ps += 1; dl.append('MACD pozitif')
        else:                                          ps -= 1; dl.append('MACD negatif')
        if rsi < 40:   ps += 1; dl.append(f'RSI {rsi:.0f} düşük')
        elif rsi > 65: ps -= 1; dl.append(f'RSI {rsi:.0f} yüksek')
        yon = '🟢 YUKARI' if ps >= 4 else '🟡 YATAY' if ps >= 1 else '🔴 AŞAĞI'
        filtre = 'GEÇTİ' if ps >= 4 else 'DİKKATLİ' if ps >= 1 else 'BEKLE'
        return {'Fiyat': round(fiyat,0), 'RSI': round(rsi,1), 'Yön': yon,
                'Haftalık': round((fiyat-close.iloc[-6])/close.iloc[-6]*100,1),
                'Durum': ' | '.join(dl)}, filtre
    except:
        return None, 'DİKKATLİ'

# ── ANA KOD ──────────────────────────────────────────────
print("="*60)
print("  📊 TEKNİK ANALİZ MODÜLÜ")
print(f"  {datetime.now().strftime('%d.%m.%Y %H:%M')}")
print("="*60)

_simdi = datetime.now()
if _simdi.weekday() < 5 and (10, 0) <= (_simdi.hour, _simdi.minute) < (18, 15):
    print("ℹ️  Seans açık: teknik göstergeler dünkü kapanışa göre hesaplanıyor.")

print("\n📈 BIST100 kontrol ediliyor...")
piyasa, filtre = piyasa_yonu()
if piyasa:
    print(f"  {piyasa['Yön']} | RSI:{piyasa['RSI']} | Haftalık:%{piyasa['Haftalık']}")
    print(f"  Durum: {piyasa['Durum']}")
    print(f"  Yeni pozisyon: {filtre}")

print(f"\n📂 Veriler yükleniyor...")
df = pd.read_excel(SISTEM_YOLU)
df = df.rename(columns={'Sembol':'Kod'}) if 'Sembol' in df.columns else df
al_tut = df[df['Toplam Skor'] >= 55].copy()
print(f"✅ {len(df)} hisse | Teknik analiz: {len(al_tut)} hisse")
print(f"⏱️  Tahmini süre: ~{len(al_tut)*2//60+2} dakika\n")

sonuclar = []
for i, (_, row) in enumerate(al_tut.iterrows(), 1):
    kod = row['Kod']
    print(f"[{i:3d}/{len(al_tut)}] {kod}...", end=' ')
    veri, dd = teknik_veri_cek(kod)
    skor, aciklama = teknik_skor_hesapla(veri)
    sonuclar.append({
        'Kod': kod,
        'Teknik Skor': skor,
        'RSI': veri['RSI'] if veri else None,
        'MACD Durumu': '📈 Pozitif' if veri and veri['MACD'] > veri['MACD_Signal'] else '📉 Negatif',
        'MACD Kesişim': '⚡ Kesişim!' if veri and veri['MACD_Kesisim'] else '-',
        'MA20 Durumu': '✅ Üstünde' if veri and veri['Fiyat'] > veri['MA20'] else '❌ Altında',
        'MA50 Durumu': '✅ Üstünde' if veri and veri['MA50'] and veri['Fiyat'] > veri['MA50'] else '❌ Altında' if veri and veri['MA50'] else '-',
        'Bollinger %': veri['BB_Yuzde'] if veri else None,
        'Hacim Oranı': veri['Hacim_Orani'] if veri else None,
        '5G Trend (%)': veri['Trend_5G'] if veri else None,
        '20G Trend (%)': veri['Trend_20G'] if veri else None,
        'Teknik Açıklama': aciklama,
        'Destek': dd['Destek'] if dd else None,
        'Direnç': dd['Direnç'] if dd else None,
        'Desteğe %': dd['Desteğe %'] if dd else None,
        'Dirençe %': dd['Dirençe %'] if dd else None,
        'Risk/Ödül': dd['Risk/Ödül'] if dd else None,
        'Pozisyon': dd['Pozisyon'] if dd else None,
    })
    emoji = '🟢' if skor >= 20 else '🟡' if skor >= 5 else '🔴'
    rsi_str = f"{veri['RSI']:.0f}" if veri else '?'
    poz_str = dd['Pozisyon'][:15] if dd else '?'
    print(f"{emoji} {skor:+d} | RSI:{rsi_str} | {poz_str}")
    time.sleep(0.3)

TEKNIK_KOLONLAR = ['Kod','Teknik Skor','RSI','MACD Durumu','MACD Kesişim','MA20 Durumu','MA50 Durumu',
                   'Bollinger %','Hacim Oranı','5G Trend (%)','20G Trend (%)','Teknik Açıklama',
                   'Destek','Direnç','Desteğe %','Dirençe %','Risk/Ödül','Pozisyon']
teknik_df = pd.DataFrame(sonuclar, columns=TEKNIK_KOLONLAR)   # AL listesi boşsa da çalışsın
if len(teknik_df) == 0:
    print("ℹ️  AL listesinde (skor ≥ 55) hisse yok; teknik analiz yapılmadı.")
df = df.merge(teknik_df, on='Kod', how='left')
df['Teknik Skor'] = df['Teknik Skor'].fillna(0).astype(int)
df['Kombine Skor'] = df['Toplam Skor'] + df['Teknik Skor']
df['Teknik Karar'] = df.apply(lambda r: '🟢 TEKNİK ONAY' if r['Kombine Skor']>=70 and r['Teknik Skor']>=15
                               else '🟡 ZAYIF ONAY' if r['Kombine Skor']>=55 and r['Teknik Skor']>=5
                               else '🔴 TEKNİK RED' if r['Teknik Skor']<0 else '⚪ NÖTR', axis=1)
df = df.sort_values('Kombine Skor', ascending=False).reset_index(drop=True)
onay = df[df['Teknik Karar']=='🟢 TEKNİK ONAY']
print(f"\n✅ Tamamlandı! Teknik Onaylı: {len(onay)} hisse")
print("\n🏆 TOP 10:")
for _, r in onay.head(10).iterrows():
    print(f"  {r['Kod']:8s} Kombine:{r['Kombine Skor']:3.0f} | {r.get('Pozisyon','?')}")

# ── EXCEL'E YAZ ──────────────────────────────────────────
print(f"\n💾 Excel güncelleniyor...")
wb = load_workbook(EXCEL_YOLU)
for s in ['📊 Teknik Analiz','✅ Teknik Onay']:
    if s in wb.sheetnames: del wb[s]

TARIH = datetime.now().strftime('%d.%m.%Y %H:%M')
piyasa_str = f"BIST100: {piyasa['Yön']} | {piyasa['Durum']}" if piyasa else "Piyasa verisi yok"
filtre_renk = '00FF88' if filtre=='GEÇTİ' else 'FFB300' if filtre=='DİKKATLİ' else 'FF4757'

def hc(ws,r,c,v,bg='051020',fg='00FF88',sz=10):
    x=ws.cell(r,c,v); x.font=Font(name='Arial',bold=True,color=fg,size=sz)
    x.fill=PatternFill('solid',fgColor=bg); x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    x.border=BORDER2; return x

def dc(ws,r,c,v,bg='091422',fg='E8F4F8',bold=False,sz=9,wrap=False,align='center'):
    try: val=None if isinstance(v,float) and np.isnan(v) else v
    except: val=v
    x=ws.cell(r,c,val); x.font=Font(name='Arial',bold=bold,color=fg,size=sz)
    x.fill=PatternFill('solid',fgColor=bg); x.alignment=Alignment(horizontal=align,vertical='center',wrap_text=wrap)
    x.border=BORDER2; return x

MACD_R = {'📈 Pozitif':('001A10','00FF88'),'📉 Negatif':('2D0000','FF4757')}
MA_R   = {'✅ Üstünde':('001A10','00FF88'),'❌ Altında':('2D0000','FF4757')}
KAR_R  = {'🟢 TEKNİK ONAY':('001A10','00FF88'),'🟡 ZAYIF ONAY':('1A1400','FFB300'),
           '🔴 TEKNİK RED':('2D0000','FF4757'),'⚪ NÖTR':('0D1B2E','AAAAAA')}
POZ_R  = {'🟢 Desteğe yakın':('001A10','00FF88'),'🔴 Dirençte':('2D0000','FF4757'),
          '🟡 Orta bölge':('1A1400','FFB300'),'⚪ Serbest bölge':('0D1B2E','AAAAAA')}

teknik_cols = ['Kod','Toplam Skor','Teknik Skor','Kombine Skor','Teknik Karar',
               'RSI','MACD Durumu','MACD Kesişim','MA20 Durumu','MA50 Durumu',
               'Bollinger %','Hacim Oranı','5G Trend (%)','20G Trend (%)',
               'Destek','Direnç','Desteğe %','Dirençe %','Risk/Ödül','Pozisyon',
               'Teknik Açıklama']
teknik_cols = [c for c in teknik_cols if c in df.columns]

for sayfa_adi, veri_df, renk, tab in [
    ('📊 Teknik Analiz', df[df['Teknik Açıklama'].notna()].sort_values('Kombine Skor',ascending=False), '00D4FF', '00D4FF'),
    ('✅ Teknik Onay', onay.sort_values('Kombine Skor',ascending=False), '00FF88', '00FF88')
]:
    ws = wb.create_sheet(sayfa_adi)
    ws.sheet_view.showGridLines = False
    ws.sheet_tab_color = tab
    ws.merge_cells(f'A1:{get_column_letter(len(teknik_cols))}1')
    ws['A1'] = f"{sayfa_adi} — {TARIH} | {len(veri_df)} hisse"
    ws['A1'].font = Font(name='Arial',bold=True,size=13,color=renk)
    ws['A1'].fill = PatternFill('solid',fgColor='020D1A')
    ws['A1'].alignment = Alignment(horizontal='center',vertical='center')
    ws.row_dimensions[1].height = 36
    ws.merge_cells(f'A2:{get_column_letter(len(teknik_cols))}2')
    ws['A2'] = piyasa_str
    ws['A2'].font = Font(name='Arial',size=9,color=filtre_renk,italic=True)
    ws['A2'].fill = PatternFill('solid',fgColor='020D1A')
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 18
    for ci,col in enumerate(teknik_cols,1): hc(ws,3,ci,col,fg=renk)
    ws.row_dimensions[3].height = 28

    for ri,(_,row) in enumerate(veri_df.iterrows(),4):
        bg = '0D1B2E' if ri%2==0 else '091422'
        for ci,cn in enumerate(teknik_cols,1):
            val = row.get(cn)
            if cn=='Kombine Skor':
                c=dc(ws,ri,ci,val,bg=bg)
                if pd.notna(val): c.font=Font(name='Arial',bold=True,size=11,color='00FF88' if val>=70 else '00D4FF' if val>=55 else 'FFB300')
            elif cn=='Teknik Skor':
                c=dc(ws,ri,ci,val,bg=bg)
                if pd.notna(val): c.font=Font(name='Arial',bold=True,size=10,color='00FF88' if val>=20 else 'FFB300' if val>=5 else 'FF4757')
            elif cn=='Teknik Karar' and str(val) in KAR_R:
                bk,fk=KAR_R[str(val)]; c=ws.cell(ri,ci,val)
                c.fill=PatternFill('solid',fgColor=bk); c.font=Font(name='Arial',bold=True,size=9,color=fk)
                c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
            elif cn=='MACD Durumu' and str(val) in MACD_R:
                bk,fk=MACD_R[str(val)]; c=ws.cell(ri,ci,val)
                c.fill=PatternFill('solid',fgColor=bk); c.font=Font(name='Arial',bold=True,size=9,color=fk)
                c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
            elif cn in ['MA20 Durumu','MA50 Durumu'] and str(val) in MA_R:
                bk,fk=MA_R[str(val)]; c=ws.cell(ri,ci,val)
                c.fill=PatternFill('solid',fgColor=bk); c.font=Font(name='Arial',bold=True,size=9,color=fk)
                c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
            elif cn=='Pozisyon' and str(val) in POZ_R:
                bk,fk=POZ_R[str(val)]; c=ws.cell(ri,ci,val)
                c.fill=PatternFill('solid',fgColor=bk); c.font=Font(name='Arial',bold=True,size=9,color=fk)
                c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
            elif cn=='RSI':
                c=dc(ws,ri,ci,val,bg=bg)
                if pd.notna(val):
                    try: c.font=Font(name='Arial',size=10,bold=True,color='00FF88' if float(val)<40 else 'FFB300' if float(val)<65 else 'FF4757')
                    except: pass
            elif cn=='Hacim Oranı':
                c=dc(ws,ri,ci,val,bg=bg)
                if pd.notna(val):
                    try: c.font=Font(name='Arial',size=9,color='00FF88' if float(val)>=2 else 'FFB300' if float(val)>=1 else 'FF4757')
                    except: pass
            elif cn=='Risk/Ödül':
                c=dc(ws,ri,ci,val,bg=bg,bold=True)
                if pd.notna(val):
                    try: c.font=Font(name='Arial',bold=True,size=10,color='00FF88' if float(val)>=2 else 'FFB300' if float(val)>=1 else 'FF4757')
                    except: pass
            elif cn in ['Desteğe %','Dirençe %']:
                c=dc(ws,ri,ci,val,bg=bg)
                if pd.notna(val):
                    try: c.font=Font(name='Arial',size=9,color='00FF88' if cn=='Desteğe %' and float(val)<=5 else 'FF4757' if cn=='Dirençe %' and float(val)<=3 else 'E8F4F8')
                    except: pass
            elif cn in ['Destek','Direnç']:
                dc(ws,ri,ci,f"{val}₺" if pd.notna(val) else '-',bg=bg,
                   fg='00FF88' if cn=='Destek' else 'FF4757')
            elif cn=='MACD Kesişim':
                c=dc(ws,ri,ci,val,bg=bg)
                if str(val)=='⚡ Kesişim!': c.font=Font(name='Arial',bold=True,size=10,color='FFD700')
            elif cn=='Teknik Açıklama':
                dc(ws,ri,ci,val,bg='051A10',fg='B2DFDB',sz=8,wrap=True,align='left')
            else:
                dc(ws,ri,ci,val,bg=bg)
        ws.row_dimensions[ri].height = 30

    gw = {'Kod':8,'Toplam Skor':8,'Teknik Skor':9,'Kombine Skor':10,'Teknik Karar':14,
          'RSI':7,'MACD Durumu':12,'MACD Kesişim':11,'MA20 Durumu':12,'MA50 Durumu':12,
          'Bollinger %':10,'Hacim Oranı':10,'5G Trend (%)':10,'20G Trend (%)':10,
          'Destek':10,'Direnç':10,'Desteğe %':10,'Dirençe %':10,'Risk/Ödül':10,
          'Pozisyon':16,'Teknik Açıklama':50}
    for ci,cn in enumerate(teknik_cols,1): ws.column_dimensions[get_column_letter(ci)].width=gw.get(cn,10)
    ws.freeze_panes='F4'
    ws.auto_filter.ref=f'A3:{get_column_letter(len(teknik_cols))}3'

wb.save(EXCEL_YOLU)
print(f"✅ Teknik Analiz ve Teknik Onay sayfaları güncellendi!")

# ── AL LİSTESİ VE HAFTALIK TAKİP GÜNCELLE ──────────────────
print("\n🔄 AL Listesi ve Haftalık Takip güncelleniyor...")
df_sistem = pd.read_excel(SISTEM_YOLU)
df_sistem = df_sistem.rename(columns={'Sembol':'Kod'}) if 'Sembol' in df_sistem.columns else df_sistem
al_listesi = df_sistem[df_sistem['Toplam Skor'] >= 55].sort_values('Toplam Skor', ascending=False).reset_index(drop=True)

wb3 = load_workbook(EXCEL_YOLU)

def yaz(ws, r, c, v, bg='091422', fg='E8F4F8', bold=False, sz=9, wrap=False, align='center'):
    try: val = None if isinstance(v, float) and np.isnan(v) else v
    except: val = v
    x = ws.cell(r, c, val)
    x.font = Font(name='Arial', bold=bold, color=fg, size=sz)
    x.fill = PatternFill('solid', fgColor=bg)
    x.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
    x.border = BORDER2
    return x

def baslik(ws, r, c, v, bg='051020', fg='00FF88', sz=10):
    x = ws.cell(r, c, v)
    x.font = Font(name='Arial', bold=True, color=fg, size=sz)
    x.fill = PatternFill('solid', fgColor=bg)
    x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    x.border = BORDER2
    return x

KARAR_RENK = {'🟢 AL':('001A10','00FF88'),'🟡 TUT':('1A1400','FFB300'),'🔴 SAT':('2D0000','FF4757')}

al_cols = ['#','Kod','Son Fiyat','Toplam Skor','KAP Alım Sinyali','Karar',
           'Aylık (%)','Haftalık (%)','Günlük (%)','Son 3 Gün (%)','Yıl İçi (%)',
           'Zirveden (%)','Dipten (%)','Hedef (Ort.)','Hedefe Uzaklık (%)','Hacim Artışı (%)','Likidite']

if '🟢 AL Listesi' in wb3.sheetnames: del wb3['🟢 AL Listesi']
ws_al = wb3.create_sheet('🟢 AL Listesi', 0)
ws_al.sheet_view.showGridLines = False
ws_al.sheet_tab_color = '00FF88'
ws_al.merge_cells(f'A1:{get_column_letter(len(al_cols))}1')
ws_al['A1'] = f"🟢 AL LİSTESİ — {TARIH} | {len(al_listesi)} hisse"
ws_al['A1'].font = Font(name='Arial', bold=True, size=13, color='00FF88')
ws_al['A1'].fill = PatternFill('solid', fgColor='020D1A')
ws_al['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws_al.row_dimensions[1].height = 36
for ci, col in enumerate(al_cols, 1): baslik(ws_al, 2, ci, col)
ws_al.row_dimensions[2].height = 28

for ri, (_, row) in enumerate(al_listesi.iterrows(), 3):
    bg = '0D1B2E' if ri%2==0 else '091422'
    karar = '🟢 AL'
    for ci, cn in enumerate(al_cols, 1):
        if cn == '#': yaz(ws_al, ri, ci, ri-2, bg=bg, bold=True)
        elif cn == 'Karar':
            bk,fk = KARAR_RENK.get(karar,('091422','E8F4F8'))
            c = ws_al.cell(ri,ci,karar); c.fill=PatternFill('solid',fgColor=bk)
            c.font=Font(name='Arial',bold=True,size=10,color=fk)
            c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
        elif cn == 'Toplam Skor':
            val=row.get('Toplam Skor'); c=yaz(ws_al,ri,ci,val,bg=bg,bold=True,sz=11)
            if pd.notna(val): c.font=Font(name='Arial',bold=True,size=11,color='00FF88' if val>=80 else '00D4FF' if val>=65 else 'FFB300')
        elif cn == 'KAP Alım Sinyali':
            val=row.get('KAP Alım Sinyali','-'); fg='00FF88' if '🟢' in str(val) else 'FFB300' if '🟡' in str(val) else 'AAAAAA'
            yaz(ws_al,ri,ci,val,bg=bg,fg=fg,bold='🟢' in str(val))
        elif cn in ['Aylık (%)','Haftalık (%)','Günlük (%)','Son 3 Gün (%)','Yıl İçi (%)','Zirveden (%)','Dipten (%)','Hedefe Uzaklık (%)','Hacim Artışı (%)']:
            val=row.get(cn); c=yaz(ws_al,ri,ci,val,bg=bg)
            if pd.notna(val):
                try: c.font=Font(name='Arial',size=9,color='00FF88' if float(val)>0 else 'FF4757' if float(val)<0 else 'AAAAAA')
                except: pass
        elif cn == 'Son Fiyat':
            val=row.get('Son Fiyat'); yaz(ws_al,ri,ci,f"{val}₺" if pd.notna(val) else '-',bg=bg,bold=True,fg='FFD700')
        else: yaz(ws_al,ri,ci,row.get(cn),bg=bg)
    ws_al.row_dimensions[ri].height = 25

gw_al = {'#':4,'Kod':8,'Son Fiyat':10,'Toplam Skor':9,'KAP Alım Sinyali':12,'Karar':8,
         'Aylık (%)':9,'Haftalık (%)':9,'Günlük (%)':9,'Son 3 Gün (%)':9,'Yıl İçi (%)':9,
         'Zirveden (%)':10,'Dipten (%)':9,'Hedef (Ort.)':10,'Hedefe Uzaklık (%)':10,
         'Hacim Artışı (%)':11,'Likidite':10}
for ci,cn in enumerate(al_cols,1): ws_al.column_dimensions[get_column_letter(ci)].width=gw_al.get(cn,10)
ws_al.freeze_panes='C3'; ws_al.auto_filter.ref=f'A2:{get_column_letter(len(al_cols))}2'

ht_cols = ['#','Kod','Son Fiyat','Toplam Skor','KAP Alım Sinyali','Karar',
           'Son 3 Gün (%)','Haftalık (%)','Aylık (%)','Yıl İçi (%)',
           'Zirveden (%)','Hedefe Uzaklık (%)','Likidite']

if '📅 Haftalık Takip' in wb3.sheetnames: del wb3['📅 Haftalık Takip']
ws_ht = wb3.create_sheet('📅 Haftalık Takip')
ws_ht.sheet_view.showGridLines = False
ws_ht.sheet_tab_color = '00D4FF'
ws_ht.merge_cells(f'A1:{get_column_letter(len(ht_cols))}1')
ws_ht['A1'] = f"📅 HAFTALIK TAKİP — {TARIH} | {len(al_listesi)} hisse"
ws_ht['A1'].font = Font(name='Arial',bold=True,size=13,color='00D4FF')
ws_ht['A1'].fill = PatternFill('solid',fgColor='020D1A')
ws_ht['A1'].alignment = Alignment(horizontal='center',vertical='center')
ws_ht.row_dimensions[1].height = 36
for ci,col in enumerate(ht_cols,1): baslik(ws_ht,2,ci,col,fg='00D4FF')
ws_ht.row_dimensions[2].height = 28

for ri,(_, row) in enumerate(al_listesi.iterrows(),3):
    bg='0D1B2E' if ri%2==0 else '091422'
    karar='🟢 AL'
    for ci,cn in enumerate(ht_cols,1):
        if cn=='#': yaz(ws_ht,ri,ci,ri-2,bg=bg,bold=True)
        elif cn=='Karar':
            bk,fk=KARAR_RENK.get(karar,('091422','E8F4F8'))
            c=ws_ht.cell(ri,ci,karar); c.fill=PatternFill('solid',fgColor=bk)
            c.font=Font(name='Arial',bold=True,size=10,color=fk)
            c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
        elif cn=='Toplam Skor':
            val=row.get('Toplam Skor'); c=yaz(ws_ht,ri,ci,val,bg=bg,bold=True,sz=11)
            if pd.notna(val): c.font=Font(name='Arial',bold=True,size=11,color='00FF88' if val>=80 else '00D4FF' if val>=65 else 'FFB300')
        elif cn in ['Son 3 Gün (%)','Haftalık (%)','Aylık (%)','Yıl İçi (%)','Zirveden (%)','Hedefe Uzaklık (%)']:
            val=row.get(cn); c=yaz(ws_ht,ri,ci,val,bg=bg)
            if pd.notna(val):
                try: c.font=Font(name='Arial',size=9,color='00FF88' if float(val)>0 else 'FF4757' if float(val)<0 else 'AAAAAA')
                except: pass
        elif cn=='Son Fiyat':
            val=row.get('Son Fiyat'); yaz(ws_ht,ri,ci,f"{val}₺" if pd.notna(val) else '-',bg=bg,bold=True,fg='FFD700')
        elif cn=='KAP Alım Sinyali':
            val=row.get('KAP Alım Sinyali','-'); fg='00FF88' if '🟢' in str(val) else 'FFB300' if '🟡' in str(val) else 'AAAAAA'
            yaz(ws_ht,ri,ci,val,bg=bg,fg=fg,bold='🟢' in str(val))
        else: yaz(ws_ht,ri,ci,row.get(cn),bg=bg)
    ws_ht.row_dimensions[ri].height = 25

gw_ht={'#':4,'Kod':8,'Son Fiyat':10,'Toplam Skor':9,'KAP Alım Sinyali':12,'Karar':8,
       'Son 3 Gün (%)':9,'Haftalık (%)':9,'Aylık (%)':9,'Yıl İçi (%)':9,
       'Zirveden (%)':10,'Hedefe Uzaklık (%)':10,'Likidite':10}
for ci,cn in enumerate(ht_cols,1): ws_ht.column_dimensions[get_column_letter(ci)].width=gw_ht.get(cn,10)
ws_ht.freeze_panes='C3'; ws_ht.auto_filter.ref=f'A2:{get_column_letter(len(ht_cols))}2'

wb3.save(EXCEL_YOLU)
print(f"✅ AL Listesi ve Haftalık Takip güncellendi!")

# ── STATİK SAYFALARI GÜNCELLE ────────────────────────────
print("\n🔄 Diğer sayfalar güncelleniyor...")
df_tam = pd.read_excel(SISTEM_YOLU)
df_tam = df_tam.rename(columns={'Sembol':'Kod'}) if 'Sembol' in df_tam.columns else df_tam
df_al_tut2 = df_tam[df_tam['Toplam Skor'] >= 35].sort_values('Toplam Skor', ascending=False).reset_index(drop=True)

wb5 = load_workbook(EXCEL_YOLU)

def s_yaz(ws,r,c,v,bg='091422',fg='E8F4F8',bold=False,sz=9,wrap=False,align='center'):
    try: val=None if isinstance(v,float) and np.isnan(v) else v
    except: val=v
    x=ws.cell(r,c,val); x.font=Font(name='Arial',bold=bold,color=fg,size=sz)
    x.fill=PatternFill('solid',fgColor=bg); x.alignment=Alignment(horizontal=align,vertical='center',wrap_text=wrap)
    x.border=BORDER2; return x

def s_baslik(ws,r,c,v,bg='051020',fg='00FF88',sz=10):
    x=ws.cell(r,c,v); x.font=Font(name='Arial',bold=True,color=fg,size=sz)
    x.fill=PatternFill('solid',fgColor=bg); x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    x.border=BORDER2; return x

def sayfa_olustur(wb,ad,renk,tab,baslik_metni,kolonlar):
    if ad in wb.sheetnames: del wb[ad]
    ws=wb.create_sheet(ad); ws.sheet_view.showGridLines=False; ws.sheet_tab_color=tab
    ws.merge_cells(f'A1:{get_column_letter(len(kolonlar))}1')
    ws['A1']=baslik_metni; ws['A1'].font=Font(name='Arial',bold=True,size=13,color=renk)
    ws['A1'].fill=PatternFill('solid',fgColor='020D1A'); ws['A1'].alignment=Alignment(horizontal='center',vertical='center')
    ws.row_dimensions[1].height=36
    for ci,col in enumerate(kolonlar,1): s_baslik(ws,2,ci,col,fg=renk)
    ws.row_dimensions[2].height=28
    return ws

# Kalite Analizi
kalite_cols=['#','Kod','Toplam Skor','KAP Alım Sinyali','PEG Skor','PEG Oranı','ROE Tutarlılığı','ROE Ort. (%)','Büyüme Tutarlılığı','Kar Büyümesi (%)','Reel Kar Büyümesi (%)','Ciro Büyümesi (%)','Reel Ciro Büyümesi (%)','Güncel F/K','Tarihi F/K Ort.','Hedef (Ort.)','Hedefe Uzaklık (%)','Likidite']
kalite_cols=[c for c in kalite_cols if c in df_al_tut2.columns or c=='#']
ws_kal=sayfa_olustur(wb5,'🎯 Kalite Analizi','FFD700','FFD700',f"🎯 KALİTE ANALİZİ — {TARIH} | {len(df_al_tut2)} hisse",kalite_cols)
for ri,(_,row) in enumerate(df_al_tut2.iterrows(),3):
    bg='0D1B2E' if ri%2==0 else '091422'
    for ci,cn in enumerate(kalite_cols,1):
        if cn=='#': s_yaz(ws_kal,ri,ci,ri-2,bg=bg,bold=True)
        elif cn=='Toplam Skor':
            c=s_yaz(ws_kal,ri,ci,row.get(cn),bg=bg,bold=True,sz=11)
            try: c.font=Font(name='Arial',bold=True,size=11,color='00FF88' if row.get(cn,0)>=80 else '00D4FF' if row.get(cn,0)>=55 else 'FFB300')
            except: pass
        elif cn in ['Kar Büyümesi (%)','Reel Kar Büyümesi (%)','Ciro Büyümesi (%)','Reel Ciro Büyümesi (%)','Hedefe Uzaklık (%)']:
            val=row.get(cn); c=s_yaz(ws_kal,ri,ci,val,bg=bg)
            if pd.notna(val):
                try: c.font=Font(name='Arial',size=9,color='00FF88' if float(val)>0 else 'FF4757')
                except: pass
        else: s_yaz(ws_kal,ri,ci,row.get(cn),bg=bg)
    ws_kal.row_dimensions[ri].height=25
for ci in range(1,len(kalite_cols)+1): ws_kal.column_dimensions[get_column_letter(ci)].width=12
ws_kal.column_dimensions['A'].width=4; ws_kal.column_dimensions['B'].width=8
ws_kal.freeze_panes='C3'; ws_kal.auto_filter.ref=f'A2:{get_column_letter(len(kalite_cols))}2'

# Finansal Sağlık
fin_cols=['#','Kod','Toplam Skor','FCF Kategori','FCF Verimi (%)','FCF (mn TL)','Faiz Karşılama Kat.','Faiz Karşılama','Borç Riski','Net Borç/Özsermaye','Borç/Özsermaye','Likidite','Ort. TL Hacim (mn)']
fin_cols=[c for c in fin_cols if c in df_al_tut2.columns or c=='#']
ws_fin=sayfa_olustur(wb5,'💵 Finansal Sağlık','00D4FF','00D4FF',f"💵 FİNANSAL SAĞLIK — {TARIH} | {len(df_al_tut2)} hisse",fin_cols)
for ri,(_,row) in enumerate(df_al_tut2.iterrows(),3):
    bg='0D1B2E' if ri%2==0 else '091422'
    for ci,cn in enumerate(fin_cols,1):
        if cn=='#': s_yaz(ws_fin,ri,ci,ri-2,bg=bg,bold=True)
        elif cn in ['FCF Kategori','Faiz Karşılama Kat.','Borç Riski']:
            val=str(row.get(cn,'')); fg='00FF88' if '🟢' in val else 'FFB300' if '🟡' in val else 'FF4757' if '🔴' in val else 'AAAAAA'
            s_yaz(ws_fin,ri,ci,row.get(cn),bg=bg,fg=fg)
        else: s_yaz(ws_fin,ri,ci,row.get(cn),bg=bg)
    ws_fin.row_dimensions[ri].height=25
for ci in range(1,len(fin_cols)+1): ws_fin.column_dimensions[get_column_letter(ci)].width=12
ws_fin.column_dimensions['A'].width=4; ws_fin.column_dimensions['B'].width=8
ws_fin.freeze_panes='C3'; ws_fin.auto_filter.ref=f'A2:{get_column_letter(len(fin_cols))}2'

# Borç & Likidite
borc_cols=['#','Kod','Toplam Skor','Borç Riski','Net Borç/Özsermaye','Borç/Özsermaye','Likidite','Ort. TL Hacim (mn)','FCF Kategori','Faiz Karşılama']
borc_cols=[c for c in borc_cols if c in df_tam.columns or c=='#']
df_borc=df_tam.sort_values('Net Borç/Özsermaye',ascending=True).reset_index(drop=True)
ws_bor=sayfa_olustur(wb5,'💰 Borç & Likidite','FF4757','FF4757',f"💰 BORÇ & LİKİDİTE — {TARIH} | {len(df_borc)} hisse",borc_cols)
for ri,(_,row) in enumerate(df_borc.iterrows(),3):
    bg='0D1B2E' if ri%2==0 else '091422'
    for ci,cn in enumerate(borc_cols,1):
        if cn=='#': s_yaz(ws_bor,ri,ci,ri-2,bg=bg,bold=True)
        elif cn=='Borç Riski':
            val=str(row.get(cn,'')); fg='00FF88' if '🟢' in val else 'FFB300' if '🟡' in val else 'FF4757' if '🔴' in val else 'AAAAAA'
            s_yaz(ws_bor,ri,ci,row.get(cn),bg=bg,fg=fg)
        elif cn=='Net Borç/Özsermaye':
            val=row.get(cn); c=s_yaz(ws_bor,ri,ci,val,bg=bg,bold=True)
            if pd.notna(val):
                try: c.font=Font(name='Arial',bold=True,size=9,color='00FF88' if float(val)<0 else 'FFB300' if float(val)<1 else 'FF4757')
                except: pass
        else: s_yaz(ws_bor,ri,ci,row.get(cn),bg=bg)
    ws_bor.row_dimensions[ri].height=22
for ci in range(1,len(borc_cols)+1): ws_bor.column_dimensions[get_column_letter(ci)].width=14
ws_bor.column_dimensions['A'].width=4; ws_bor.column_dimensions['B'].width=8
ws_bor.freeze_panes='C3'; ws_bor.auto_filter.ref=f'A2:{get_column_letter(len(borc_cols))}2'

# Sektör Analizi
df_al_only=df_tam[df_tam['Toplam Skor']>=55].copy()
if 'İSY Sektör' in df_al_only.columns:
    df_al_only['Sektör']=df_al_only['İSY Sektör'].fillna('').replace('', 'Bilinmiyor')
if 'Sektör' in df_al_only.columns and len(df_al_only)>0:
    sektor_ozet=df_al_only.groupby('Sektör').agg(AL_Hisse=('Kod','count'),Ort_Skor=('Toplam Skor','mean'),Ort_FK=('Güncel F/K','mean'),Ort_PDDD=('PD/DD','mean'),Ort_Aylik=('Aylık (%)','mean')).reset_index().sort_values('AL_Hisse',ascending=False)
    sektor_ozet.columns=['Sektör','AL Hisse','Ort Skor','Ort F/K','Ort PD/DD','Ort Aylık %']
    if '🏭 Sektör Analizi' in wb5.sheetnames: del wb5['🏭 Sektör Analizi']
    ws_sek=wb5.create_sheet('🏭 Sektör Analizi'); ws_sek.sheet_view.showGridLines=False; ws_sek.sheet_tab_color='AA44FF'
    s_cols=['Sektör','AL Hisse','Ort Skor','Ort F/K','Ort PD/DD','Ort Aylık %']
    ws_sek.merge_cells(f'A1:{get_column_letter(len(s_cols))}1')
    ws_sek['A1']=f"🏭 SEKTÖR ANALİZİ — {TARIH}"; ws_sek['A1'].font=Font(name='Arial',bold=True,size=13,color='AA44FF')
    ws_sek['A1'].fill=PatternFill('solid',fgColor='020D1A'); ws_sek['A1'].alignment=Alignment(horizontal='center',vertical='center')
    ws_sek.row_dimensions[1].height=36
    for ci,col in enumerate(s_cols,1): s_baslik(ws_sek,2,ci,col,fg='AA44FF')
    ws_sek.row_dimensions[2].height=28
    for ri,(_,row) in enumerate(sektor_ozet.iterrows(),3):
        bg='0D1B2E' if ri%2==0 else '091422'
        for ci,cn in enumerate(s_cols,1):
            val=row.get(cn)
            if cn=='Sektör': s_yaz(ws_sek,ri,ci,val,bg=bg,bold=True,align='left')
            elif cn=='AL Hisse':
                c=s_yaz(ws_sek,ri,ci,val,bg=bg,bold=True,sz=11)
                try: c.font=Font(name='Arial',bold=True,size=11,color='00FF88' if val>=5 else 'FFB300' if val>=2 else 'AAAAAA')
                except: pass
            elif cn=='Ort Aylık %':
                c=s_yaz(ws_sek,ri,ci,round(val,1) if pd.notna(val) else None,bg=bg)
                if pd.notna(val):
                    try: c.font=Font(name='Arial',size=9,color='00FF88' if val>0 else 'FF4757')
                    except: pass
            else: s_yaz(ws_sek,ri,ci,round(val,1) if pd.notna(val) else None,bg=bg)
        ws_sek.row_dimensions[ri].height=28
    ws_sek.column_dimensions['A'].width=25
    for ci in range(2,len(s_cols)+1): ws_sek.column_dimensions[get_column_letter(ci)].width=12

# Elenenler
df_elen=df_tam[df_tam['Toplam Skor']<18].sort_values('Toplam Skor',ascending=True).reset_index(drop=True)
elen_cols=['#','Kod','Toplam Skor','Borç Riski','Likidite','Kar Büyümesi (%)','Aylık (%)']
elen_cols=[c for c in elen_cols if c in df_elen.columns or c=='#']
ws_el=sayfa_olustur(wb5,'⚠️ Elenenler','FF4757','FF4757',f"⚠️ ELENENLER — {TARIH} | {len(df_elen)} hisse",elen_cols)
for ri,(_,row) in enumerate(df_elen.iterrows(),3):
    bg='2D0000' if ri%2==0 else '1A0000'
    for ci,cn in enumerate(elen_cols,1):
        if cn=='#': s_yaz(ws_el,ri,ci,ri-2,bg=bg,bold=True)
        elif cn=='Toplam Skor': s_yaz(ws_el,ri,ci,row.get(cn),bg=bg,fg='FF4757',bold=True)
        else: s_yaz(ws_el,ri,ci,row.get(cn),bg=bg)
    ws_el.row_dimensions[ri].height=22
for ci in range(1,len(elen_cols)+1): ws_el.column_dimensions[get_column_letter(ci)].width=14
ws_el.column_dimensions['A'].width=4; ws_el.column_dimensions['B'].width=8
ws_el.freeze_panes='C3'

wb5.save(EXCEL_YOLU)
print(f"✅ Kalite, Finansal, Borç, Sektör, Elenenler güncellendi!")

# ── KAP İÇERİDEN ALIM SAYFASI ────────────────────────────
print("\n📌 KAP İçeriden Alım güncelleniyor...")
# Hücre 2'nin okuduğu KAP verisi (tek kaynak)
kap_verisi = {}
if 'KAP Alım Sayısı' in df_tam.columns:
    for _, r in df_tam[df_tam['KAP Alım Sayısı'].notna()].iterrows():
        try: kap_verisi[str(r['Kod'])] = int(r['KAP Alım Sayısı'])
        except: pass
print(f"  ✅ {len(kap_verisi)} hissede KAP bildirimi (bilgi amaçlı, puana katılmıyor)")

kap_liste = sorted(kap_verisi.items(), key=lambda x: -x[1])
wb6 = load_workbook(EXCEL_YOLU)
if '📌 KAP İçeriden Alım' in wb6.sheetnames: del wb6['📌 KAP İçeriden Alım']
ws_kap = wb6.create_sheet('📌 KAP İçeriden Alım', 1)
ws_kap.sheet_view.showGridLines = False; ws_kap.sheet_tab_color = 'FF6B35'
kap_cols = ['#','Kod','Alım Adedi','KAP Sinyali','Toplam Skor','Aylık (%)','Son Fiyat']
ws_kap.merge_cells(f'A1:{get_column_letter(len(kap_cols))}1')
ws_kap['A1'] = f"📌 KAP İÇERİDEN ALIM — {TARIH} | {len(kap_liste)} hisse | doğrulanmadı, puana katılmıyor"
ws_kap['A1'].font = Font(name='Arial',bold=True,size=13,color='FF6B35')
ws_kap['A1'].fill = PatternFill('solid',fgColor='020D1A')
ws_kap['A1'].alignment = Alignment(horizontal='center',vertical='center')
ws_kap.row_dimensions[1].height = 36
for ci,col in enumerate(kap_cols,1):
    c=ws_kap.cell(2,ci,col); c.font=Font(name='Arial',bold=True,color='FF6B35',size=10)
    c.fill=PatternFill('solid',fgColor='051020'); c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
ws_kap.row_dimensions[2].height = 28
for ri,(kod,adet) in enumerate(kap_liste,3):
    bg='0D1B2E' if ri%2==0 else '091422'
    sinyal = f'🔥 {adet} Alım' if adet>=10 else f'✅ {adet} Alım' if adet>=3 else f'🟡 {adet} Alım'
    hisse_row = df_tam[df_tam['Kod']==kod]
    skor = hisse_row['Toplam Skor'].iloc[0] if len(hisse_row)>0 else '-'
    aylik = hisse_row['Aylık (%)'].iloc[0] if len(hisse_row)>0 else '-'
    fiyat = hisse_row['Son Fiyat'].iloc[0] if len(hisse_row)>0 else '-'
    for ci,val in enumerate([ri-2,kod,adet,sinyal,skor,aylik,f"{fiyat}₺" if fiyat!='-' else '-'],1):
        c=ws_kap.cell(ri,ci,val); c.fill=PatternFill('solid',fgColor=bg); c.border=BORDER2
        c.alignment=Alignment(horizontal='center',vertical='center')
        if ci==2: c.font=Font(name='Arial',bold=True,size=10,color='E8F4F8')
        elif ci==3: c.font=Font(name='Arial',bold=True,size=11,color='FF4500' if adet>=10 else '00FF88' if adet>=3 else 'FFB300')
        elif ci==4: c.font=Font(name='Arial',bold=True,size=10,color='FF4500' if adet>=10 else '00FF88' if adet>=3 else 'FFB300')
        else: c.font=Font(name='Arial',size=9,color='E8F4F8')
    ws_kap.row_dimensions[ri].height=28
for ci,g in enumerate([4,8,10,12,10,10,10],1): ws_kap.column_dimensions[get_column_letter(ci)].width=g
ws_kap.freeze_panes='A3'
wb6.save(EXCEL_YOLU)
print(f"✅ KAP İçeriden Alım güncellendi!")

# ── TÜM HİSSELER ─────────────────────────────────────────
print("\n📊 Tüm Hisseler güncelleniyor...")
df_tum_yeni=df_tam.sort_values('Toplam Skor',ascending=False).reset_index(drop=True)
al_say=len(df_tum_yeni[df_tum_yeni['Toplam Skor']>=55])
tut_say=len(df_tum_yeni[(df_tum_yeni['Toplam Skor']>=35)&(df_tum_yeni['Toplam Skor']<55)])
izle_say=len(df_tum_yeni[(df_tum_yeni['Toplam Skor']>=18)&(df_tum_yeni['Toplam Skor']<35)])
sat_say=len(df_tum_yeni[df_tum_yeni['Toplam Skor']<18])
wb7=load_workbook(EXCEL_YOLU)
if '📊 Tüm Hisseler' in wb7.sheetnames: del wb7['📊 Tüm Hisseler']
ws_tum=wb7.create_sheet('📊 Tüm Hisseler',2)
ws_tum.sheet_view.showGridLines=False; ws_tum.sheet_tab_color='00D4FF'
tum_cols=['#','Kod','Toplam Skor','KAP Alım Sinyali','Karar','Son Fiyat','Hedef (Ort.)','Hedefe Uzaklık (%)','Güncel F/K','Tarihi F/K Ort.','Aylık (%)','Haftalık (%)','Zirveden (%)','Likidite','Kar Büyümesi (%)','Reel Kar Büyümesi (%)','PEG Skor','ROE Tutarlılığı','Borç Riski','İSY Sektör']
tum_cols=[c for c in tum_cols if c in df_tum_yeni.columns or c in ['#','Karar']]
ws_tum.merge_cells(f'A1:{get_column_letter(len(tum_cols))}1')
ws_tum['A1']=f"📊 TÜM HİSSELER — {TARIH} | {len(df_tum_yeni)} Hisse | AL:{al_say} TUT:{tut_say} İZLE:{izle_say} SAT:{sat_say}"
ws_tum['A1'].font=Font(name='Arial',bold=True,size=13,color='00D4FF')
ws_tum['A1'].fill=PatternFill('solid',fgColor='020D1A'); ws_tum['A1'].alignment=Alignment(horizontal='center',vertical='center')
ws_tum.row_dimensions[1].height=36
for ci,col in enumerate(tum_cols,1):
    c=ws_tum.cell(2,ci,col); c.font=Font(name='Arial',bold=True,color='00D4FF',size=10)
    c.fill=PatternFill('solid',fgColor='051020'); c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
ws_tum.row_dimensions[2].height=28
KARAR_RENK2={'🟢 AL':('001A10','00FF88'),'🟡 TUT':('1A1400','FFB300'),'🟠 İZLE':('1A0A00','FF8C00'),'🔴 SAT':('2D0000','FF4757')}
for ri,(_,row) in enumerate(df_tum_yeni.iterrows(),3):
    bg='0D1B2E' if ri%2==0 else '091422'
    skor=row.get('Toplam Skor',0)
    karar='🟢 AL' if skor>=55 else '🟡 TUT' if skor>=35 else '🟠 İZLE' if skor>=18 else '🔴 SAT'
    for ci,cn in enumerate(tum_cols,1):
        val=row.get(cn) if cn not in ['#','Karar'] else None
        if cn=='#':
            c=ws_tum.cell(ri,ci,ri-2); c.fill=PatternFill('solid',fgColor=bg)
            c.font=Font(name='Arial',size=8,color='AAAAAA'); c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
        elif cn=='Karar':
            bk,fk=KARAR_RENK2.get(karar,('091422','E8F4F8')); c=ws_tum.cell(ri,ci,karar)
            c.fill=PatternFill('solid',fgColor=bk); c.font=Font(name='Arial',bold=True,size=9,color=fk)
            c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
        elif cn=='Toplam Skor':
            c=ws_tum.cell(ri,ci,val); c.fill=PatternFill('solid',fgColor=bg)
            c.font=Font(name='Arial',bold=True,size=11,color='00FF88' if skor>=55 else 'FFB300' if skor>=35 else 'FF8C00' if skor>=18 else 'FF4757')
            c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
        elif cn in ['Aylık (%)','Haftalık (%)','Zirveden (%)','Hedefe Uzaklık (%)','Kar Büyümesi (%)','Reel Kar Büyümesi (%)']:
            c=ws_tum.cell(ri,ci,val); c.fill=PatternFill('solid',fgColor=bg)
            try: c.font=Font(name='Arial',size=9,color='00FF88' if float(val)>0 else 'FF4757' if float(val)<0 else 'AAAAAA')
            except: c.font=Font(name='Arial',size=9,color='AAAAAA')
            c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
        elif cn=='Son Fiyat':
            c=ws_tum.cell(ri,ci,f"{val}₺" if pd.notna(val) else '-'); c.fill=PatternFill('solid',fgColor=bg)
            c.font=Font(name='Arial',bold=True,size=9,color='FFD700'); c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
        else:
            c=ws_tum.cell(ri,ci,val); c.fill=PatternFill('solid',fgColor=bg)
            c.font=Font(name='Arial',size=9,color='E8F4F8'); c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
    ws_tum.row_dimensions[ri].height=22
for ci,cn in enumerate(tum_cols,1):
    ws_tum.column_dimensions[get_column_letter(ci)].width={'#':4,'Kod':8,'Toplam Skor':9,'KAP Alım Sinyali':12,'Karar':8,'Son Fiyat':10,'Hedef (Ort.)':10,'Hedefe Uzaklık (%)':10,'Güncel F/K':8,'Aylık (%)':9,'Haftalık (%)':9,'Zirveden (%)':10,'Likidite':10,'Kar Büyümesi (%)':11,'PEG Skor':10,'ROE Tutarlılığı':14,'Borç Riski':12}.get(cn,10)
ws_tum.freeze_panes='C3'; ws_tum.auto_filter.ref=f'A2:{get_column_letter(len(tum_cols))}2'
wb7.save(EXCEL_YOLU)
print(f"✅ Tüm Hisseler güncellendi! {len(df_tum_yeni)} hisse")

# ── BİLANÇO TAKVİMİ ──────────────────────────────────────
print("\n📅 Bilanço takvimi güncelleniyor...")
from datetime import date
import re
bugun = datetime.now().date()
ay = bugun.month
if ay <= 3:   q_tarih=date(bugun.year,3,31); q_ad='Q4 (Ekim-Aralık)'
elif ay <= 5: q_tarih=date(bugun.year,5,31); q_ad='Q1 (Ocak-Mart)'
elif ay <= 8: q_tarih=date(bugun.year,8,31); q_ad='Q2 (Nisan-Haziran)'
else:         q_tarih=date(bugun.year,11,30); q_ad='Q3 (Temmuz-Eylül)'

try:
    df_haberler_b=pd.read_excel(EXCEL_YOLU,sheet_name='📰 Haberler',header=1)
    df_haberler_b.columns=['Hisse','Kaynak','Tarih','Başlık','Link']
except: df_haberler_b=pd.DataFrame(columns=['Hisse','Kaynak','Tarih','Başlık','Link'])

BILANCO_KEL=['finansal sonuç','bilanço','mali tablo','faaliyet raporu','kar açıkla','dönem sonuç','ara dönem','yıllık rapor','konsolide']
kodlar_bilanco=df[df['Teknik Karar']=='🟢 TEKNİK ONAY']['Kod'].tolist()
bilanco_sonuclar=[]
for kod in kodlar_bilanco:
    bilanco_tarihi=None; bilanco_kaynak='Çeyrek Tahmini'; bilanco_notu=q_ad
    if len(df_haberler_b)>0:
        for _,haber in df_haberler_b[df_haberler_b['Hisse']==kod].iterrows():
            if any(k in str(haber['Başlık']).lower() for k in BILANCO_KEL):
                try:
                    ht=pd.to_datetime(str(haber['Tarih'])[:10],dayfirst=True).date()
                    if ht>=bugun: bilanco_tarihi=ht; bilanco_kaynak='KAP Haberi'; bilanco_notu=str(haber['Başlık'])[:60]; break
                except: continue
    if bilanco_tarihi is None: bilanco_tarihi=q_tarih
    kalan_gun=(bilanco_tarihi-bugun).days
    if kalan_gun<=7: durum='🔴 Bu hafta!'
    elif kalan_gun<=14: durum='🟠 2 hafta içinde'
    elif kalan_gun<=30: durum='🟡 Bu ay içinde'
    elif kalan_gun<=60: durum='🟢 1-2 ay sonra'
    else: durum='⚪ Uzak'
    bilanco_sonuclar.append({'Kod':kod,'Bilanço Tarihi':bilanco_tarihi.strftime('%d.%m.%Y'),'Kalan Gün':kalan_gun,'Durum':durum,'Kaynak':bilanco_kaynak,'Not':bilanco_notu})

df_bilanco=pd.DataFrame(bilanco_sonuclar,columns=['Kod','Bilanço Tarihi','Kalan Gün','Durum','Kaynak','Not']).sort_values('Kalan Gün').reset_index(drop=True)
wb8=load_workbook(EXCEL_YOLU)
if '📅 Bilanço Takvimi' in wb8.sheetnames: del wb8['📅 Bilanço Takvimi']
ws_bil=wb8.create_sheet('📅 Bilanço Takvimi'); ws_bil.sheet_view.showGridLines=False; ws_bil.sheet_tab_color='FFB300'
b_cols=['Kod','Bilanço Tarihi','Kalan Gün','Durum','Kaynak','Not']
ws_bil.merge_cells(f'A1:{get_column_letter(len(b_cols))}1')
ws_bil['A1']=f"📅 BİLANÇO TAKVİMİ — {TARIH} | {len(df_bilanco)} hisse | {q_ad}"
ws_bil['A1'].font=Font(name='Arial',bold=True,size=13,color='FFB300')
ws_bil['A1'].fill=PatternFill('solid',fgColor='020D1A'); ws_bil['A1'].alignment=Alignment(horizontal='center',vertical='center')
ws_bil.row_dimensions[1].height=36
for ci,col in enumerate(b_cols,1):
    c=ws_bil.cell(2,ci,col); c.font=Font(name='Arial',bold=True,color='FFB300',size=10)
    c.fill=PatternFill('solid',fgColor='051020'); c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
ws_bil.row_dimensions[2].height=28
DURUM_R={'🔴 Bu hafta!':('2D0000','FF4757'),'🟠 2 hafta içinde':('1A0A00','FF6B35'),'🟡 Bu ay içinde':('1A1400','FFB300'),'🟢 1-2 ay sonra':('001A10','00FF88'),'⚪ Uzak':('0D1B2E','AAAAAA')}
for ri,(_,row) in enumerate(df_bilanco.iterrows(),3):
    bg='0D1B2E' if ri%2==0 else '091422'
    for ci,cn in enumerate(b_cols,1):
        val=row.get(cn)
        if cn=='Durum' and str(val) in DURUM_R:
            bk,fk=DURUM_R[str(val)]; c=ws_bil.cell(ri,ci,val)
            c.fill=PatternFill('solid',fgColor=bk); c.font=Font(name='Arial',bold=True,size=10,color=fk)
            c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
        elif cn=='Kalan Gün':
            c=ws_bil.cell(ri,ci,val); c.fill=PatternFill('solid',fgColor=bg)
            c.font=Font(name='Arial',bold=True,size=11,color='FF4757' if val<=7 else 'FFB300' if val<=30 else '00FF88')
            c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
        elif cn=='Not':
            c=ws_bil.cell(ri,ci,val); c.fill=PatternFill('solid',fgColor='051A10')
            c.font=Font(name='Arial',size=8,color='B2DFDB'); c.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True); c.border=BORDER2
        else:
            c=ws_bil.cell(ri,ci,val); c.fill=PatternFill('solid',fgColor=bg)
            c.font=Font(name='Arial',size=9,color='E8F4F8'); c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER2
    ws_bil.row_dimensions[ri].height=35
for ci,cn in enumerate(b_cols,1): ws_bil.column_dimensions[get_column_letter(ci)].width={'Kod':8,'Bilanço Tarihi':14,'Kalan Gün':10,'Durum':18,'Kaynak':12,'Not':55}.get(cn,10)
ws_bil.freeze_panes='A3'; ws_bil.auto_filter.ref=f'A2:{get_column_letter(len(b_cols))}2'
wb8.save(EXCEL_YOLU)
print(f"✅ Bilanço Takvimi güncellendi!")
print(f"\n🎉 TÜM SAYFALAR GÜNCELLENDİ!")