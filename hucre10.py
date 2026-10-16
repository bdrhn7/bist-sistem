# ═══════════════════════════════════════════════════════════════
#  HÜCRE 10 — ZİRVE STRATEJİSİ  (v1.3 — 16.10.2026)
#  Backtestten en sağlam geçen kural (hücre 9 v2 + bağımsız istatistik incelemesi):
#    Her ayın ilk işlem günü, günlük işlem hacmi ≥ 50 mn TL olan hisselerden
#    52 haftalık zirvesine EN YAKIN 30 hisse — eşit ağırlık — ay sonuna kadar tut.
#  Çalıştırma: hücre 1 → hücre 2 → bu hücre. Ayda bir yeter (ayın ilk işlem günü).
#  Yapay zekâ kullanmaz. Hücre 3-8'e ihtiyaç duymaz.
#
#  Ek güvenlik kuralları (bağımsız incelemeden):
#   - 5 TL altı hisseler ve son 1 haftada %20'den fazla yükselenler alınmaz
#     (backtestteki kazancın bir kısmı alınması zor tavan serilerinden geliyordu)
#   - Borsa İstanbul tedbirli (VBTS: brüt takas, tek fiyat, açığa satış yasağı) hisseler
#     alınmaz → aracı kurum uygulamandan kontrol et, listede elle işaretle
#   - Hisse başına dar stop YOK. Portföy düzeyinde: zirveden −%25 → yarıya indir,
#     −%35 → dur. Son 12 ayda likit hisse ortalamasının 10 puandan fazla gerisindeyse → bırak
#   - Gerçekçi beklenti: likit hisse ortalamasından ayda +%0,5 … +%1 fazla (backtestteki
#     +%2,5 değil); %25-35 ihtimalle avantaj hiç yok. Bu bir deneme: ilk 6 ay küçük tutarla.
#  v1.1: Hücre 2'nin verisi bozuksa (fiyatlar boş, hisse sayısı çok az) liste OLUŞTURULMUYOR
#        ve kaydedilmiyor; bozuk veriyle yapılmış eski bir kayıt varsa siliniyor.
#        Fiyatların hangi günün kapanışı olduğu yazılıyor.
#  v1.2: GERÇEK PORTFÖY: Drive'daki gercek_portfoy.csv'den (Tarih, Kod, Islem, Adet, Fiyat)
#        aldığın hisseleri okuyup maliyet, kâr/zarar ve TUT / SAT durumunu gösteriyor;
#        kâğıt üzerindeki 30 hisseyle karşılaştırıyor.
#  v1.3: Gerçek portföyün fiyatı İsyatırım'ın son kapanışından (Yahoo bir günü boş bıraksa da güncel).
# ═══════════════════════════════════════════════════════════════
import os
import pandas as pd
import numpy as np
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

DRIVE       = '/content/drive/MyDrive/isyatirim'
SISTEM_YOLU = f'{DRIVE}/BIST_Sistem.xlsx'
EXCEL_YOLU  = f'{DRIVE}/BIST_Tam_KAP_v9_11mart.xlsx'
KAYIT_YOLU  = f'{DRIVE}/zirve_kayit.csv'
PORTFOY_YOLU = f'{DRIVE}/gercek_portfoy.csv'   # gerçek alım-satımların: Tarih,Kod,Islem(AL/SAT),Adet,Fiyat

# ── AYARLAR ──────────────────────────────────────────────────
STRATEJI_PARASI = 0        # bu stratejiye ayırdığın TL (0 = sadece liste; girersen adet hesaplanır)
HISSE_SAYISI    = 30
MIN_TL_HACIM    = 50       # mn TL, 20 günlük ortalama
MIN_FIYAT       = 5        # TL
MAX_HAFTALIK    = 20       # %, son 1 hafta
HARIC           = []       # tedbirli (VBTS) ya da almak istemediğin hisseler, ör. ['ABCDE', 'XYZAB']
YENIDEN_SEC     = False    # True: bu ay liste zaten seçilmiş olsa bile yeniden seç
YARIYA_INDIR    = -25      # strateji değeri zirvesinden bu kadar düşerse: yarıya indir
DUR             = -35      # bu kadar düşerse: dur
BIRAK_12AY      = -10      # son 12 ayda likit ortalamanın bu kadar puan gerisindeyse: bırak

print("=" * 65)
print("  🧭 ZİRVE STRATEJİSİ — aylık liste ve gerçek sonuç takibi")
print(f"  {datetime.now():%d.%m.%Y %H:%M}")
print("=" * 65)

df = pd.read_excel(SISTEM_YOLU)
df = df.rename(columns={'Sembol': 'Kod'}) if 'Sembol' in df.columns else df
gerekli = ['Kod', 'Son Fiyat', 'Zirveden (%)', 'Haftalık (%)', 'Ort. TL Hacim (mn)']
eksik = [c for c in gerekli if c not in df.columns]
if eksik:
    raise SystemExit(f"❌ BIST_Sistem.xlsx'te eksik sütun: {eksik} — önce hücre 2'yi çalıştır.")
# ── VERİ KONTROLÜ: bozuk veriyle liste seçilmesin ──
MIN_HISSE = 250        # hücre 2 normalde ~370 hisse çıkarır
_bos = df['Son Fiyat'].isna() | pd.to_numeric(df['Zirveden (%)'], errors='coerce').isna()
if _bos.mean() > 0.10:
    raise SystemExit(f"❌ {int(_bos.sum())}/{len(df)} hissenin fiyatı ya da zirve verisi BOŞ. Hücre 2'nin verisi bozuk "
                     f"(Yahoo fiyat göndermemiş). Liste oluşturulmadı, hiçbir şey kaydedilmedi.\n"
                     f"   → Hücre 2'yi (v2.5 ve üstü) tekrar çalıştır, sonra bu hücreyi.")
if _bos.any():
    print(f"ℹ️  Fiyatı boş {int(_bos.sum())} hisse dışarıda bırakıldı: {', '.join(df.loc[_bos, 'Kod'].astype(str).head(10))}")
    df = df[~_bos].copy()
if len(df) < MIN_HISSE:
    raise SystemExit(f"❌ BIST_Sistem'de sadece {len(df)} hisse var (normalde ~370). Hücre 2 yarım kalmış olabilir. "
                     f"Liste oluşturulmadı, hiçbir şey kaydedilmedi.")
if 'Veri Tarihi' in df.columns:
    _vt = pd.to_datetime(df['Veri Tarihi']).mode().iloc[0]
    _simdi = datetime.now(); _bek = pd.Timestamp(_simdi.date())
    if not (_bek.weekday() < 5 and (_simdi.hour, _simdi.minute) >= (18, 15)):
        _bek = _bek - pd.offsets.BDay(1)
    print(f"📅 Fiyatlar {_vt:%d.%m.%Y} kapanışından")
    if _vt < _bek:
        print(f"⚠️  Beklenen son kapanış {_bek:%d.%m.%Y} — fiyatlar 1 gün eski (bayram tatili değilse). "
              f"Mümkünse 1-2 saat sonra hücre 2'yi tekrar çalıştırıp sonra bu hücreyi çalıştır.")
yas_saat = (datetime.now().timestamp() - os.path.getmtime(SISTEM_YOLU)) / 3600
print(f"✅ BIST_Sistem: {len(df)} hisse | dosya {yas_saat:.0f} saat önce güncellenmiş")
if yas_saat > 24:
    print("⚠️  Veri 1 günden eski. Liste için önce hücre 2'yi çalıştır.")

bugun = pd.Timestamp(datetime.now().date())
fiyat_simdi = df.set_index('Kod')['Son Fiyat'].astype(float)

# ── KAYIT ────────────────────────────────────────────────────
kayit = pd.read_csv(KAYIT_YOLU, parse_dates=['Tarih']) if os.path.exists(KAYIT_YOLU) else \
        pd.DataFrame(columns=['Tarih', 'Kod', 'Fiyat', 'Evren', 'Secildi'])
# bozuk veriyle yapılmış kayıt varsa (seçilen hisse yok / fiyatlar boş) sil
_bozuk = [t for t, g in kayit.groupby('Tarih')
          if g['Secildi'].astype(bool).sum() < 10 or pd.to_numeric(g['Fiyat'], errors='coerce').isna().mean() > 0.10]
if _bozuk:
    kayit = kayit[~kayit['Tarih'].isin(_bozuk)]
    kayit.to_csv(KAYIT_YOLU, index=False)
    print(f"🧹 Bozuk veriyle yapılmış kayıt silindi: {', '.join(pd.Timestamp(t).strftime('%d.%m.%Y') for t in _bozuk)}")
tarihler = sorted(kayit['Tarih'].unique()) if len(kayit) else []
son_secim = pd.Timestamp(tarihler[-1]) if tarihler else None
bu_ay_var = son_secim is not None and (son_secim.year, son_secim.month) == (bugun.year, bugun.month)

# ── LİSTE ────────────────────────────────────────────────────
evren = df[pd.to_numeric(df['Ort. TL Hacim (mn)'], errors='coerce') >= MIN_TL_HACIM].copy()
aday = evren[(evren['Son Fiyat'] >= MIN_FIYAT) & (evren['Haftalık (%)'] <= MAX_HAFTALIK)
             & (~evren['Kod'].isin([h.upper() for h in HARIC]))]
yeni_liste = aday.sort_values('Zirveden (%)', ascending=False).head(HISSE_SAYISI)

if bu_ay_var and not YENIDEN_SEC:
    secili_kodlar = kayit[(kayit['Tarih'] == son_secim) & (kayit['Secildi'])]['Kod'].tolist()
    print(f"ℹ️  Bu ayın listesi {son_secim:%d.%m.%Y}'de seçilmiş; aynı liste gösteriliyor "
          f"(yeniden seçmek için YENIDEN_SEC = True).")
    liste = df[df['Kod'].isin(secili_kodlar)].sort_values('Zirveden (%)', ascending=False).copy()
else:
    liste = yeni_liste.copy()
    onceki = set(kayit[(kayit['Tarih'] == son_secim) & (kayit['Secildi'])]['Kod']) if son_secim is not None else set()
    # bu ayın seçimini kaydet: tüm hisselerin fiyatı (sonraki ay getiriyi ölçmek için)
    yeni = pd.DataFrame({'Tarih': bugun, 'Kod': df['Kod'], 'Fiyat': df['Son Fiyat'],
                         'Evren': df['Kod'].isin(evren['Kod']), 'Secildi': df['Kod'].isin(liste['Kod'])})
    kayit = pd.concat([kayit[kayit['Tarih'] != bugun], yeni], ignore_index=True)
    kayit.to_csv(KAYIT_YOLU, index=False)
    tarihler = sorted(kayit['Tarih'].unique())
    print(f"📝 Yeni liste seçildi ve kaydedildi ({len(liste)} hisse).")
    if onceki:
        cik = sorted(onceki - set(liste['Kod'])); gir = sorted(set(liste['Kod']) - onceki)
        print(f"   🔴 Listeden çıkan (sat): {', '.join(cik) if cik else 'yok'}")
        print(f"   🟢 Listeye giren (al):  {', '.join(gir) if gir else 'yok'}")

onceki_kodlar = set()
if len(tarihler) >= 2:
    onceki_kodlar = set(kayit[(kayit['Tarih'] == pd.Timestamp(tarihler[-2])) & (kayit['Secildi'])]['Kod'])

# ── GERÇEK SONUÇ TAKİBİ (kağıt üzerinde, eşit ağırlık) ──────
donemler = []
for i, t in enumerate(tarihler):
    t = pd.Timestamp(t)
    g = kayit[kayit['Tarih'] == t].set_index('Kod')
    if i + 1 < len(tarihler):
        son = kayit[kayit['Tarih'] == pd.Timestamp(tarihler[i + 1])].set_index('Kod')['Fiyat']
        bitis = pd.Timestamp(tarihler[i + 1]); acik = False
    else:
        son = fiyat_simdi; bitis = bugun; acik = True
    if bitis <= t:
        continue
    r = (son.reindex(g.index) / g['Fiyat'] - 1) * 100
    supheli = r[r < -45].index.tolist()          # bedelsiz/bölünme olabilir → ölçüme katma
    r = r.drop(supheli)
    s = r[g['Secildi'].reindex(r.index).fillna(False).astype(bool)].dropna()
    e = r[g['Evren'].reindex(r.index).fillna(False).astype(bool)].dropna()
    if len(s) == 0:
        continue
    donemler.append({'Başlangıç': t, 'Bitiş': bitis, 'Açık': acik, 'Strateji (%)': s.mean(),
                     'Likit Ort. (%)': e.mean(), 'Fark (%)': s.mean() - e.mean(),
                     'Hisse': len(s), 'Şüpheli': ', '.join(supheli)})
tk = pd.DataFrame(donemler)
durum = []
if len(tk):
    deger = (1 + tk['Strateji (%)'] / 100).cumprod() * 100
    zirve = deger.cummax()
    dusus = (deger / zirve - 1) * 100
    tk['Strateji Değeri'] = deger.round(1); tk['Zirveden Düşüş (%)'] = dusus.round(1)
    son12 = tk[tk['Başlangıç'] >= bugun - pd.DateOffset(months=12)]
    fark12 = ((1 + son12['Strateji (%)'] / 100).prod() - (1 + son12['Likit Ort. (%)'] / 100).prod()) * 100
    durum.append(f"Strateji değeri: {deger.iloc[-1]:.1f} (başlangıç 100) | zirvesinden düşüş %{dusus.iloc[-1]:.1f}")
    durum.append(f"Son 12 ay: likit hisse ortalamasına göre {fark12:+.1f} puan ({len(son12)} dönem)")
    if dusus.iloc[-1] <= DUR:
        durum.append(f"🔴 DUR: strateji zirvesinden %{dusus.iloc[-1]:.0f} düştü (sınır %{DUR}). Pozisyonları kapat, yeniden değerlendir.")
    elif dusus.iloc[-1] <= YARIYA_INDIR:
        durum.append(f"🟠 YARIYA İNDİR: strateji zirvesinden %{dusus.iloc[-1]:.0f} düştü (sınır %{YARIYA_INDIR}).")
    elif len(son12) >= 10 and fark12 <= BIRAK_12AY:
        durum.append(f"🔴 BIRAK: son 12 ayda likit ortalamanın {fark12:.0f} puan gerisinde (sınır {BIRAK_12AY}).")
    else:
        durum.append("🟢 Kurallar içinde — devam.")
    if len(tk) < 6:
        durum.append(f"ℹ️  Henüz {len(tk)} dönem var; ilk 6 ay küçük tutarla devam et. 12 ay bile kesin sonuç vermez.")
else:
    durum.append("ℹ️  İlk liste: takip bir sonraki çalıştırmada başlar.")
for d in durum:
    print('  ' + d)

# ── TUTAR / ADET ─────────────────────────────────────────────
liste = liste.copy()
liste['Ağırlık (%)'] = round(100 / max(len(liste), 1), 2)
liste['Durum'] = liste['Kod'].apply(lambda k: 'Devam' if k in onceki_kodlar else 'Yeni')
if len(liste) == 0:
    raise SystemExit("❌ Listede hisse çıkmadı — filtreleri (MIN_TL_HACIM, MIN_FIYAT, HARIC) kontrol et.")
if STRATEJI_PARASI and STRATEJI_PARASI > 0:
    pay = STRATEJI_PARASI / len(liste)
    liste['Tutar (TL)'] = round(pay, 0)
    liste['Adet'] = (pay // liste['Son Fiyat']).astype(int)
kol = ['Kod', 'Durum', 'Son Fiyat', 'Zirveden (%)', 'Haftalık (%)', 'Aylık (%)', 'Ort. TL Hacim (mn)',
       'İSY Sektör', 'Ağırlık (%)'] + (['Tutar (TL)', 'Adet'] if 'Adet' in liste else [])
kol = [c for c in kol if c in liste.columns]
print(f"\n📋 BU AYIN LİSTESİ ({len(liste)} hisse) — tedbirli (VBTS) olanları aracı kurumundan kontrol et, HARIC'e yaz:")
print(liste[kol].to_string(index=False))
if 'İSY Sektör' in liste:
    sk = liste['İSY Sektör'].value_counts()
    if sk.iloc[0] > len(liste) * 0.25:
        print(f"⚠️  Sektör yoğunlaşması: {sk.index[0]} {sk.iloc[0]} hisse")

# ── GERÇEK PORTFÖY ───────────────────────────────────────────
gp = None; gp_ozet = []
if os.path.exists(PORTFOY_YOLU):
    ip = pd.read_csv(PORTFOY_YOLU)
    ip['Tarih'] = pd.to_datetime(ip['Tarih'])
    ip['Kod'] = ip['Kod'].astype(str).str.upper().str.strip()
    ip['Islem'] = ip['Islem'].astype(str).str.upper().str.strip()
    tum_fiyat = pd.read_excel(SISTEM_YOLU).rename(columns={'Sembol': 'Kod'}).set_index('Kod')['Son Fiyat']
    fiyat_kaynak = 'BIST_Sistem (Yahoo) kapanışı'
    ozet_yolu = f'{DRIVE}/temelozet.xlsx'               # İsyatırım: her gün hücre 1 indiriyor, en güncel kapanış
    try:
        if os.path.exists(ozet_yolu) and (datetime.now().timestamp() - os.path.getmtime(ozet_yolu)) < 36 * 3600:
            oz = pd.read_excel(ozet_yolu)
            kk = next(c for c in oz.columns if str(c).replace(' ', '').startswith('Kapanış'))
            isy = pd.to_numeric(oz.set_index(oz['Kod'].astype(str).str.upper().str.strip())[kk], errors='coerce').dropna()
            tum_fiyat = isy.combine_first(tum_fiyat)
            fiyat_kaynak = 'İsyatırım son kapanış'
    except Exception:
        pass
    satirlar = []; gerceklesen = 0.0
    for kod, g in ip.sort_values('Tarih', kind='stable').groupby('Kod', sort=False):
        adet = 0.0; maliyet = 0.0                      # ortalama maliyet yöntemi
        for _, r in g.iterrows():
            if r['Islem'] == 'AL':
                maliyet += r['Adet'] * r['Fiyat']; adet += r['Adet']
            elif r['Islem'] == 'SAT' and adet > 0:
                n = min(r['Adet'], adet); ort = maliyet / adet
                gerceklesen += n * (r['Fiyat'] - ort); maliyet -= n * ort; adet -= n
        if adet <= 0:
            continue
        fiyat = tum_fiyat.get(kod, np.nan)
        if pd.isna(fiyat):                               # hücre 2'de yoksa (ör. zarar eden şirket) Yahoo'dan
            try:
                import yfinance as yf
                fiyat = float(yf.Ticker(f'{kod}.IS').history(period='10d')['Close'].dropna().iloc[-1])
            except Exception:
                fiyat = np.nan
        ort = maliyet / adet
        listede = kod in set(liste['Kod'])
        satirlar.append({'Kod': kod, 'Adet': int(adet), 'Ort. Maliyet': round(ort, 3), 'Son Fiyat': fiyat,
                         'Maliyet (TL)': round(maliyet, 0), 'Değer (TL)': round(adet * fiyat, 0) if pd.notna(fiyat) else None,
                         'K/Z (TL)': round(adet * (fiyat - ort), 0) if pd.notna(fiyat) else None,
                         'K/Z (%)': round((fiyat / ort - 1) * 100, 1) if pd.notna(fiyat) else None,
                         'İlk Alış': g['Tarih'].min().strftime('%d.%m.%Y'),
                         'Karar': '✅ TUT (listede)' if listede else '🔴 SAT (bu ayın listesinde yok)'})
    if satirlar:
        gp = pd.DataFrame(satirlar)
        fiyatli = gp[gp['Değer (TL)'].notna()]
        top_m = fiyatli['Maliyet (TL)'].sum(); top_d = fiyatli['Değer (TL)'].sum(min_count=1)
        print(f"\n💼 GERÇEK PORTFÖYÜN ({len(gp)} hisse) — fiyatlar: {fiyat_kaynak}")
        print(gp.to_string(index=False))
        if pd.notna(top_d) and top_m > 0:
            gp_ozet.append(f"Toplam: maliyet {top_m:,.0f} TL → değer {top_d:,.0f} TL | K/Z {top_d - top_m:+,.0f} TL (%{(top_d / top_m - 1) * 100:+.1f})")
        fiyatsiz = gp[gp['Değer (TL)'].isna()]['Kod'].tolist()
        if fiyatsiz:
            gp_ozet.append(f"⚠️  Fiyatı alınamayan (toplama katılmadı): {', '.join(fiyatsiz)}")
        if gerceklesen:
            gp_ozet.append(f"Satışlardan gerçekleşen K/Z: {gerceklesen:+,.0f} TL")
        if len(tk):
            son = tk.iloc[-1]
            gp_ozet.append(f"Kâğıt üzerindeki strateji, son dönem ({son['Başlangıç']:%d.%m} → {son['Bitiş']:%d.%m}): "
                           f"30 hisse %{son['Strateji (%)']:+.1f} | likit hisse ortalaması %{son['Likit Ort. (%)']:+.1f}")
        sat = gp[gp['Karar'].str.startswith('🔴')]['Kod'].tolist()
        if sat:
            gp_ozet.append(f"🔴 Bu ayın listesinde olmayan: {', '.join(sat)} → kurala göre sat (ayın ilk işlem günü)")
        for x in gp_ozet:
            print('  ' + x)
else:
    print(f"\nℹ️  Gerçek portföy dosyası yok ({os.path.basename(PORTFOY_YOLU)}). Aldığın hisseleri takip etmek için "
          f"Drive'a bu adla Tarih,Kod,Islem,Adet,Fiyat sütunlu bir CSV koy.")

# ── EXCEL ────────────────────────────────────────────────────
wb = load_workbook(EXCEL_YOLU)
ad = '🧭 Zirve Stratejisi'
if ad in wb.sheetnames:
    del wb[ad]
ws = wb.create_sheet(ad, 0)
ws['A1'] = f"🧭 ZİRVE STRATEJİSİ — {datetime.now():%d.%m.%Y %H:%M} | {len(liste)} hisse | eşit ağırlık | ayda bir yenile"
ws['A1'].font = Font(bold=True, size=13, color='FFFFFFFF'); ws['A1'].fill = PatternFill('solid', fgColor='FF0F3460')
s = 2
for d in durum:
    s += 1; ws.cell(s, 1, d).font = Font(bold=True)
s += 2
for ci, c in enumerate(['#'] + kol, 1):
    x = ws.cell(s, ci, c); x.font = Font(bold=True, color='FFFFFFFF'); x.fill = PatternFill('solid', fgColor='FF1A1A2E')
for i, (_, r) in enumerate(liste[kol].iterrows(), 1):
    ws.cell(s + i, 1, i)
    for ci, c in enumerate(kol, 2):
        v = r[c]
        v = None if isinstance(v, float) and np.isnan(v) else (v.item() if hasattr(v, 'item') else v)
        ws.cell(s + i, ci, v)
s += len(liste) + 3
if gp is not None:
    ws.cell(s, 1, 'GERÇEK PORTFÖYÜN').font = Font(bold=True); s += 1
    for ci, c in enumerate(gp.columns, 1):
        x = ws.cell(s, ci, c); x.font = Font(bold=True, color='FFFFFFFF'); x.fill = PatternFill('solid', fgColor='FF0E7566')
    for i, (_, r) in enumerate(gp.iterrows(), 1):
        for ci, c in enumerate(gp.columns, 1):
            v = r[c]; v = None if isinstance(v, float) and np.isnan(v) else (v.item() if hasattr(v, 'item') else v)
            ws.cell(s + i, ci, v)
    s += len(gp) + 1
    for x in gp_ozet:
        s += 1; ws.cell(s, 1, x).font = Font(bold=True)
    s += 3
kurallar = [
    "KURALLAR (backtest + bağımsız istatistik incelemesi):",
    f"• Ayın ilk işlem günü yenile: listeden çıkanı sat, gireni al. Her hisse eşit (~%{100 / max(len(liste), 1):.1f}); hiçbiri %5'i geçmesin.",
    f"• {MIN_FIYAT} TL altı, son haftada %{MAX_HAFTALIK}+ yükselen ve VBTS tedbirli hisseler alınmaz.",
    f"• Hisse başına dar stop yok. Strateji değeri zirvesinden %{YARIYA_INDIR} → yarıya indir, %{DUR} → dur.",
    f"• Son 12 ayda likit hisse ortalamasının {abs(BIRAK_12AY)} puandan fazla gerisindeyse stratejiyi bırak.",
    "• Kayıp olursa kuralları DEĞİŞTİRME (geçmişe uydurma olur). Değişiklik ancak 12 ay sonra, kayda bakarak.",
    "• Gerçekçi beklenti: likit ortalamadan ayda +%0,5 … +%1 fazla; %25-35 ihtimalle avantaj yok.",
    "• %30-40'lık düşüşleri göze al: bu stratejiye sadece böyle bir düşüşte birikiminin %5-8'ini kaybedeceğin kadar para koy.",
    "• Eğitim amaçlıdır, yatırım tavsiyesi değildir.",
]
for k in kurallar:
    ws.cell(s, 1, k).font = Font(bold=k.startswith('KURALLAR')); s += 1
if len(tk):
    s += 1
    t2 = tk.copy()
    t2['Başlangıç'] = t2['Başlangıç'].dt.strftime('%d.%m.%Y'); t2['Bitiş'] = t2['Bitiş'].dt.strftime('%d.%m.%Y')
    t2['Açık'] = t2['Açık'].map({True: 'devam ediyor', False: 'kapandı'})
    for c in ['Strateji (%)', 'Likit Ort. (%)', 'Fark (%)']:
        t2[c] = t2[c].round(2)
    ws.cell(s, 1, 'GERÇEK SONUÇ (kağıt üzerinde, temettü hariç)').font = Font(bold=True); s += 1
    for ci, c in enumerate(t2.columns, 1):
        x = ws.cell(s, ci, c); x.font = Font(bold=True, color='FFFFFFFF'); x.fill = PatternFill('solid', fgColor='FF1A1A2E')
    for i, (_, r) in enumerate(t2.iterrows(), 1):
        for ci, c in enumerate(t2.columns, 1):
            v = r[c]; v = v.item() if hasattr(v, 'item') else v
            ws.cell(s + i, ci, v)
for i, w in enumerate([5, 9, 8, 10, 12, 12, 10, 16, 28, 11, 11, 8], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
wb.save(EXCEL_YOLU)
print(f"\n✅ '{ad}' sayfası yazıldı (Excel'in ilk sayfası). Kayıt: zirve_kayit.csv")
