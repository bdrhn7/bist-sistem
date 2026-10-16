# ═══════════════════════════════════════════════════════════════
#  HÜCRE 7 — KAPSAMLI ANALİZ MOTORU  (v4.9 — 16.10.2026)
#  Güven + İSYATİRIM + WACC/DCF + Veto + Giriş/Çıkış + Portföy
#  v4 değişiklikleri:
#   - DCF iki aşamalı, net borç düşülüyor, 3 senaryo, tersten DCF
#   - RF %35,6 ve kurumlar vergisi %25 güncellendi
#   - Portföy boşaltıldı, SERMAYE en üste taşındı
#   - Yabancı oranı sütunlarında tüm aylar tanınıyor
#   - v4.1: büyüme tavanı %30, sektör İsyatırım'dan, GYO'da DCF yok
#   - v4.2: DCF'de hisse sayısı İsyatırım piyasa değeriyle kontrol ediliyor;
#           seans açıkken bugünün yarım günlük verisi kullanılmıyor (ATR, stop, hedef)
#   - v4.3: Makro Radar 3 saatten eskiyse uyarı veriyor (önce hücre 8 çalışmalı)
#   - v4.4: dolar/euro ile rapor veren şirketlerde DCF rakamları TL'ye çevriliyor;
#           hukuki risk vetosu artık yapay zekânın açık 'RİSK: VAR' cevabına bakıyor
#           (gerekçede 'dava içermiyor' yazınca yanlışlıkla veto oluyordu)
#   - v4.5: Risk/Ödül artık gerçek: (sistem hedefi − fiyat) ÷ (fiyat − normal stop).
#           Eskiden 4 ATR ÷ 2,5 ATR = her hissede 1,6 çıkıyor ve herkese +1 puan veriyordu.
#   - v4.5: DCF'de son 3 yılın serbest nakit akışı bugünkü TL'ye çevrilip ortalanıyor
#           (eski yılların düşük enflasyonlu rakamları ortalamayı ~%25 aşağı çekiyordu)
#   - v4.6: Risk/Ödül eşikleri: 5+ → +2, 2-5 → +1, 1-2 → 0, 1 altı → −1. Hedef 12 aylık,
#           stop kısa vadeli olduğu için oranlar büyük çıkıyor; ortalama katkı eskisi gibi ~+1.
#   - v4.6: Süper AL'e girmek için giriş kararı da ŞİMDİ GİR ya da HAZIR OL olmalı
#           (makro "yeni pozisyon açma" derken listede hisse görünüyordu)
#   - v4.7: DCF'de son yılın nakit akışı 3 yıllık ortalamanın altındaysa son yıl kullanılıyor
#           (nakit akışı düşen şirkette eski yüksek seviyeye döneceği varsayılmıyor)
#   - v4.8: her çalıştırmada kararlar Drive'da sistem_kayit.csv'ye yazılıyor (aynı gün
#           tekrar çalışırsa o günün kaydı yenilenir). Hücre 9 bu kaydı ileride değerlendirir:
#           yapay zekânın AL/NÖTR/SAT kararları ve Öneri/Giriş kararları gerçekten işe yaradı mı?
#   v4.9: Yahoo fiyatı boş (NaN) satır gönderirse o satır atılıyor (16.10 sabahı yaşandı)
# ═══════════════════════════════════════════════════════════════

import yfinance as yf
import pandas as pd
import numpy as np
import os
import time
import warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# ══════════════════════════════════════════════════════════════
#  AYARLAR
# ══════════════════════════════════════════════════════════════
SISTEM_YOLU = '/content/drive/MyDrive/isyatirim/BIST_Sistem.xlsx'
EXCEL_YOLU  = '/content/drive/MyDrive/isyatirim/BIST_Tam_KAP_v9_11mart.xlsx'
DRIVE       = '/content/drive/MyDrive/isyatirim'

# WACC parametreleri — ayda bir kontrol et
RF           = 0.356   # Türkiye 10 yıllık tahvil (14.10.2026: ~%35,6)
PIYASA_RISKI = 0.055   # Piyasa risk primi ~%5.5
VERGI_ORANI  = 0.25    # Kurumlar vergisi (genel oran %25)

# TÜİK yıllık TÜFE (Aralık, %) — hücre 2 ile aynı. Her ocakta yeni yılı ekle.
YILLIK_TUFE = {2020: 14.60, 2021: 36.08, 2022: 64.27, 2023: 64.77,
               2024: 44.38, 2025: 30.89}
VARSAYILAN_TUFE = 30.0

# BİLEŞİK DÖNDÜRME PARAMETRELERİ
BILESIK_HEDEF_YUZ = 12
BILESIK_SURE_GUN  = 45

# Portföy — hisse alınca buraya ekle. Örnek:
# {'Kod': 'XXXXX', 'Adet': 100, 'Alis': 10.50, 'Stop': 9.20, 'Tarih': '2026-10-15'},
PORTFOY = [
]

# Pozisyon büyüklüğü hesabında kullanılan toplam sermaye (TL)
SERMAYE = 25000

BORDER = Border(
    left=Side(style='thin', color='1E3A5F'),
    right=Side(style='thin', color='1E3A5F'),
    top=Side(style='thin', color='1E3A5F'),
    bottom=Side(style='thin', color='1E3A5F')
)

# ══════════════════════════════════════════════════════════════
#  YARDIMCI EXCEL FONKSİYONLARI
# ══════════════════════════════════════════════════════════════
def hc(ws, r, c, v, bg='051020', fg='00FF88', sz=10):
    x = ws.cell(r, c, v)
    x.font = Font(name='Arial', bold=True, color=fg, size=sz)
    x.fill = PatternFill('solid', fgColor=bg)
    x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    x.border = BORDER
    return x

def dc(ws, r, c, v, bg='091422', fg='E8F4F8', bold=False, sz=9, wrap=False, align='center'):
    try:
        val = None if isinstance(v, float) and np.isnan(v) else v
    except:
        val = v
    x = ws.cell(r, c, val)
    x.font = Font(name='Arial', bold=bold, color=fg, size=sz)
    x.fill = PatternFill('solid', fgColor=bg)
    x.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
    x.border = BORDER
    return x

def tamamlanmis_gunler(df_h):
    """Fiyatı boş gelen satırları atar; seans kapanmadan (18:15'ten önce) bugünün
    satırlarını (yarım ya da boş mum) çıkarır."""
    try:
        if 'Close' in df_h.columns:
            df_h = df_h[df_h['Close'].notna() & (df_h['Close'] > 0)]
        simdi = datetime.now()
        if len(df_h) and (simdi.hour, simdi.minute) < (18, 15):
            df_h = df_h[[d.date() != simdi.date() for d in df_h.index]]
    except: pass
    return df_h

# ══════════════════════════════════════════════════════════════
#  MODÜL 1: GÜVEN SKORU
#  Verinin ne kadar güvenilir olduğunu ölçer (0-100)
# ══════════════════════════════════════════════════════════════
def guven_skoru_hesapla(kod):
    skor = 0
    detay = []
    uyarilar = []
    try:
        t    = yf.Ticker(f'{kod}.IS')
        info = t.info
        inc  = t.income_stmt
        hist = tamamlanmis_gunler(t.history(period='2y'))

        # Finansal veri yeterliliği (0-25)
        if inc is not None and not inc.empty and 'Net Income' in inc.index:
            yil = len([c for c in inc.columns if pd.notna(inc.loc['Net Income', c])])
            if yil >= 4:   skor += 25; detay.append(f'✅ {yil} yıl veri')
            elif yil >= 3: skor += 18; detay.append(f'🟡 {yil} yıl veri')
            elif yil >= 2: skor += 10; detay.append(f'🟠 {yil} yıl veri')
            else:          skor +=  0; detay.append('🔴 Veri az'); uyarilar.append('Veri az')
        else:
            detay.append('🔴 Finansal veri yok'); uyarilar.append('Finansal veri yok')

        # Hacim yeterliliği (0-20)
        if hist is not None and len(hist) > 20:
            hist['TL'] = hist['Close'] * hist['Volume']
            tl = hist['TL'].iloc[-20:].mean() / 1_000_000
            if tl >= 100:   skor += 20; detay.append(f'✅ Hacim {tl:.0f}mn')
            elif tl >= 30:  skor += 15; detay.append(f'🟡 Hacim {tl:.0f}mn')
            elif tl >= 5:   skor +=  8; detay.append(f'🟠 Hacim {tl:.0f}mn')
            else:           skor +=  0; detay.append(f'🔴 Hacim {tl:.1f}mn'); uyarilar.append('Likidite riski')
        else:
            detay.append('🔴 Fiyat verisi yok'); uyarilar.append('Fiyat verisi yok')

        # Kar tutarlılığı (0-20)
        if inc is not None and not inc.empty and 'Net Income' in inc.index:
            karlar = [float(inc.loc['Net Income', c]) for c in inc.columns[:4] if pd.notna(inc.loc['Net Income', c])]
            if len(karlar) >= 2:
                poz = sum(1 for k in karlar if k > 0)
                oran = poz / len(karlar)
                if oran == 1.0:    skor += 20; detay.append('✅ Tüm dönemler karlı')
                elif oran >= 0.75: skor += 14; detay.append(f'🟡 {poz}/{len(karlar)} dönem karlı')
                elif oran >= 0.5:  skor +=  7; detay.append(f'🟠 {poz}/{len(karlar)} dönem karlı')
                else:              skor +=  0; detay.append(f'🔴 {poz}/{len(karlar)} dönem karlı'); uyarilar.append('Zarar riski')

        # Bilanço tazeliği (0-20)
        if inc is not None and not inc.empty:
            try:
                fark = (pd.Timestamp.now() - pd.Timestamp(inc.columns[0])).days / 30
                if fark <= 4:   skor += 20; detay.append(f'✅ Bilanço {fark:.0f}ay')
                elif fark <= 6: skor += 14; detay.append(f'🟡 Bilanço {fark:.0f}ay')
                elif fark <= 9: skor +=  7; detay.append(f'🟠 Bilanço {fark:.0f}ay')
                else:           skor +=  0; detay.append(f'🔴 Bilanço {fark:.0f}ay eski'); uyarilar.append('Eski bilanço')
            except:
                skor += 10; detay.append('🟡 Bilanço tarihi belirsiz')

        # Piyasa değeri (0-15)
        mc = info.get('marketCap', 0)
        if mc and mc > 0:
            mc_mn = mc / 1_000_000
            if mc_mn >= 1000:   skor += 15; detay.append(f'✅ PD {mc_mn:.0f}mn')
            elif mc_mn >= 200:  skor += 10; detay.append(f'🟡 PD {mc_mn:.0f}mn')
            elif mc_mn >= 50:   skor +=  5; detay.append(f'🟠 PD {mc_mn:.0f}mn')
            else:               skor +=  0; detay.append(f'🔴 PD {mc_mn:.0f}mn küçük'); uyarilar.append('Küçük şirket')
        else:
            skor += 5; detay.append('🟡 PD bilinmiyor')

    except Exception as e:
        return 0, '🔴 Güvenilmez', f'Hata: {str(e)[:40]}', ['Veri hatası']

    if skor >= 85:   kat = '🟢 Çok Güvenilir'
    elif skor >= 70: kat = '🟢 Güvenilir'
    elif skor >= 50: kat = '🟡 Orta Güven'
    elif skor >= 30: kat = '🟠 Düşük Güven'
    else:            kat = '🔴 Güvenilmez'

    return skor, kat, ' | '.join(detay), uyarilar

def guven_duzelt(skor, guven):
    if guven < 30:   return round(skor * 0.60), '-%40'
    elif guven < 50: return round(skor * 0.80), '-%20'
    elif guven < 70: return round(skor * 0.90), '-%10'
    else:            return skor, 'OK'

# ══════════════════════════════════════════════════════════════
#  MODÜL 2: WACC + DCF  (v2 — iki aşamalı, net borç düşülmüş)
#  - 5 yıl büyüme + uzun vade büyümesi (iki aşama)
#  - Şirket değerinden NET BORÇ düşülür → hissedar değeri
#  - Serbest nakit akışı: son 3 yılın ortalaması
#  - 3 senaryo: Kötü / Orta / İyi → değer aralığı
#  - Tersten DCF: fiyat yılda % kaç büyüme bekliyor?
#  - Banka/sigorta gibi finans şirketlerinde DCF uygulanmaz
# ══════════════════════════════════════════════════════════════
DCF_YIL          = 5      # Birinci aşama kaç yıl
UZUN_VADE_BUYUME = 0.15   # 5. yıldan sonra yıllık büyüme (TL, enflasyon dahil)
BUYUME_TAVAN     = 0.30   # İlk 5 yıl için en fazla büyüme (geçmiş yüksek enflasyon şişirmesin)
BUYUME_TABAN     = 0.00   # İlk 5 yıl için en az büyüme

_KUR_CACHE = {}
def kur_serisi(para):
    """'USD' → USD/TRY günlük kapanış serisi. TRY için None."""
    para = (para or 'TRY').upper()
    if para in ('TRY', 'TL'):
        return None
    if para not in _KUR_CACHE:
        k = yf.Ticker(f'{para}TRY=X').history(period='6y')['Close']
        if getattr(k.index, 'tz', None) is not None:
            k.index = k.index.tz_localize(None)
        _KUR_CACHE[para] = k
    return _KUR_CACHE[para]

def kur(seri, tarih=None):
    """Tarihteki (ya da öncesindeki son) kur; tarih yoksa son kur. seri None ise 1."""
    if seri is None:
        return 1.0
    if tarih is None:
        return float(seri.iloc[-1])
    o = seri[seri.index <= pd.Timestamp(tarih)]
    return float(o.iloc[-1]) if len(o) else float(seri.iloc[0])

def _ilk_deger(df, anahtarlar, sutun=0):
    """Tablodan verilen satır isimlerinden ilk bulunanı döndürür."""
    if df is None or df.empty:
        return None
    for k in anahtarlar:
        if k in df.index and len(df.columns) > sutun:
            v = df.loc[k].iloc[sutun]
            if pd.notna(v):
                return float(v)
    return None

def _enf_carpani(yil_eski, yil_yeni):
    """yil_eski'deki 1 TL, yil_yeni sonunda kaç TL eder (TÜFE ile)."""
    c = 1.0
    for y in range(int(yil_eski) + 1, int(yil_yeni) + 1):
        c *= 1 + YILLIK_TUFE.get(y, VARSAYILAN_TUFE) / 100
    return c

def _fcf_ortalama(cf, yil=3, kur_s=None):
    """Son 'yil' yılın serbest nakit akışı ortalaması (faaliyet nakdi - yatırım), bugünkü TL ile.
    TL raporlayanda eski yıllar TÜFE ile bugüne taşınır; döviz raporlayanda hepsi son kurla çevrilir."""
    if cf is None or cf.empty:
        return None, 0, False
    degerler = []; son_deger = None
    son_yil = pd.Timestamp(cf.columns[0]).year
    for i in range(min(yil, len(cf.columns))):
        op  = _ilk_deger(cf, ['Operating Cash Flow', 'Cash From Operations'], i)
        cap = _ilk_deger(cf, ['Capital Expenditure', 'Purchase Of PPE', 'Capital Expenditures'], i)
        if op is not None and cap is not None:
            if kur_s is not None:
                carpan = kur(kur_s)
            else:
                carpan = _enf_carpani(pd.Timestamp(cf.columns[i]).year, son_yil)
            degerler.append((op - abs(cap)) * carpan)
            if i == 0:
                son_deger = degerler[-1]
    if not degerler:
        return None, 0, False
    ort = sum(degerler) / len(degerler)
    if son_deger is not None and 0 < son_deger < ort:
        return son_deger, len(degerler), True     # düşüş var: eski seviyeye dönüş varsayma
    return ort, len(degerler), False

def _ciro_buyume(inc, kur_s=None):
    """Ciro (hasılat) yıllık bileşik büyümesi (TL). Net kârdan daha istikrarlı."""
    if inc is None or inc.empty:
        return None
    for k in ['Total Revenue', 'Operating Revenue']:
        if k in inc.index:
            ks = [float(inc.loc[k, c]) * kur(kur_s, c) for c in inc.columns[:4]
                  if pd.notna(inc.loc[k, c]) and float(inc.loc[k, c]) > 0]
            if len(ks) >= 2:
                return (ks[0] / ks[-1]) ** (1 / (len(ks) - 1)) - 1
    return None

def _hisse_degeri(fcf0, g1, g_uzun, wacc, net_borc, hisse):
    """İki aşamalı DCF → hisse başı değer."""
    if wacc - g_uzun < 0.05:          # uzun vade büyümesi WACC'a çok yakınsa formül patlar
        return None
    nakit = fcf0
    bd = 0.0
    for t in range(1, DCF_YIL + 1):
        nakit *= (1 + g1)
        bd += nakit / (1 + wacc) ** t
    son_deger = nakit * (1 + g_uzun) / (wacc - g_uzun)
    bd += son_deger / (1 + wacc) ** DCF_YIL
    ozsermaye = bd - net_borc
    return ozsermaye / hisse

def _ima_edilen_buyume(fcf0, g_uzun, wacc, net_borc, hisse, fiyat):
    """Tersten DCF: bugünkü fiyatı haklı çıkaran 5 yıllık büyüme oranı."""
    alt, ust = -0.30, 1.50
    f_alt = _hisse_degeri(fcf0, alt, g_uzun, wacc, net_borc, hisse)
    f_ust = _hisse_degeri(fcf0, ust, g_uzun, wacc, net_borc, hisse)
    if f_alt is None or f_ust is None:
        return None
    if f_alt >= fiyat: return -30.0     # fiyat çok düşük: küçülme bile haklı çıkarıyor
    if f_ust <= fiyat: return 150.0     # fiyat çok yüksek: %150 büyüme bile yetmiyor
    for _ in range(60):
        orta = (alt + ust) / 2
        if _hisse_degeri(fcf0, orta, g_uzun, wacc, net_borc, hisse) < fiyat:
            alt = orta
        else:
            ust = orta
    return round(orta * 100, 1)

def _sektor_turu(sektor_isy, sektor_yf):
    """'finans', 'gyo', 'holding' veya '' döndürür. Önce İsyatırım sektörüne bakar."""
    if sektor_isy:
        s = str(sektor_isy).lower()
        if 'gyo' in s or 'gayrimenkul' in s:
            return 'gyo'
        if any(k in s for k in ['banka', 'sigorta', 'kiralama', 'faktoring',
                                 'aracı kurum', 'menkul', 'finansman', 'emeklilik']):
            return 'finans'
        if 'holding' in s:
            return 'holding'
        return ''
    # İsyatırım sektörü yoksa yfinance'a bak
    return 'finans' if sektor_yf == 'Financial Services' else ''

def wacc_dcf_hesapla(kod, sektor_isy=None, pd_isy_mn=None):
    try:
        t    = yf.Ticker(f'{kod}.IS')
        info = t.info
        inc  = t.income_stmt
        bal  = t.balance_sheet
        cf   = t.cashflow
        if inc is None or inc.empty or bal is None or bal.empty:
            return None

        res = {}

        # Banka/sigorta ve GYO'da nakit akışı yöntemi anlamsız
        tur = _sektor_turu(sektor_isy, str(info.get('sector', '') or ''))
        if tur == 'finans':
            res['DCF Sinyal'] = '⚪ Finans şirketi (DCF yok)'
            res['DCF Skoru']  = 0
            res['DCF Not']    = 'Banka/sigortada nakit akışı yöntemi uygulanmaz; F/K ve PD/DD kullan'
            return res
        if tur == 'gyo':
            res['DCF Sinyal'] = '⚪ GYO (DCF yok)'
            res['DCF Skoru']  = 0
            res['DCF Not']    = 'GYO değeri portföy değerine göre ölçülür; PD/DD kullan'
            return res

        # ── Özsermaye maliyeti (CAPM) ──
        beta = info.get('beta', 1.0) or 1.0
        beta = max(0.1, min(float(beta), 3.0))
        res['Beta'] = round(beta, 2)
        ke = RF + beta * PIYASA_RISKI
        res['Ke (%)'] = round(ke * 100, 2)

        # ── Raporlama para birimi → TL çevirisi ──
        rapor_para = (info.get('financialCurrency') or 'TRY').upper()
        kur_s = kur_serisi(rapor_para)
        kur_son = kur(kur_s)
        if rapor_para != 'TRY':
            res['Rapor Para Birimi'] = rapor_para

        # ── Borç ve nakit (sadece finansal borç, ticari borç değil), TL ──
        borc  = (_ilk_deger(bal, ['Total Debt']) or 0.0) * kur_son
        nakit = (_ilk_deger(bal, ['Cash Cash Equivalents And Short Term Investments',
                                  'Cash And Cash Equivalents']) or 0.0) * kur_son
        net_borc = borc - nakit
        res['Net Borç (mn)'] = round(net_borc / 1_000_000, 1)

        # ── Borç maliyeti ──
        faiz = _ilk_deger(inc, ['Interest Expense', 'Interest Expense Non Operating'])
        if faiz: faiz *= kur_son
        if faiz and borc > 0:
            kd = min(max(abs(faiz) / borc, RF * 0.8), 0.60)
        else:
            kd = RF * 1.2
        res['Kd (%)'] = round(kd * 100, 2)

        # ── Fiyat, hisse sayısı ──
        shares = info.get('sharesOutstanding', 0)
        fiyat  = info.get('currentPrice') or info.get('regularMarketPrice')
        if fiyat is None:
            try:
                h = t.history(period='5d')['Close'].dropna()
                fiyat = h.iloc[-1] if len(h) > 0 else None
            except: pass

        # ── Hisse sayısı kontrolü: İsyatırım piyasa değeriyle ──
        hisse_duzeltildi = False
        try:
            if pd_isy_mn and fiyat and fiyat > 0:
                isy_hisse = float(pd_isy_mn) * 1_000_000 / fiyat
                if not shares:
                    shares = isy_hisse; hisse_duzeltildi = True
                else:
                    oran = isy_hisse / shares
                    if 0.05 <= oran <= 20 and abs(oran - 1) > 0.25:
                        shares = isy_hisse; hisse_duzeltildi = True
        except: pass

        # ── WACC ağırlıkları (piyasa değeri bazlı) ──
        pd_val = (fiyat * shares) if fiyat and shares else None
        oz = _ilk_deger(bal, ['Common Stock Equity'])
        if oz: oz *= kur_son
        if pd_val and borc > 0:
            tv = pd_val + borc; ev = pd_val / tv; dv = borc / tv
        elif oz and borc > 0:
            tv = abs(oz) + borc; ev = abs(oz) / tv; dv = borc / tv
        elif borc == 0:
            ev = 1.0; dv = 0.0
        else:
            ev = 0.6; dv = 0.4

        wacc = ke * ev + kd * (1 - VERGI_ORANI) * dv
        res['WACC (%)'] = round(wacc * 100, 2)
        res['E/V (%)']  = round(ev * 100, 1)
        res['D/V (%)']  = round(dv * 100, 1)

        # ── Serbest nakit akışı (3 yıl ortalama) ──
        fcf0, fcf_yil, fcf_son_yil = _fcf_ortalama(cf, 3, kur_s)
        if fcf0 is not None:
            res['FCF Ort. (mn)'] = round(fcf0 / 1_000_000, 1)

        # ── Büyüme: ciro büyümesinden, sınırlı ──
        g_ciro = _ciro_buyume(inc, kur_s)
        g1 = g_ciro if g_ciro is not None else UZUN_VADE_BUYUME
        g1 = max(BUYUME_TABAN, min(g1, BUYUME_TAVAN))
        res['Büyüme (%)'] = round(g1 * 100, 1)     # eski sütun adı korunuyor
        res['Uzun Vade Büyüme (%)'] = round(UZUN_VADE_BUYUME * 100, 1)

        if fcf0 is None or not shares or shares <= 0 or not fiyat or fiyat <= 0:
            res['DCF Sinyal'] = '⚪ FCF Yok'; res['DCF Skoru'] = 0
            return res
        if fcf0 <= 0:
            res['DCF Sinyal'] = '⚪ Nakit yakıyor (FCF negatif)'; res['DCF Skoru'] = 0
            res['DCF Not'] = f'Son {fcf_yil} yıl ortalama serbest nakit negatif'
            return res

        # ── 3 senaryo ──
        kotu = _hisse_degeri(fcf0, g1 * 0.5,              UZUN_VADE_BUYUME - 0.03, wacc + 0.03, net_borc, shares)
        orta = _hisse_degeri(fcf0, g1,                    UZUN_VADE_BUYUME,        wacc,        net_borc, shares)
        iyi  = _hisse_degeri(fcf0, min(g1 * 1.25, 0.60),  UZUN_VADE_BUYUME + 0.02, wacc - 0.02, net_borc, shares)

        if orta is None:
            res['DCF Sinyal'] = '⚠️ WACC<Büyüme'; res['DCF Skoru'] = 0
            return res

        res['DCF İçsel Değer'] = round(max(orta, 0), 2)
        if kotu is not None: res['DCF Kötü'] = round(max(kotu, 0), 2)
        if iyi  is not None: res['DCF İyi']  = round(max(iyi, 0), 2)

        # ── Tersten DCF ──
        res['Fiyatın Beklediği Büyüme (%)'] = _ima_edilen_buyume(
            fcf0, UZUN_VADE_BUYUME, wacc, net_borc, shares, fiyat)

        # ── Sinyal (eski etiketler aynen korunuyor) ──
        if orta <= 0:
            res['DCF İskonto (%)'] = -100.0
            res['DCF Sinyal'] = '🔴 Çok Pahalı'; res['DCF Skoru'] = -8
            res['DCF Not'] = 'Net borç, şirket değerinden büyük'
            return res

        isk = ((orta - fiyat) / fiyat) * 100
        res['DCF İskonto (%)'] = round(isk, 1)
        kotu_bile_ucuz = kotu is not None and kotu >= fiyat   # güvenlik payı

        if isk >= 50 and kotu_bile_ucuz:
            res['DCF Sinyal'] = '🟢 Çok Ucuz';   res['DCF Skoru'] = 15
        elif isk >= 20:
            res['DCF Sinyal'] = '🟢 Ucuz';       res['DCF Skoru'] = 10
        elif isk >= 0:
            res['DCF Sinyal'] = '🟡 Makul';      res['DCF Skoru'] = 5
        elif isk >= -20:
            res['DCF Sinyal'] = '🟠 Pahalı';     res['DCF Skoru'] = -3
        else:
            res['DCF Sinyal'] = '🔴 Çok Pahalı'; res['DCF Skoru'] = -8

        notlar = []
        if fcf_son_yil:      notlar.append('Son yıl FCF ortalamanın altında: son yıl kullanıldı')
        if hisse_duzeltildi: notlar.append('Hisse sayısı İsyatırım piyasa değeriyle düzeltildi')
        if tur == 'holding': notlar.append('Holding: konsolide rakamlar, DCF sınırlı güvenilir')
        if g_ciro is None:   notlar.append('Ciro verisi yok, büyüme varsayıldı')
        if fcf_yil < 3:      notlar.append(f'FCF sadece {fcf_yil} yıllık')
        if notlar:
            res['DCF Not'] = ' | '.join(notlar)
        return res
    except Exception as e:
        return {'DCF Sinyal': '⚪ Hata', 'DCF Skoru': 0, 'DCF Not': str(e)[:60]}

# ══════════════════════════════════════════════════════════════
#  MODÜL 3: İSYATIRIM VERİLERİ
# ══════════════════════════════════════════════════════════════
def isyatirim_veri_cek(kod, ozet_df, yabanci_df, tarihsel_df, finansal_df):
    v = {}
    try:
        if yabanci_df is not None:
            row = yabanci_df[yabanci_df['Kod'].astype(str).str.upper() == kod]
            if len(row) > 0:
                row = row.iloc[0]
                tcols = [c for c in yabanci_df.columns if any(x in str(c) for x in
                         ['Ocak','Şubat','Mart','Nisan','Mayıs','Haziran','Temmuz','Ağustos','Eylül','Ekim','Kasım','Aralık','2027','2026','2025'])]
                if len(tcols) >= 1:
                    son = row.get(tcols[-1])
                    if pd.notna(son): v['Yabancı Oranı (%)'] = round(float(son), 2)
                if len(tcols) >= 2:
                    onc = row.get(tcols[-2])
                    if pd.notna(son) and pd.notna(onc):
                        v['Yabancı Trend'] = round(float(son) - float(onc), 2)
                bp = row.get('Değişim (Baz Puan)')
                if pd.notna(bp): v['Yabancı Baz Puan'] = round(float(bp), 1)
    except: pass

    try:
        if ozet_df is not None:
            row = ozet_df[ozet_df['Kod'].astype(str).str.upper() == kod]
            if len(row) > 0:
                row = row.iloc[0]
                for col, key in [('Sektör','Sektör'), ('Halka AçıklıkOranı (%)','Halka Açıklık (%)'),
                                  ('Piyasa Değeri(mn TL)','Piyasa Değeri (mn TL)')]:
                    val = row.get(col)
                    if pd.notna(val):
                        try: v[key] = round(float(val), 1) if key != 'Sektör' else str(val)
                        except: v[key] = str(val)
    except: pass

    try:
        if tarihsel_df is not None:
            row = tarihsel_df[tarihsel_df['Kod'].astype(str).str.upper() == kod]
            if len(row) > 0:
                row = row.iloc[0]
                for col, key in [('F/K Tahmin','İSY F/K Tahmin'), ('PD/DD Tahmin','İSY PD/DD Tahmin'),
                                  ('Prim/İsk. (%)','F/K Prim/İsk (%)'), ('Prim/İsk. (%).2','PD/DD Prim/İsk (%)')]:
                    val = row.get(col)
                    if pd.notna(val):
                        try: v[key] = round(float(val), 1)
                        except: pass
    except: pass

    try:
        if finansal_df is not None:
            row = finansal_df[finansal_df['Kod'].astype(str).str.upper() == kod]
            if len(row) > 0:
                row = row.iloc[0]
                for col, key in [('F/K','İSY F/K'), ('PD/DD','İSY PD/DD'),
                                  ('FD/FAVÖK','FD/FAVÖK'), ('Son Dönem','Son Dönem')]:
                    val = row.get(col)
                    if pd.notna(val):
                        try: v[key] = round(float(val), 2) if key != 'Son Dönem' else str(val)
                        except: pass
    except: pass

    return v

def isyatirim_skor(veri):
    s = 0; d = []
    yab = veri.get('Yabancı Oranı (%)')
    if yab is not None:
        if yab >= 15:  s += 8;  d.append(f'✅ Yab %{yab}')
        elif yab >= 8: s += 4;  d.append(f'🟡 Yab %{yab}')
        elif yab >= 3: s += 1;  d.append(f'🟠 Yab %{yab}')
        else:          s -= 3;  d.append(f'🔴 Yab %{yab}')

    bp = veri.get('Yabancı Baz Puan')
    if bp is not None:
        if bp >= 50:    s += 8;  d.append(f'✅ +{bp}bp')
        elif bp >= 10:  s += 4;  d.append(f'🟡 +{bp}bp')
        elif bp >= -10: s += 1;  d.append(f'🟡 {bp}bp')
        elif bp >= -50: s -= 4;  d.append(f'🟠 {bp}bp')
        else:           s -= 8;  d.append(f'🔴 {bp}bp kaçış')

    fk = veri.get('F/K Prim/İsk (%)')
    if fk is not None:
        if fk <= -20:   s += 8;  d.append(f'✅ F/K %{fk}isk')
        elif fk <= -10: s += 5;  d.append(f'🟡 F/K %{fk}')
        elif fk <= 0:   s += 2;  d.append(f'🟡 F/K normal')
        elif fk <= 20:  s -= 3;  d.append(f'🟠 F/K %{fk}prim')
        else:           s -= 7;  d.append(f'🔴 F/K %{fk}')

    pd3 = veri.get('PD/DD Prim/İsk (%)')
    if pd3 is not None:
        if pd3 <= -20:  s += 5;  d.append(f'✅ PD/DD %{pd3}isk')
        elif pd3 <= 0:  s += 2;  d.append(f'🟡 PD/DD normal')
        else:           s -= 3;  d.append(f'🟠 PD/DD %{pd3}')

    halka = veri.get('Halka Açıklık (%)')
    if halka is not None:
        if halka >= 30:  s += 4;  d.append(f'✅ Halka %{halka}')
        elif halka >= 15: s += 2; d.append(f'🟡 Halka %{halka}')
        elif halka >= 5:  s -= 2; d.append(f'🟠 Halka %{halka}')
        else:             s -= 6; d.append(f'🔴 Halka %{halka}')

    fd = veri.get('FD/FAVÖK')
    if fd is not None:
        try:
            fd = float(fd)
            if fd <= 5:    s += 5;  d.append(f'✅ FD/F {fd}')
            elif fd <= 10: s += 3;  d.append(f'🟡 FD/F {fd}')
            elif fd <= 15: s += 0;  d.append(f'🟠 FD/F {fd}')
            else:          s -= 4;  d.append(f'🔴 FD/F {fd}')
        except: pass

    return s, ' | '.join(d) if d else 'ISY verisi yok'

# ══════════════════════════════════════════════════════════════
#  MODÜL 4: SERMAYE OLAYLARI
# ══════════════════════════════════════════════════════════════
def sermaye_kontrol(kod, sermaye_df):
    if sermaye_df is None:
        return '-', 0
    try:
        bugun = pd.Timestamp.now()
        rows  = sermaye_df[sermaye_df['Kod'].astype(str).str.upper() == kod]
        if len(rows) == 0:
            return '-', 0

        gelecek = rows[rows['Tarih'] >= bugun].sort_values('Tarih')
        if len(gelecek) == 0:
            gecmis = rows[rows['Tarih'] < bugun].sort_values('Tarih', ascending=False)
            if len(gecmis) > 0 and (bugun - gecmis.iloc[0]['Tarih']).days <= 30:
                return f"🕐 {(bugun - gecmis.iloc[0]['Tarih']).days}g önce", -2
            return '-', 0

        r = gelecek.iloc[0]
        kalan = (r['Tarih'] - bugun).days

        try:
            bik  = float(r.get('BedelsizIKOran(%)', 0) or 0)
            btem = float(r.get('BedelsizTemettüOran(%)', 0) or 0)
            ntem = float(r.get('NakitTem.Brüt(%)', 0) or 0)
            bed  = float(r.get('BedelliOran(%)', 0) or 0)
        except:
            bik = btem = ntem = bed = 0

        parcalar = []
        if bik + btem > 0:  parcalar.append(f'🎁 %{bik+btem:.0f} Bedelsiz')
        if ntem > 0:        parcalar.append(f'💰 %{ntem:.0f} Tem.')
        if bed > 0:         parcalar.append(f'📋 %{bed:.0f} Bedelli')

        if not parcalar:
            return '-', 0

        tarih_str = r['Tarih'].strftime('%d.%m.%Y')
        txt = ' | '.join(parcalar)
        if kalan <= 7:    return f'🔥 {txt} — {kalan}g! ({tarih_str})', 15
        elif kalan <= 30: return f'⚡ {txt} — {kalan}g ({tarih_str})', 10
        elif kalan <= 90: return f'📅 {txt} — {kalan}g ({tarih_str})', 5
        else:             return f'🗓️ {txt} — {tarih_str}', 2
    except:
        return '-', 0

# ══════════════════════════════════════════════════════════════
#  MODÜL 5: VOLATİLİTE BAZLI STOP/HEDEF/POZİSYON
#  Geçmiş fiyat hareketinden stop ve pozisyon büyüklüğü hesaplar
# ══════════════════════════════════════════════════════════════
def volatilite_analiz(kod):
    try:
        t    = yf.Ticker(f'{kod}.IS')
        hist = t.history(period='1y', interval='1d')
        if hist is None or len(hist) < 61:
            return None
        hist = tamamlanmis_gunler(hist)

        close = hist['Close']
        fiyat = close.iloc[-1]

        # Günlük volatilite
        gunluk_getiri = close.pct_change().dropna()
        vol_gunluk    = gunluk_getiri.std()
        vol_yillik    = vol_gunluk * np.sqrt(252)

        # ATR (Average True Range) — 14 günlük
        high  = hist['High']
        low   = hist['Low']
        tr    = pd.concat([high - low,
                           (high - close.shift()).abs(),
                           (low  - close.shift()).abs()], axis=1).max(axis=1)
        atr   = tr.rolling(14).mean().iloc[-1]
        atr_yuzde = (atr / fiyat) * 100

        # Tipik düşüş ve toparlanma analizi (son 1 yıl)
        pencere = 20  # 20 günlük pencereler
        max_dususler = []
        for i in range(0, len(close) - pencere, 5):
            pencere_veri = close.iloc[i:i+pencere]
            tepe = pencere_veri.max()
            dip  = pencere_veri.min()
            if tepe > 0:
                max_dususler.append(((dip - tepe) / tepe) * 100)

        ort_dusus = np.median(max_dususler) if max_dususler else -10

        # Momentum (son 20 gün vs 60 gün)
        momentum_20 = ((fiyat / close.iloc[-20]) - 1) * 100 if len(close) >= 20 else 0
        momentum_60 = ((fiyat / close.iloc[-60]) - 1) * 100 if len(close) >= 60 else 0

        # ATR bazlı stop seviyeleri
        stop_dar    = round(fiyat - 1.5 * atr, 2)   # Dar stop: 1.5x ATR
        stop_normal = round(fiyat - 2.5 * atr, 2)   # Normal: 2.5x ATR
        stop_genis  = round(fiyat - 4.0 * atr, 2)   # Geniş: 4x ATR

        # Hedef seviyeleri (ATR bazlı)
        hedef_1 = round(fiyat + 2.0 * atr, 2)   # Kısa hedef
        hedef_2 = round(fiyat + 4.0 * atr, 2)   # Orta hedef
        hedef_3 = round(fiyat + 7.0 * atr, 2)   # Uzun hedef

        # Risk/Ödül ana döngüde sistem hedefiyle hesaplanıyor (ATR'yle hesaplanınca hep 1,6 çıkıyordu)
        risk_odul = None

        return {
            'ATR':           round(atr, 2),
            'ATR (%)':       round(atr_yuzde, 2),
            'Vol Günlük (%)': round(vol_gunluk * 100, 2),
            'Vol Yıllık (%)': round(vol_yillik * 100, 1),
            'Ort Düşüş (%)': round(ort_dusus, 1),
            'Momentum 20G':  round(momentum_20, 1),
            'Momentum 60G':  round(momentum_60, 1),
            'Stop Dar':      stop_dar,
            'Stop Normal':   stop_normal,
            'Stop Geniş':    stop_genis,
            'Hedef 1':       hedef_1,
            'Hedef 2':       hedef_2,
            'Hedef 3':       hedef_3,
            'Risk/Ödül':     risk_odul,
            'Son Fiyat':     round(fiyat, 2),
        }
    except:
        return None

def pozisyon_buyuklugu(sermaye, vol_veri, risk_yuzdesi=0.02):
    """
    Volatilite bazlı pozisyon büyüklüğü
    risk_yuzdesi: toplam sermayenin kaçı riske girecek (varsayılan %2)
    """
    if vol_veri is None:
        return None, None

    fiyat       = vol_veri['Son Fiyat']
    stop        = vol_veri['Stop Normal']
    stop_mesafe = fiyat - stop

    if stop_mesafe <= 0 or not np.isfinite(stop_mesafe):
        return None, None

    risk_tl    = sermaye * risk_yuzdesi
    adet       = int(risk_tl / stop_mesafe)
    maliyet_tl = round(adet * fiyat, 0)

    return adet, maliyet_tl

# ══════════════════════════════════════════════════════════════
#  MODÜL 6: GİRİŞ ZAMANLAMASI
#  Temel + Teknik + Makro üçü uyumluysa GİR
# ══════════════════════════════════════════════════════════════
def giris_zamanlama(row_sistem, wacc_veri, isy_veri, vol_veri, makro_skor=None):
    """
    Giriş zamanlaması analizi
    Temel, teknik ve makro sinyalleri birleştirir
    """
    sinyaller = []
    puan = 0

    # TEMEL ANALİZ SİNYALLERİ
    toplam_skor = float(row_sistem.get('Toplam Skor', 0) or 0)
    if toplam_skor >= 80:   puan += 3; sinyaller.append('✅ Temel güçlü')
    elif toplam_skor >= 55: puan += 2; sinyaller.append('🟡 Temel iyi')
    elif toplam_skor >= 35: puan += 1; sinyaller.append('🟠 Temel orta')
    else:                   puan -= 2; sinyaller.append('🔴 Temel zayıf')

    # DCF SİNYALİ
    dcf = str(wacc_veri.get('DCF Sinyal', '') if wacc_veri else '')
    if '🟢 Çok Ucuz' in dcf:   puan += 3; sinyaller.append('✅ DCF çok ucuz')
    elif '🟢 Ucuz' in dcf:     puan += 2; sinyaller.append('🟡 DCF ucuz')
    elif '🟡 Makul' in dcf:    puan += 1; sinyaller.append('🟡 DCF makul')
    elif '🔴' in dcf:          puan -= 2; sinyaller.append('🔴 DCF pahalı')

    # İSYATİRİM SİNYALİ
    fk_prim = isy_veri.get('F/K Prim/İsk (%)')
    if fk_prim is not None:
        if fk_prim <= -20:   puan += 2; sinyaller.append('✅ İSY ucuz')
        elif fk_prim <= 0:   puan += 1; sinyaller.append('🟡 İSY normal')
        elif fk_prim > 50:   puan -= 2; sinyaller.append('🔴 İSY pahalı')

    # VOLATİLİTE / TEKNİK SİNYAL
    if vol_veri:
        mom20 = vol_veri.get('Momentum 20G', 0)
        mom60 = vol_veri.get('Momentum 60G', 0)
        risko = vol_veri.get('Risk/Ödül')

        if mom20 > 0 and mom60 > 0:   puan += 2; sinyaller.append('✅ Momentum pozitif')
        elif mom20 > 0:               puan += 1; sinyaller.append('🟡 Kısa vade pozitif')
        elif mom20 < -10:             puan -= 2; sinyaller.append('🔴 Momentum negatif')

        if risko is None:  sinyaller.append('⚪ Risk/Ödül yok (hedef yok)')
        elif risko >= 5:   puan += 2; sinyaller.append(f'✅ Risk/Ödül {risko}')
        elif risko >= 2:   puan += 1; sinyaller.append(f'🟡 Risk/Ödül {risko}')
        elif risko >= 1:   sinyaller.append(f'🟠 Risk/Ödül {risko}')
        else:              puan -= 1; sinyaller.append(f'🔴 Risk/Ödül {risko}')

    # MAKRO SİNYAL
    if makro_skor is not None:
        if makro_skor >= 20:    puan += 3; sinyaller.append('✅ Makro uygun')
        elif makro_skor >= 0:   puan += 1; sinyaller.append('🟡 Makro nötr')
        elif makro_skor >= -20: puan -= 1; sinyaller.append('🟠 Makro riskli')
        else:                   puan -= 3; sinyaller.append('🔴 Makro kötü')

    # MAKRO BAZLI KARAR
    # Makro < -20: çok kötü, sadece en güçlü hisseler HAZIR OL
    # Makro -20 ile -5 arası: riskli ama toparlanıyor, puan 7+ ise ŞİMDİ GİR
    # Makro -5 ile 0 arası: toparlanıyor, puan 6+ ise ŞİMDİ GİR
    # Makro 0+: iyi ortam, normal çalış
    if makro_skor is not None and makro_skor < -20:
        if puan >= 10:   karar = '🟡 HAZIR OL';   renk = 'FFB300'
        elif puan >= 5:  karar = '⚪ İZLE';       renk = 'AAAAAA'
        else:            karar = '🔴 UZAK DUR';   renk = 'FF4757'
    elif makro_skor is not None and makro_skor < -5:
        # Riskli ama toparlanıyor
        if puan >= 7:    karar = '🟢 ŞİMDİ GİR'; renk = '00FF88'
        elif puan >= 5:  karar = '🟡 HAZIR OL';   renk = 'FFB300'
        elif puan >= 3:  karar = '🟠 BEKLE';      renk = 'FF8C00'
        else:            karar = '⚪ İZLE';       renk = 'AAAAAA'
    elif makro_skor is not None and makro_skor < 0:
        # Toparlanıyor, yakında pozitife geçecek
        if puan >= 6:    karar = '🟢 ŞİMDİ GİR'; renk = '00FF88'
        elif puan >= 4:  karar = '🟡 HAZIR OL';   renk = 'FFB300'
        elif puan >= 2:  karar = '🟠 BEKLE';      renk = 'FF8C00'
        else:            karar = '⚪ İZLE';       renk = 'AAAAAA'
    else:
        # Makro iyi, normal çalış
        if puan >= 8:    karar = '🟢 ŞİMDİ GİR'; renk = '00FF88'
        elif puan >= 5:  karar = '🟡 HAZIR OL';   renk = 'FFB300'
        elif puan >= 3:  karar = '🟠 BEKLE';      renk = 'FF8C00'
        elif puan >= 0:  karar = '⚪ İZLE';       renk = 'AAAAAA'
        else:            karar = '🔴 UZAK DUR';   renk = 'FF4757'

    return puan, karar, renk, ' | '.join(sinyaller)

# ══════════════════════════════════════════════════════════════
#  MODÜL 7: ÇIKIŞ STRATEJİSİ (GELİŞTİRİLMİŞ)
#  Backtest bulgularını baz alır:
#  - 60 gün ortalama +%13 getiri
#  - Stop -%15, Hedef +%30 optimal
#  - Trailing stop ile karı koru
# ══════════════════════════════════════════════════════════════
def cikis_stratejisi(alis_fiyat, son_fiyat, stop, hedef_ort, vol_veri, alis_tarih=None):
    if alis_fiyat is None or son_fiyat is None:
        return '-', 'AAAAAA', '-', '-', '-'

    kz_yuz = ((son_fiyat - alis_fiyat) / alis_fiyat) * 100

    # ATR bazlı hedefler
    atr_h1 = vol_veri.get('Hedef 1', 0) if vol_veri else 0
    atr_h2 = vol_veri.get('Hedef 2', 0) if vol_veri else 0
    atr    = vol_veri.get('ATR', 0) if vol_veri else 0

    # Trailing stop hesapla (backtest: -%15 optimal)
    # Kar arttıkça stop'u yukarı çek
    if kz_yuz >= 30:
        trailing_stop = round(son_fiyat * 0.92, 2)   # %8 altı
        trailing_not  = f'Trailing stop: {trailing_stop:.2f}₺ (%8 altı)'
    elif kz_yuz >= 20:
        trailing_stop = round(alis_fiyat * 1.10, 2)  # alışın %10 üstü
        trailing_not  = f'Trailing stop: {trailing_stop:.2f}₺ (alış+%10)'
    elif kz_yuz >= 10:
        trailing_stop = round(alis_fiyat * 1.03, 2)  # alışın %3 üstü
        trailing_not  = f'Trailing stop: {trailing_stop:.2f}₺ (alış+%3)'
    else:
        trailing_stop = stop
        trailing_not  = f'Stop: {stop:.2f}₺'

    # Hedef mesafeleri
    if hedef_ort and hedef_ort > 0:
        hedefe_uzak = ((hedef_ort - son_fiyat) / hedef_ort) * 100
    else:
        hedefe_uzak = None

    # 0. BİLEŞİK DÖNDÜRME — 45 gün içinde %12 hedefi
    if alis_tarih:
        try:
            gun = (datetime.now() - datetime.strptime(alis_tarih, '%Y-%m-%d')).days
            if kz_yuz >= BILESIK_HEDEF_YUZ and gun <= BILESIK_SURE_GUN:
                mesaj = f'🎯 %{BILESIK_HEDEF_YUZ} HEDEFİ — SAT & DÖNDÜR! ({gun}. gün, %{kz_yuz:.1f} kar)'
                return mesaj, '00FF88', '🎯 SAT & DÖNDÜR', f'Kar: %{kz_yuz:+.1f} | {gun}. gün | Bileşik döndür!', trailing_not
            elif gun > BILESIK_SURE_GUN and kz_yuz < BILESIK_HEDEF_YUZ:
                mesaj = f'⏰ {BILESIK_SURE_GUN} GÜN DOLDU — Hedef gelmedi! ({gun}. gün, %{kz_yuz:+.1f})'
                return mesaj, 'FF8C00', '⏰ DÖNDÜR', f'{gun}. gün | Kar: %{kz_yuz:+.1f} | Alternatif ara', trailing_not
        except: pass

        # 1. STOP KIRILDI
    if son_fiyat <= stop:
        mesaj  = f'🛑 STOP KIRILDI — HEMEN SAT! ({son_fiyat:.2f}₺ <= {stop:.2f}₺)'
        islem  = '🛑 HEMEN SAT'
        detay  = f'Zarar: %{kz_yuz:.1f} | Stop kırıldı'
        return mesaj, 'FF4757', islem, detay, trailing_not

    # 2. TRAILING STOP KIRILDI
    if kz_yuz >= 10 and son_fiyat <= trailing_stop:
        mesaj  = f'🛑 TRAİLİNG STOP — Karı Koru! ({son_fiyat:.2f}₺ <= {trailing_stop:.2f}₺)'
        islem  = '🛑 KARI KOR'
        detay  = f'Kar: %{kz_yuz:+.1f} | Trailing stop kırıldı'
        return mesaj, 'FF8C00', islem, detay, trailing_not

    # 3. TEMEL HEDEF AŞILDI
    if hedef_ort and hedef_ort > 0 and son_fiyat >= hedef_ort * 0.95:
        mesaj  = f'🎯 HEDEFE ULAŞTI — TAM ÇIK! Hedef:{hedef_ort:.2f}₺ | Şu an:{son_fiyat:.2f}₺'
        islem  = '🎯 TAM ÇIK'
        detay  = f'Kar: %{kz_yuz:+.1f} | Hedef aşıldı'
        return mesaj, '00FF88', islem, detay, trailing_not

    # 4. ATR HEDEF 2 AŞILDI — Kısmi çıkış
    if atr_h2 and son_fiyat >= atr_h2:
        mesaj  = f'🎯 ATR H2 AŞILDI — YARIYI SAT! H2:{atr_h2:.2f}₺ | Şu an:{son_fiyat:.2f}₺'
        islem  = '⚡ YARIYI SAT'
        detay  = f'Kar: %{kz_yuz:+.1f} | {trailing_not}'
        return mesaj, '00FF88', islem, detay, trailing_not

    # 5. ATR HEDEF 1 AŞILDI — Hazırlan
    if atr_h1 and son_fiyat >= atr_h1:
        mesaj  = f'⚡ ATR H1 AŞILDI — HAZIR OL! H1:{atr_h1:.2f}₺ | H2 hedef:{atr_h2:.2f}₺'
        islem  = '⚡ HAZIRLAN'
        detay  = f'Kar: %{kz_yuz:+.1f} | {trailing_not}'
        return mesaj, 'FFB300', islem, detay, trailing_not

    # 6. HEDEFE YAKLAŞTI
    if hedefe_uzak is not None and hedefe_uzak <= 10:
        mesaj  = f'⚡ Hedefe %{hedefe_uzak:.0f} kaldı — Hazır ol! Hedef:{hedef_ort:.2f}₺'
        islem  = '⚡ HAZIRLAN'
        detay  = f'Kar: %{kz_yuz:+.1f} | {trailing_not}'
        return mesaj, 'FFB300', islem, detay, trailing_not

    # 7. GÜÇLÜ KAR (+%20 üstü) — Trailing stop uygula
    if kz_yuz >= 20:
        mesaj  = f'✅ +%{kz_yuz:.1f} KAR — {trailing_not}'
        islem  = '✅ TUT & KORUMA'
        detay  = f'Trailing stop aktif | H2:{atr_h2:.2f}₺ hedef' if atr_h2 else f'Trailing stop aktif'
        return mesaj, '00FF88', islem, detay, trailing_not

    # 8. İYİ KAR (+%10 üstü)
    if kz_yuz >= 10:
        mesaj  = f'🟢 +%{kz_yuz:.1f} karda — {trailing_not}'
        islem  = '🟢 TUT'
        detay  = f'H1:{atr_h1:.2f}₺ | H2:{atr_h2:.2f}₺' if atr_h1 else 'Bekle'
        return mesaj, '00CC66', islem, detay, trailing_not

    # 9. ZARAR — Stop yakın mı?
    if kz_yuz <= -12:
        mesaj  = f'⚠️ -%{abs(kz_yuz):.1f} — STOP GÖZDEN GEÇİR! Stop:{stop:.2f}₺'
        islem  = '⚠️ DİKKAT'
        detay  = f'Stop mesafe: %{((son_fiyat-stop)/son_fiyat*100):.1f}'
        return mesaj, 'FF8C00', islem, detay, trailing_not
    elif kz_yuz <= -7:
        mesaj  = f'🟠 -%{abs(kz_yuz):.1f} — Dikkat et'
        islem  = '🟠 İZLE'
        detay  = f'Stop: {stop:.2f}₺'
        return mesaj, 'FF8C00', islem, detay, trailing_not

    # 10. NORMAL POZİSYON
    if kz_yuz >= 0:
        mesaj  = f'🟡 +%{kz_yuz:.1f} karda — bekle | H1:{atr_h1:.2f}₺ | H2:{atr_h2:.2f}₺' if atr_h1 else f'🟡 +%{kz_yuz:.1f} karda — bekle'
    else:
        mesaj  = f'🟡 -%{abs(kz_yuz):.1f} zararda — bekle | Stop:{stop:.2f}₺'
    islem  = '🟡 BEKLE'
    detay  = f'H1:{atr_h1:.2f}₺ | H2:{atr_h2:.2f}₺ | {trailing_not}' if atr_h1 else trailing_not
    return mesaj, 'FFB300', islem, detay, trailing_not

# ══════════════════════════════════════════════════════════════
#  MODÜL 8: VETO KONTROLÜ
# ══════════════════════════════════════════════════════════════
def veto_kontrol(row, isy_veri, dcf_veri):
    vetolar = []

    # Hedef < fiyat
    try:
        hedef = float(str(row.get('Hedef (Ort.)', 0)).replace('₺', ''))
        fiyat = float(str(row.get('Son Fiyat', 0)).replace('₺', ''))
        if hedef > 0 and fiyat > 0 and hedef < fiyat:
            vetolar.append(f'🚫 Hedef({hedef:.2f}₺) < Fiyat')
    except: pass

    # AI SAT
    haber = str(row.get('Haber Kararı', ''))
    if '🔴 SAT' in haber:
        vetolar.append('🚫 AI: SAT')

    # F/K aşırı primli
    fk = isy_veri.get('F/K Prim/İsk (%)')
    if fk is not None:
        try:
            if float(fk) > 300: vetolar.append(f'🚫 F/K %{fk:.0f} aşırı primli')
        except: pass

    # Güven düşük
    guven = row.get('Güven Skoru', 100)
    try:
        if float(guven) < 40: vetolar.append(f'🚫 Güven {guven:.0f} düşük')
    except: pass

    # Likidite
    if '🔴 Çok Düşük' in str(row.get('Likidite', '')):
        vetolar.append('🚫 Likidite çok düşük')

    # Hukuki risk: yapay zekânın açık 'RİSK: VAR' cevabı (hücre 6 v2.1+)
    hukuki = row.get('Hukuki Risk', None)
    if hukuki is not None and pd.notna(hukuki) and str(hukuki).strip() not in ('', '-'):
        if 'VAR' in str(hukuki).upper():
            vetolar.append('🚫 Hukuki/SPK riski')
    else:
        # eski hücre 6 sürümü: gerekçede kelime ara, ama "içermiyor/yok" geçen cümleleri sayma
        ai_g = str(row.get('AI Gerekçe', '')).lower()
        olumsuz = any(k in ai_g for k in ['içermiyor', 'içermediği', 'yoktur', ' yok', 'bulunmuyor'])
        if not olumsuz and any(k in ai_g for k in ['spk', 'suç duyurusu', 'manipülasyon', 'soruşturma']):
            vetolar.append('🚫 Hukuki/SPK riski')

    # DCF pahalı + hedef altı
    dcf_s = str(dcf_veri.get('DCF Sinyal', '') if dcf_veri else '')
    hedefe = row.get('Hedefe Uzaklık (%)', 0)
    try:
        if '🔴 Çok Pahalı' in dcf_s and float(hedefe) < 0:
            vetolar.append('🚫 DCF pahalı + hedef altı')
    except: pass

    return vetolar

def oneri_uret(row, vetolar, giris_puan):
    veto_n = len(vetolar)
    if veto_n >= 2:   return '🔴 ALMA',      'FF4757', '2D0000'
    if veto_n == 1:   return '🟠 DİKKAT',    'FF8C00', '1A0A00'

    if giris_puan >= 10:  return '🟢 GÜÇLÜ AL', '00FF88', '001A10'
    if giris_puan >= 6:   return '🟢 AL',        '00CC66', '002010'
    if giris_puan >= 3:   return '🟡 İZLE',      'FFB300', '1A1400'
    if giris_puan >= 0:   return '⚪ BEKLE',     'AAAAAA', '0D1B2E'
    return '🔴 UZAK DUR', 'FF4757', '2D0000'

# ══════════════════════════════════════════════════════════════
#  ANA KOD BAŞLANGICI
# ══════════════════════════════════════════════════════════════
print("=" * 65)
print("  🔬 KAPSAMLI ANALİZ MOTORU — HÜCRE 7")
print(f"  {datetime.now().strftime('%d.%m.%Y %H:%M')}")
print("=" * 65)

# Verileri oku
print("\n📂 Veriler yükleniyor...")
df_sistem = pd.read_excel(SISTEM_YOLU)
df_sistem = df_sistem.rename(columns={'Sembol': 'Kod'}) if 'Sembol' in df_sistem.columns else df_sistem

# Teknik analiz sayfasından teknik verileri al
df_teknik = pd.DataFrame()
try:
    for h in range(6):
        df_tmp = pd.read_excel(EXCEL_YOLU, sheet_name='📊 Teknik Analiz', header=h)
        if 'Kod' in df_tmp.columns:
            df_teknik = df_tmp
            break
except: pass

# AI analiz sayfasından haber kararlarını al
df_ai = pd.DataFrame()
try:
    for h in range(6):
        df_tmp = pd.read_excel(EXCEL_YOLU, sheet_name='🤖 AI Analiz', header=h)
        if 'Kod' in df_tmp.columns:
            df_ai = df_tmp
            break
except: pass

# Makro skoru oku
makro_skor_global = None
makro_zaman = None
try:
    wb_tmp = load_workbook(EXCEL_YOLU)
    if '📡 Makro Radar' in wb_tmp.sheetnames:
        ws_tmp = wb_tmp['📡 Makro Radar']
        for row in ws_tmp.iter_rows(min_row=2, max_row=3):
            for cell in row:
                if cell.value and 'SKOR:' in str(cell.value):
                    import re
                    m = re.search(r'SKOR:\s*([+-]?\d+)', str(cell.value))
                    if m: makro_skor_global = int(m.group(1))
        import re
        m_t = re.search(r'(\d{2}\.\d{2}\.\d{4} \d{2}:\d{2})', str(ws_tmp['A1'].value or ''))
        if m_t:
            makro_zaman = datetime.strptime(m_t.group(1), '%d.%m.%Y %H:%M')
    wb_tmp.close()
except: pass

makro_yas_saat = None
if makro_zaman is not None:
    makro_yas_saat = (datetime.now() - makro_zaman).total_seconds() / 3600
    if makro_yas_saat > 3:
        print(f"⚠️ DİKKAT: Makro Radar {makro_yas_saat:.0f} saat önce güncellenmiş!")
        print("   Giriş kararları eski makro skora göre çıkacak. Önce hücre 8'i, sonra bu hücreyi çalıştır.")
elif makro_skor_global is None:
    print("⚠️ Makro Radar bulunamadı. Önce hücre 8'i çalıştır.")

print(f"✅ Sistem: {len(df_sistem)} hisse")
print(f"✅ Teknik: {len(df_teknik)} hisse")
print(f"✅ AI: {len(df_ai)} hisse")
print(f"✅ Makro skor: {makro_skor_global}")

# İsyatirim dosyaları
print("\n📊 İsyatirim dosyaları...")
ozet_df = yabanci_df = tarihsel_df = finansal_df = sermaye_df = performans_df = None

for dosya, var_adi in [
    ('temelozet.xlsx',        'ozet_df'),
    ('temelyabancioran.xlsx', 'yabanci_df'),
    ('temeltarihsel.xlsx',    'tarihsel_df'),
    ('temelfinansal.xlsx',    'finansal_df'),
    ('temelsermaye.xlsx',     'sermaye_df'),
    ('temelperformans.xlsx',  'performans_df'),
]:
    try:
        df_tmp = pd.read_excel(f'{DRIVE}/{dosya}')
        if var_adi == 'sermaye_df':
            df_tmp['Tarih'] = pd.to_datetime(df_tmp['Tarih'], dayfirst=True, errors='coerce')
        exec(f'{var_adi} = df_tmp')
        print(f"  ✅ {dosya} — {len(df_tmp)} satır")
    except Exception as e:
        print(f"  ⚠️ {dosya}: {e}")

# Analiz edilecek hisseler — AL listesi + portföy
al_listesi = df_sistem[df_sistem['Toplam Skor'] >= 55]['Kod'].tolist()
portfoy_kodlar = [p['Kod'] for p in PORTFOY]
kodlar = list(dict.fromkeys(portfoy_kodlar + al_listesi))
print(f"\n🎯 {len(kodlar)} hisse analiz edilecek ({len(portfoy_kodlar)} portföy + {len(al_listesi)} AL listesi)")
print(f"⏱️  Tahmini süre: ~{len(kodlar) * 3 // 60 + 2} dakika\n")

# ══════════════════════════════════════════════════════════════
#  ANA ANALİZ DÖNGÜSÜ
# ══════════════════════════════════════════════════════════════
sonuclar = []

for i, kod in enumerate(kodlar, 1):
    portfoy_mi = kod in portfoy_kodlar
    print(f"[{i:3d}/{len(kodlar)}] {'📌 ' if portfoy_mi else ''}{kod}...", end=' ')

    # Sistem verisini al
    row_s  = df_sistem[df_sistem['Kod'] == kod].iloc[0] if len(df_sistem[df_sistem['Kod'] == kod]) > 0 else pd.Series()
    row_t  = df_teknik[df_teknik['Kod'] == kod].iloc[0] if len(df_teknik) > 0 and 'Kod' in df_teknik.columns and len(df_teknik[df_teknik['Kod'] == kod]) > 0 else pd.Series()
    row_ai = df_ai[df_ai['Kod'] == kod].iloc[0] if len(df_ai) > 0 and 'Kod' in df_ai.columns and len(df_ai[df_ai['Kod'] == kod]) > 0 else pd.Series()

    toplam_skor   = float(row_s.get('Toplam Skor', 0) or 0)
    teknik_skor   = float(row_t.get('Teknik Skor', 0) or 0) if len(row_t) > 0 else 0
    haber_skor    = float(row_ai.get('Haber Skoru', 0) or 0) if len(row_ai) > 0 else 0
    sistem_hedef  = float(row_s.get('Hedef (Ort.)', 0) or 0) if len(row_s) > 0 else 0
    son_fiyat_raw = float(row_s.get('Son Fiyat', 0) or 0) if len(row_s) > 0 else 0

    # 1. Güven skoru
    guven_puan, guven_kat, guven_detay, guven_uyari = guven_skoru_hesapla(kod)

    # 2. İsyatirim (sektör bilgisi DCF'de kullanılıyor, o yüzden önce)
    isy_veri = isyatirim_veri_cek(kod, ozet_df, yabanci_df, tarihsel_df, finansal_df)
    isy_skor, isy_detay = isyatirim_skor(isy_veri)

    # 3. WACC + DCF
    wacc_veri = wacc_dcf_hesapla(kod, isy_veri.get('Sektör'), isy_veri.get('Piyasa Değeri (mn TL)'))
    dcf_skor  = wacc_veri.get('DCF Skoru', 0) if wacc_veri else 0

    # 4. Sermaye olayı
    sermaye_aciklama, sermaye_skor = sermaye_kontrol(kod, sermaye_df)

    # 5. Volatilite analizi
    vol_veri = volatilite_analiz(kod)
    # 5b. Risk/Ödül: hedefe kadar kazanç ÷ normal stopa kadar kayıp
    if vol_veri:
        _f, _st = vol_veri.get('Son Fiyat', 0), vol_veri.get('Stop Normal', 0)
        if sistem_hedef > 0 and _f > _st > 0:
            vol_veri['Risk/Ödül'] = round((sistem_hedef - _f) / (_f - _st), 1)

    # 6. Güven düzeltmesi
    temel_toplam = toplam_skor + teknik_skor + haber_skor
    temel_duzelt, duzelt_acikla = guven_duzelt(temel_toplam, guven_puan)

    # 7. Süper final skor
    super_final = temel_duzelt + isy_skor + dcf_skor + sermaye_skor

    # 8. Giriş zamanlaması
    giris_puan, giris_karar, giris_renk, giris_detay = giris_zamanlama(
        row_s, wacc_veri, isy_veri, vol_veri, makro_skor_global
    )

    # 9. Veto kontrolü
    row_combined = {
        'Hedef (Ort.)':       sistem_hedef,
        'Son Fiyat':          son_fiyat_raw,
        'Haber Kararı':       row_ai.get('Haber Kararı', '') if len(row_ai) > 0 else '',
        'AI Gerekçe':         row_ai.get('AI Gerekçe', '') if len(row_ai) > 0 else '',
        'Hukuki Risk':        row_ai.get('Hukuki Risk', None) if len(row_ai) > 0 else None,
        'Güven Skoru':        guven_puan,
        'Likidite':           row_s.get('Likidite', '') if len(row_s) > 0 else '',
        'Hedefe Uzaklık (%)': row_s.get('Hedefe Uzaklık (%)', 0) if len(row_s) > 0 else 0,
    }
    vetolar = veto_kontrol(row_combined, isy_veri, wacc_veri)
    oneri, oneri_renk, oneri_bg = oneri_uret(row_combined, vetolar, giris_puan)

    # 10. Yabancı sinyali
    yab = isy_veri.get('Yabancı Oranı (%)')
    yab_bp = isy_veri.get('Yabancı Baz Puan')
    if yab_bp is not None and yab_bp >= 50:    yab_sinyal = f'🟢 Giriş +{yab_bp}bp'
    elif yab_bp is not None and yab_bp <= -50: yab_sinyal = f'🔴 Kaçış {yab_bp}bp'
    elif yab is not None:                      yab_sinyal = f'🟡 %{yab} sabit'
    else:                                       yab_sinyal = 'Veri yok'

    # 11. Pozisyon büyüklüğü (SERMAYE en üstte ayarlanır)
    poz_adet, poz_maliyet = pozisyon_buyuklugu(SERMAYE, vol_veri)

    sonuclar.append({
        # Kimlik
        'Kod':                kod,
        'Portföyde':          '📌 Evet' if portfoy_mi else '-',
        # Ana kararlar
        'Öneri':              oneri,
        'Giriş Kararı':       giris_karar,
        'Giriş Puanı':        giris_puan,
        'Veto Sayısı':        len(vetolar),
        'Veto Detayı':        ' | '.join(vetolar) if vetolar else '✅ Temiz',
        # Fiyat
        'Son Fiyat':          son_fiyat_raw,
        'Hedef (Ort.)':       sistem_hedef,
        'Hedefe Uzaklık (%)': float(row_s.get('Hedefe Uzaklık (%)', 0) or 0) if len(row_s) > 0 else 0,
        # Skorlar
        'Toplam Skor':        toplam_skor,
        'Teknik Skor':        teknik_skor,
        'Haber Skoru':        haber_skor,
        'Güven Skoru':        guven_puan,
        'Güven Kategorisi':   guven_kat,
        'İSY Skoru':          isy_skor,
        'DCF Skoru':          dcf_skor,
        'Sermaye Skoru':      sermaye_skor,
        'Süper Final Skor':   super_final,
        # DCF / WACC
        'DCF İçsel Değer':    wacc_veri.get('DCF İçsel Değer') if wacc_veri else None,
        'DCF İskonto (%)':    wacc_veri.get('DCF İskonto (%)') if wacc_veri else None,
        'DCF Sinyal':         wacc_veri.get('DCF Sinyal', '⚪') if wacc_veri else '⚪',
        'WACC (%)':           wacc_veri.get('WACC (%)') if wacc_veri else None,
        'Beta':               wacc_veri.get('Beta') if wacc_veri else None,
        'DCF Kötü':           wacc_veri.get('DCF Kötü') if wacc_veri else None,
        'DCF İyi':            wacc_veri.get('DCF İyi') if wacc_veri else None,
        'Fiyatın Beklediği Büyüme (%)': wacc_veri.get('Fiyatın Beklediği Büyüme (%)') if wacc_veri else None,
        'Büyüme (%)':         wacc_veri.get('Büyüme (%)') if wacc_veri else None,
        'Net Borç (mn)':      wacc_veri.get('Net Borç (mn)') if wacc_veri else None,
        'FCF Ort. (mn)':      wacc_veri.get('FCF Ort. (mn)') if wacc_veri else None,
        'DCF Not':            wacc_veri.get('DCF Not', '') if wacc_veri else '',
        # Volatilite / Stop / Hedef
        'ATR':                vol_veri.get('ATR') if vol_veri else None,
        'ATR (%)':            vol_veri.get('ATR (%)') if vol_veri else None,
        'Vol Yıllık (%)':     vol_veri.get('Vol Yıllık (%)') if vol_veri else None,
        'Stop (Normal)':      vol_veri.get('Stop Normal') if vol_veri else None,
        'Stop (Dar)':         vol_veri.get('Stop Dar') if vol_veri else None,
        'ATR Hedef 1':        vol_veri.get('Hedef 1') if vol_veri else None,
        'ATR Hedef 2':        vol_veri.get('Hedef 2') if vol_veri else None,
        'Risk/Ödül':          vol_veri.get('Risk/Ödül') if vol_veri else None,
        'Momentum 20G':       vol_veri.get('Momentum 20G') if vol_veri else None,
        'Momentum 60G':       vol_veri.get('Momentum 60G') if vol_veri else None,
        # Pozisyon büyüklüğü
        'Önerilen Adet':      poz_adet,
        'Önerilen Maliyet':   poz_maliyet,
        # Giriş detayı
        'Giriş Detayı':       giris_detay,
        # İsyatirim
        'Yabancı Oranı (%)':  yab,
        'Yabancı Sinyali':    yab_sinyal,
        'F/K Prim/İsk (%)':   isy_veri.get('F/K Prim/İsk (%)'),
        'FD/FAVÖK':           isy_veri.get('FD/FAVÖK'),
        'Sektör':             isy_veri.get('Sektör', '-'),
        'Halka Açıklık (%)':  isy_veri.get('Halka Açıklık (%)'),
        # Sermaye olayı
        'Sermaye Olayı':      sermaye_aciklama,
        # Diğer
        'Haber Kararı':       row_ai.get('Haber Kararı', '-') if len(row_ai) > 0 else '-',
        'AI Gerekçe':         row_ai.get('AI Gerekçe', '') if len(row_ai) > 0 else '',
        'Teknik Karar':       row_t.get('Teknik Karar', '-') if len(row_t) > 0 else '-',
        'RSI':                row_t.get('RSI') if len(row_t) > 0 else None,
        'Aylık (%)':          float(row_s.get('Aylık (%)', 0) or 0) if len(row_s) > 0 else 0,
        'Haftalık (%)':       float(row_s.get('Haftalık (%)', 0) or 0) if len(row_s) > 0 else 0,
        'Zirveden (%)':       float(row_s.get('Zirveden (%)', 0) or 0) if len(row_s) > 0 else 0,
        'Likidite':           row_s.get('Likidite', '-') if len(row_s) > 0 else '-',
        'Güven Detayı':       guven_detay,
        'İSY Detayı':         isy_detay,
        'Güven Uyarıları':    ', '.join(guven_uyari) if guven_uyari else '✅ Temiz',
    })

    print(f"{oneri[:12]} | Giriş:{giris_karar[:10]} | SFinal:{super_final:.0f} | {yab_sinyal[:15]}")
    time.sleep(0.5)

df_sonuc = pd.DataFrame(sonuclar)

# Sıralama: önce portföy, sonra öneri, sonra skor
oneri_sira = {'🟢 GÜÇLÜ AL':0,'🟢 AL':1,'🟡 İZLE':2,'🟠 DİKKAT':3,'⚪ BEKLE':4,'🔴 UZAK DUR':5,'🔴 ALMA':6}
df_sonuc['_sira']  = df_sonuc['Öneri'].map(oneri_sira).fillna(7)
df_sonuc['_pfoy']  = df_sonuc['Portföyde'].apply(lambda x: 0 if '📌' in str(x) else 1)
df_sonuc['_final'] = pd.to_numeric(df_sonuc['Süper Final Skor'], errors='coerce').fillna(0)
df_sonuc = df_sonuc.sort_values(['_pfoy','_sira','_final'], ascending=[True,True,False]).reset_index(drop=True)
df_sonuc = df_sonuc.drop(columns=['_sira','_pfoy','_final'])

TARIH = datetime.now().strftime('%d.%m.%Y %H:%M')

print(f"\n✅ {len(df_sonuc)} hisse tamamlandı!")
print(f"\n🟢 GÜÇLÜ AL / AL:")
for _, r in df_sonuc[df_sonuc['Öneri'].isin(['🟢 GÜÇLÜ AL','🟢 AL'])].iterrows():
    print(f"  {r['Kod']:8s} | {r['Giriş Kararı']:15s} | SFinal:{r['Süper Final Skor']:.0f} | {r.get('Sermaye Olayı','-')[:30]}")
print(f"\n🔴 ALMA (veto):")
for _, r in df_sonuc[df_sonuc['Öneri']=='🔴 ALMA'].iterrows():
    print(f"  {r['Kod']:8s} | {r['Veto Detayı'][:60]}")

# ══════════════════════════════════════════════════════════════
#  EXCEL'E YAZ
# ══════════════════════════════════════════════════════════════
print(f"\n💾 Excel yazılıyor...")
wb = load_workbook(EXCEL_YOLU)

ONERI_R = {
    '🟢 GÜÇLÜ AL':('001A10','00FF88'), '🟢 AL':('002010','00CC66'),
    '🟡 İZLE':    ('1A1400','FFB300'), '🟠 DİKKAT':('1A0A00','FF8C00'),
    '⚪ BEKLE':   ('0D1B2E','AAAAAA'), '🔴 UZAK DUR':('2D0000','FF4757'),
    '🔴 ALMA':    ('2D0000','FF4757'),
}
GIRIS_R = {
    '🟢 ŞİMDİ GİR':('001A10','00FF88'), '🟡 HAZIR OL':('1A1400','FFB300'),
    '🟠 BEKLE':    ('1A0A00','FF8C00'), '⚪ İZLE':    ('0D1B2E','AAAAAA'),
    '🔴 UZAK DUR': ('2D0000','FF4757'),
}
TEK_R = {
    '🟢 TEKNİK ONAY':('001A10','00FF88'), '🟡 ZAYIF ONAY':('1A1400','FFB300'),
    '🔴 TEKNİK RED': ('2D0000','FF4757'), '⚪ NÖTR':      ('0D1B2E','AAAAAA'),
}
HABER_R = {
    '🟢 AL':('001A10','00FF88'), '🟡 TUT':('1A1400','FFB300'),
    '🔴 SAT':('2D0000','FF4757'), '⚪ NÖTR':('0D1B2E','AAAAAA'),
}
GUVEN_R = {
    '🟢 Çok Güvenilir':('001A10','00FF88'), '🟢 Güvenilir':('001A10','00FF88'),
    '🟡 Orta Güven':   ('1A1400','FFB300'), '🟠 Düşük Güven':('1A0A00','FF8C00'),
    '🔴 Güvenilmez':   ('2D0000','FF4757'),
}
DCF_R = {
    '🟢 Çok Ucuz':('001A10','00FF88'), '🟢 Ucuz':('002010','00CC66'),
    '🟡 Makul':   ('1A1400','FFB300'), '🟠 Pahalı':('1A0A00','FF8C00'),
    '🔴 Çok Pahalı':('2D0000','FF4757'),
}

def yaz_sayfa(wb, sayfa_adi, df, kolon_listesi, baslik, tab_renk, piyasa_bilgi=''):
    if sayfa_adi in wb.sheetnames: del wb[sayfa_adi]
    ws = wb.create_sheet(sayfa_adi)
    ws.sheet_tab_color = tab_renk
    ws.sheet_view.showGridLines = False

    k_cols = [c for c in kolon_listesi if c in df.columns]

    ws.merge_cells(f'A1:{get_column_letter(len(k_cols))}1')
    ws['A1'] = baslik
    ws['A1'].font = Font(name='Arial', bold=True, size=13, color=tab_renk)
    ws['A1'].fill = PatternFill('solid', fgColor='020D1A')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 36

    if piyasa_bilgi:
        ws.merge_cells(f'A2:{get_column_letter(len(k_cols))}2')
        ws['A2'] = piyasa_bilgi
        renk = 'FF4757' if '🔴' in piyasa_bilgi else 'FFB300' if '🟡' in piyasa_bilgi else '00FF88'
        ws['A2'].font = Font(name='Arial', size=9, italic=True, color=renk)
        ws['A2'].fill = PatternFill('solid', fgColor='020D1A')
        ws['A2'].alignment = Alignment(horizontal='center')
        ws.row_dimensions[2].height = 16
        baslik_satir = 3
    else:
        baslik_satir = 2

    for ci, col in enumerate(k_cols, 1):
        hc(ws, baslik_satir, ci, col, fg=tab_renk)
    ws.row_dimensions[baslik_satir].height = 28

    for ri, (_, row) in enumerate(df.iterrows(), baslik_satir + 1):
        bg = '2D0000' if '🔴 ALMA' in str(row.get('Öneri','')) else \
             '150800' if '🟠 DİKKAT' in str(row.get('Öneri','')) else \
             '001510' if '📌' in str(row.get('Portföyde','')) else \
             '0D1B2E' if ri % 2 == 0 else '091422'

        for ci, cn in enumerate(k_cols, 1):
            val = row.get(cn)

            if cn in ['Öneri'] and str(val) in ONERI_R:
                bk,fk = ONERI_R[str(val)]
                c = ws.cell(ri, ci, val)
                c.fill = PatternFill('solid', fgColor=bk)
                c.font = Font(name='Arial', bold=True, size=11, color=fk)
                c.alignment = Alignment(horizontal='center', vertical='center')
                c.border = BORDER

            elif cn == 'Giriş Kararı' and str(val) in GIRIS_R:
                bk,fk = GIRIS_R[str(val)]
                c = ws.cell(ri, ci, val)
                c.fill = PatternFill('solid', fgColor=bk)
                c.font = Font(name='Arial', bold=True, size=10, color=fk)
                c.alignment = Alignment(horizontal='center', vertical='center')
                c.border = BORDER

            elif cn == 'Teknik Karar' and str(val) in TEK_R:
                bk,fk = TEK_R[str(val)]
                c = ws.cell(ri, ci, val)
                c.fill = PatternFill('solid', fgColor=bk)
                c.font = Font(name='Arial', bold=True, size=9, color=fk)
                c.alignment = Alignment(horizontal='center', vertical='center')
                c.border = BORDER

            elif cn == 'Haber Kararı' and str(val) in HABER_R:
                bk,fk = HABER_R[str(val)]
                c = ws.cell(ri, ci, val)
                c.fill = PatternFill('solid', fgColor=bk)
                c.font = Font(name='Arial', bold=True, size=10, color=fk)
                c.alignment = Alignment(horizontal='center', vertical='center')
                c.border = BORDER

            elif cn == 'Güven Kategorisi' and str(val) in GUVEN_R:
                bk,fk = GUVEN_R[str(val)]
                c = ws.cell(ri, ci, val)
                c.fill = PatternFill('solid', fgColor=bk)
                c.font = Font(name='Arial', bold=True, size=9, color=fk)
                c.alignment = Alignment(horizontal='center', vertical='center')
                c.border = BORDER

            elif cn == 'DCF Sinyal':
                val_str = str(val) if val else ''
                bg2 = '001A10' if '🟢' in val_str else '1A1400' if '🟡' in val_str else \
                      '1A0A00' if '🟠' in val_str else '2D0000' if '🔴' in val_str else bg
                fk  = '00FF88' if '🟢' in val_str else 'FFB300' if '🟡' in val_str else \
                      'FF8C00' if '🟠' in val_str else 'FF4757' if '🔴' in val_str else 'AAAAAA'
                dc(ws, ri, ci, val, bg=bg2, fg=fk, bold='🟢' in val_str, sz=9)

            elif cn == 'Süper Final Skor':
                c = dc(ws, ri, ci, val, bg=bg, bold=True, sz=13)
                try:
                    v = float(val)
                    c.font = Font(name='Arial', bold=True, size=13,
                        color='FFD700' if v>=150 else '00FF88' if v>=120 else
                              '00D4FF' if v>=100 else 'FFB300' if v>=70 else 'FF4757')
                except: pass

            elif cn in ['Toplam Skor','Teknik Skor','Haber Skoru','İSY Skoru',
                        'DCF Skoru','Sermaye Skoru','Güven Skoru','Giriş Puanı']:
                c = dc(ws, ri, ci, val, bg=bg, sz=9)
                try:
                    v = float(val)
                    c.font = Font(name='Arial', size=9,
                        color='00FF88' if v>0 else 'FF4757' if v<0 else 'AAAAAA')
                except: pass

            elif cn == 'Son Fiyat':
                dc(ws, ri, ci, f"{val}₺" if pd.notna(val) and val else '-',
                   bg=bg, bold=True, fg='FFD700', sz=11)

            elif cn in ['Hedef (Ort.)','DCF İçsel Değer']:
                dc(ws, ri, ci, f"{val}₺" if pd.notna(val) and val else '-',
                   bg=bg, fg='00FF88', bold=True)

            elif cn in ['Stop (Normal)','Stop (Dar)']:
                dc(ws, ri, ci, f"{val}₺" if pd.notna(val) and val else '-',
                   bg=bg, fg='FF4757', bold=True)

            elif cn in ['ATR Hedef 1','ATR Hedef 2']:
                dc(ws, ri, ci, f"{val}₺" if pd.notna(val) and val else '-',
                   bg=bg, fg='00D4FF')

            elif cn in ['Hedefe Uzaklık (%)','DCF İskonto (%)','Aylık (%)','Haftalık (%)',
                        'Zirveden (%)','Momentum 20G','Momentum 60G','F/K Prim/İsk (%)']:
                c = dc(ws, ri, ci, val, bg=bg, sz=9)
                try:
                    v = float(val)
                    c.font = Font(name='Arial', size=9,
                        color='00FF88' if v>0 else 'FF4757' if v<0 else 'AAAAAA')
                except: pass

            elif cn == 'Risk/Ödül':
                c = dc(ws, ri, ci, val, bg=bg, bold=True)
                try:
                    v = float(val)
                    c.font = Font(name='Arial', bold=True, size=10,
                        color='00FF88' if v>=5 else 'FFB300' if v>=2 else 'FF8C00' if v>=1 else 'FF4757')
                except: pass

            elif cn == 'Sermaye Olayı':
                val_str = str(val) if val else ''
                if '🔥' in val_str:   dc(ws, ri, ci, val, bg='2D1500', fg='FF4500', bold=True, sz=8, wrap=True)
                elif '⚡' in val_str:  dc(ws, ri, ci, val, bg='1A1400', fg='FFB300', sz=8, wrap=True)
                elif '📅' in val_str:  dc(ws, ri, ci, val, bg='051020', fg='00D4FF', sz=8, wrap=True)
                else:                  dc(ws, ri, ci, val if val != '-' else '-', bg=bg, fg='AAAAAA', sz=8)

            elif cn == 'Yabancı Sinyali':
                val_str = str(val) if val else ''
                fg = '00FF88' if '🟢' in val_str else 'FF4757' if '🔴' in val_str else 'FFB300'
                dc(ws, ri, ci, val, bg=bg, fg=fg)

            elif cn in ['Veto Detayı','Güven Uyarıları']:
                val_str = str(val) if val else ''
                if '✅' in val_str: dc(ws, ri, ci, val, bg='001A10', fg='00FF88', sz=8)
                else:               dc(ws, ri, ci, val, bg='2D0500', fg='FF6B35', sz=8, wrap=True, align='left')

            elif cn in ['AI Gerekçe','Giriş Detayı','İSY Detayı','Güven Detayı']:
                dc(ws, ri, ci, val, bg='051A10', fg='B2DFDB', sz=8, wrap=True, align='left')

            elif cn == 'RSI':
                c = dc(ws, ri, ci, val, bg=bg, bold=True, sz=10)
                try:
                    v = float(val)
                    c.font = Font(name='Arial', bold=True, size=10,
                        color='00FF88' if v<40 else 'FFB300' if v<65 else 'FF4757')
                except: pass

            elif cn == 'WACC (%)':
                c = dc(ws, ri, ci, val, bg=bg, bold=True, sz=9)
                try:
                    v = float(val)
                    c.font = Font(name='Arial', bold=True, size=9,
                        color='00FF88' if v<25 else 'FFB300' if v<40 else 'FF4757')
                except: pass

            elif cn == 'Portföyde':
                val_str = str(val) if val else ''
                fg = 'FFD700' if '📌' in val_str else 'AAAAAA'
                dc(ws, ri, ci, val, bg=bg, fg=fg, bold='📌' in val_str)

            elif cn == 'Önerilen Adet':
                c = dc(ws, ri, ci, val, bg=bg, bold=True, sz=10)
                if pd.notna(val) and val:
                    c.font = Font(name='Arial', bold=True, size=10, color='00D4FF')

            elif cn == 'Önerilen Maliyet':
                dc(ws, ri, ci, f"{val:,.0f}₺" if pd.notna(val) and val else '-',
                   bg=bg, fg='E8F4F8')

            else:
                dc(ws, ri, ci, val, bg=bg)

        ws.row_dimensions[ri].height = 38

    ws.freeze_panes = f'D{baslik_satir+1}'
    ws.auto_filter.ref = f'A{baslik_satir}:{get_column_letter(len(k_cols))}{baslik_satir}'
    return ws, k_cols

# Piyasa bilgisi
piyasa_bilgi = f"📡 Makro Skor: {makro_skor_global:+d}" if makro_skor_global is not None else ''
if piyasa_bilgi and makro_zaman is not None:
    piyasa_bilgi += f" ({makro_zaman.strftime('%d.%m %H:%M')}" + (" — ESKİ!)" if makro_yas_saat and makro_yas_saat > 3 else ")")
if makro_skor_global is not None:
    if makro_skor_global >= 20:    piyasa_bilgi += ' | 🟢 UYGUN ORTAM'
    elif makro_skor_global >= 0:   piyasa_bilgi += ' | 🟡 DİKKATLİ OL'
    elif makro_skor_global >= -20: piyasa_bilgi += ' | 🟠 RİSKLİ'
    else:                          piyasa_bilgi += ' | 🔴 RİSK YÜKSEK — Yeni pozisyon açma!'

# ── 🎯 KOMBİNE AL — ANA SAYFA ──────────────────────────────
kombine_cols = [
    'Kod','Portföyde','Öneri','Giriş Kararı','Veto Sayısı',
    'Son Fiyat','Hedef (Ort.)','Hedefe Uzaklık (%)','Süper Final Skor',
    'Toplam Skor','Teknik Skor','Haber Skoru','İSY Skoru','DCF Skoru','Sermaye Skoru','Güven Skoru',
    'DCF Sinyal','DCF İçsel Değer','DCF Kötü','DCF İyi','DCF İskonto (%)','Fiyatın Beklediği Büyüme (%)','Büyüme (%)','WACC (%)','Beta','Net Borç (mn)','FCF Ort. (mn)','DCF Not',
    'Stop (Normal)','Stop (Dar)','ATR Hedef 1','ATR Hedef 2','Risk/Ödül',
    'Önerilen Adet','Önerilen Maliyet',
    'Sermaye Olayı','Teknik Karar','RSI','Haber Kararı',
    'Yabancı Oranı (%)','Yabancı Sinyali','F/K Prim/İsk (%)','FD/FAVÖK',
    'Güven Kategorisi','Güven Uyarıları',
    'Aylık (%)','Haftalık (%)','Zirveden (%)','Likidite','Sektör',
    'AI Gerekçe','Giriş Detayı','Veto Detayı','İSY Detayı',
]

ws_k, kc = yaz_sayfa(wb, '🎯 Kombine AL', df_sonuc, kombine_cols,
    f"🎯 KOMBİNE AL — {TARIH} | {len(df_sonuc)} hisse", '00FF88', piyasa_bilgi)

# Kolon genişlikleri
gw = {
    'Kod':8,'Portföyde':8,'Öneri':14,'Giriş Kararı':14,'Veto Sayısı':8,
    'Son Fiyat':10,'Hedef (Ort.)':10,'Hedefe Uzaklık (%)':11,'Süper Final Skor':13,
    'Toplam Skor':9,'Teknik Skor':9,'Haber Skoru':9,'İSY Skoru':8,
    'DCF Skoru':8,'Sermaye Skoru':9,'Güven Skoru':9,
    'DCF Sinyal':13,'DCF İçsel Değer':12,'DCF İskonto (%)':12,'WACC (%)':9,'Beta':7,
    'Stop (Normal)':11,'Stop (Dar)':10,'ATR Hedef 1':11,'ATR Hedef 2':11,'Risk/Ödül':9,
    'Önerilen Adet':12,'Önerilen Maliyet':14,
    'Sermaye Olayı':35,'Teknik Karar':14,'RSI':7,'Haber Kararı':12,
    'Yabancı Oranı (%)':11,'Yabancı Sinyali':15,'F/K Prim/İsk (%)':12,'FD/FAVÖK':9,
    'Güven Kategorisi':16,'Güven Uyarıları':28,
    'Aylık (%)':9,'Haftalık (%)':9,'Zirveden (%)':10,'Likidite':10,'Sektör':18,
    'AI Gerekçe':40,'Giriş Detayı':40,'Veto Detayı':45,'İSY Detayı':35,
}
for ci, cn in enumerate(kc, 1):
    ws_k.column_dimensions[get_column_letter(ci)].width = gw.get(cn, 10)

# ── 🏆 SÜPER AL — Filtreli liste ────────────────────────────
super_al_df = df_sonuc[
    (df_sonuc['Öneri'].isin(['🟢 GÜÇLÜ AL','🟢 AL'])) &
    (df_sonuc['Giriş Kararı'].isin(['🟢 ŞİMDİ GİR','🟡 HAZIR OL'])) &
    (df_sonuc['Güven Skoru'] >= 50) &
    (df_sonuc['Süper Final Skor'] >= 50)
].copy()

super_cols = [
    'Kod','Portföyde','Öneri','Giriş Kararı',
    'Son Fiyat','Hedef (Ort.)','Hedefe Uzaklık (%)','Süper Final Skor',
    'DCF Sinyal','WACC (%)','Beta',
    'Stop (Normal)','ATR Hedef 2','Risk/Ödül',
    'Önerilen Adet','Önerilen Maliyet',
    'Sermaye Olayı','Güven Kategorisi','Yabancı Sinyali',
    'F/K Prim/İsk (%)','Sektör','Giriş Detayı','Veto Detayı',
]
ws_s, sc = yaz_sayfa(wb, '🏆 Süper AL', super_al_df, super_cols,
    f"🏆 SÜPER AL — {TARIH} | {len(super_al_df)} hisse | {piyasa_bilgi}", 'FFD700')
for ci, cn in enumerate(sc, 1):
    ws_s.column_dimensions[get_column_letter(ci)].width = gw.get(cn, 10)

# ── 🔬 GÜVEN & İSY — Detay sayfası ─────────────────────────
guven_cols = [
    'Kod','Son Fiyat','Hedef (Ort.)','Süper Final Skor',
    'Güven Skoru','Güven Kategorisi','Güven Uyarıları',
    'DCF Sinyal','DCF İçsel Değer','DCF Kötü','DCF İyi','DCF İskonto (%)','Fiyatın Beklediği Büyüme (%)','Büyüme (%)','WACC (%)','Beta','Net Borç (mn)','FCF Ort. (mn)','DCF Not',
    'İSY Skoru','Yabancı Oranı (%)','Yabancı Sinyali',
    'F/K Prim/İsk (%)','PD/DD Prim/İsk (%)' if 'PD/DD Prim/İsk (%)' in df_sonuc.columns else 'FD/FAVÖK',
    'Halka Açıklık (%)','Sektör','Sermaye Olayı',
    'Güven Detayı','İSY Detayı',
]
guven_cols = [c for c in guven_cols if c in df_sonuc.columns]
ws_g, gc = yaz_sayfa(wb, '🔬 Güven & İSY', df_sonuc, guven_cols,
    f"🔬 GÜVEN & İSY — {TARIH} | {len(df_sonuc)} hisse", 'AA44FF')
for ci, cn in enumerate(gc, 1):
    ws_g.column_dimensions[get_column_letter(ci)].width = gw.get(cn, 10)

wb.save(EXCEL_YOLU)
print(f"\n✅ Excel güncellendi!")
print(f"   🎯 Kombine AL — {len(df_sonuc)} hisse | {len(kombine_cols)} kolon")
print(f"   🏆 Süper AL   — {len(super_al_df)} hisse")
print(f"   🔬 Güven & İSY")

# ── KAYIT: bugünkü kararlar (hücre 9 ileride bunların sonucunu ölçer) ──
KAYIT_YOLU = f'{DRIVE}/sistem_kayit.csv'
try:
    _kol = [c for c in ['Kod', 'Öneri', 'Giriş Kararı', 'Giriş Puanı', 'Son Fiyat', 'Hedef (Ort.)',
                        'Toplam Skor', 'Teknik Skor', 'Haber Kararı', 'Haber Skoru', 'Hukuki Risk',
                        'Süper Final Skor', 'DCF Sinyal', 'Risk/Ödül'] if c in df_sonuc.columns]
    kayit = df_sonuc[_kol].copy()
    kayit.insert(0, 'Tarih', datetime.now().strftime('%Y-%m-%d'))
    kayit['Makro Skor'] = makro_skor_global
    kayit['Süper AL'] = kayit['Kod'].isin(super_al_df['Kod']) if len(super_al_df) else False
    if os.path.exists(KAYIT_YOLU):
        eski = pd.read_csv(KAYIT_YOLU)
        kayit = pd.concat([eski[eski['Tarih'] != kayit['Tarih'].iloc[0]], kayit], ignore_index=True)
    kayit.to_csv(KAYIT_YOLU, index=False)
    print(f"   📝 Kayıt: bugünkü {len(df_sonuc)} karar sistem_kayit.csv'ye yazıldı "
          f"(toplam {kayit['Tarih'].nunique()} gün)")
except Exception as e:
    print(f"   ⚠️ Kayıt yazılamadı: {e}")

# ══════════════════════════════════════════════════════════════
#  PORTFÖY TAKİP
# ══════════════════════════════════════════════════════════════
print(f"\n{'='*60}")
print(f"  📌 PORTFÖY TAKİP")
print(f"{'='*60}")

portfoy_data = []
for p in PORTFOY:
    kod = p['Kod']
    try:
        t = yf.Ticker(f"{kod}.IS")
        hist = t.history(period='5d')['Close'].dropna()
        son_f = hist.iloc[-1] if len(hist) > 0 else p['Alis']
    except:
        son_f = p['Alis']

    kz_yuz = ((son_f - p['Alis']) / p['Alis']) * 100
    kz_tl  = (son_f - p['Alis']) * p['Adet']
    stop_m = ((son_f - p['Stop']) / son_f) * 100

    if stop_m <= 3:    stop_d = '🔴 STOP YAKLAŞTI!'
    elif stop_m <= 8:  stop_d = '🟠 Dikkat'
    else:              stop_d = '🟢 Güvende'

    if kz_yuz >= 20:    durum = '🚀 Güçlü Kar'
    elif kz_yuz >= 5:   durum = '✅ Karda'
    elif kz_yuz >= 0:   durum = '🟡 Başabaş'
    elif kz_yuz >= -8:  durum = '🟠 Hafif Zarar'
    else:               durum = '🔴 Zararda'

    # Çıkış stratejisi
    hedef = float(df_sistem[df_sistem['Kod']==kod]['Hedef (Ort.)'].iloc[0]) \
            if len(df_sistem[df_sistem['Kod']==kod]) > 0 else 0
    vol = next((r for r in sonuclar if r['Kod'] == kod), None)
    if vol:
        vol_d = {
            'Hedef 1':     vol.get('ATR Hedef 1'),
            'Hedef 2':     vol.get('ATR Hedef 2'),
            'Stop Normal': vol.get('Stop (Normal)'),
            'Son Fiyat':   vol.get('Son Fiyat'),
        }
    else:
        vol_d = None
    cikis_mesaj, cikis_renk, cikis_islem, cikis_detay, trailing_not = cikis_stratejisi(
        p['Alis'], son_f, p['Stop'], hedef, vol_d, p.get('Tarih')
    )

    # ATR hedefleri
    atr_h1 = vol_d.get('Hedef 1', '-') if vol_d else '-'
    atr_h2 = vol_d.get('Hedef 2', '-') if vol_d else '-'

    # Sistem önerisi
    sistem_oneri = next((r['Öneri'] for r in sonuclar if r['Kod'] == kod), '-')
    giris_k      = next((r['Giriş Kararı'] for r in sonuclar if r['Kod'] == kod), '-')
    super_final  = next((r['Süper Final Skor'] for r in sonuclar if r['Kod'] == kod), '-')

    portfoy_data.append({
        'Kod':              kod,
        'Adet':             p['Adet'],
        'Alış':             p['Alis'],
        'Son Fiyat':        round(son_f, 2),
        'Stop':             p['Stop'],
        'ATR Hedef 1':      atr_h1,
        'ATR Hedef 2':      atr_h2,
        'Stop Mesafe (%)':  round(stop_m, 1),
        'Stop Durumu':      stop_d,
        'K/Z (%)':          round(kz_yuz, 1),
        'K/Z (TL)':         round(kz_tl, 0),
        'Maliyet (TL)':     round(p['Alis'] * p['Adet'], 0),
        'Güncel Değer':     round(son_f * p['Adet'], 0),
        'Durum':            durum,
        'Ne Yapmalıyım':    cikis_islem,
        'Çıkış Stratejisi': cikis_mesaj,
        'Çıkış Detayı':     cikis_detay,
        'Trailing Stop':    trailing_not,
        'Sistem Öneri':     sistem_oneri,
        'Giriş Kararı':     giris_k,
        'Süper Final':      super_final,
    })

    emoji = '🟢' if kz_yuz >= 0 else '🔴'
    print(f"\n  📌 {kod}")
    print(f"     Alış:{p['Alis']}₺ → Şu an:{son_f:.2f}₺ | {emoji} K/Z:%{kz_yuz:+.1f} ({kz_tl:+.0f}₺)")
    print(f"     Stop:{p['Stop']}₺ | ATR H1:{atr_h1} | ATR H2:{atr_h2}")
    print(f"     👉 {cikis_islem} — {cikis_mesaj[:55]}")

df_portfoy = pd.DataFrame(portfoy_data)

if len(df_portfoy) == 0:
    print("\n  Portföy boş — yeni hisse eklemek için PORTFOY listesini güncelle")
    top_maliyet = 0; top_deger = 0; top_kz = 0; top_kz_yuz = 0
else:
    top_maliyet = df_portfoy['Maliyet (TL)'].sum()
    top_deger   = df_portfoy['Güncel Değer'].sum()
    top_kz      = top_deger - top_maliyet
    top_kz_yuz  = (top_kz / top_maliyet) * 100

print(f"\n  {'─'*45}")
print(f"  Toplam: {top_maliyet:,.0f}₺ → {top_deger:,.0f}₺ | K/Z:{top_kz:+,.0f}₺ (%{top_kz_yuz:+.1f})")

# Portföy Excel sayfası
wb2 = load_workbook(EXCEL_YOLU)
if '📌 Portföy' in wb2.sheetnames: del wb2['📌 Portföy']
ws_p = wb2.create_sheet('📌 Portföy', 0)
ws_p.sheet_tab_color = 'FFD700'
ws_p.sheet_view.showGridLines = False

p_cols = ['Kod','Adet','Alış','Son Fiyat','Stop','ATR Hedef 1','ATR Hedef 2',
          'Stop Mesafe (%)','Stop Durumu','K/Z (%)','K/Z (TL)',
          'Maliyet (TL)','Güncel Değer','Durum',
          'Ne Yapmalıyım','Çıkış Stratejisi','Çıkış Detayı','Trailing Stop',
          'Sistem Öneri','Giriş Kararı','Süper Final']

ws_p.merge_cells(f'A1:{get_column_letter(len(p_cols))}1')
ws_p['A1'] = f"📌 PORTFÖY — {TARIH} | K/Z: {top_kz:+,.0f}₺ (%{top_kz_yuz:+.1f})"
renk_p = '00FF88' if top_kz >= 0 else 'FF4757'
ws_p['A1'].font = Font(name='Arial', bold=True, size=13, color=renk_p)
ws_p['A1'].fill = PatternFill('solid', fgColor='020D1A')
ws_p['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws_p.row_dimensions[1].height = 36

for ci, col in enumerate(p_cols, 1):
    hc(ws_p, 2, ci, col, fg='FFD700')
ws_p.row_dimensions[2].height = 28

DURUM_R = {
    '🚀 Güçlü Kar':  ('001A10','00FF88'), '✅ Karda':      ('001A10','00D4FF'),
    '🟡 Başabaş':    ('1A1400','FFB300'), '🟠 Hafif Zarar':('1A0A00','FF8C00'),
    '🔴 Zararda':    ('2D0000','FF4757'),
}
STOP_R = {
    '🔴 STOP YAKLAŞTI!':('2D0000','FF4757'),
    '🟠 Dikkat':         ('1A0A00','FF8C00'),
    '🟢 Güvende':        ('001A10','00FF88'),
}

for ri, (_, row) in enumerate(df_portfoy.iterrows(), 3):
    bg = '0D1B2E' if ri % 2 == 0 else '091422'
    for ci, cn in enumerate(p_cols, 1):
        val = row.get(cn)
        if cn == 'Durum' and str(val) in DURUM_R:
            bk,fk = DURUM_R[str(val)]
            c = ws_p.cell(ri, ci, val)
            c.fill = PatternFill('solid', fgColor=bk)
            c.font = Font(name='Arial', bold=True, size=11, color=fk)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = BORDER
        elif cn == 'Stop Durumu' and str(val) in STOP_R:
            bk,fk = STOP_R[str(val)]
            c = ws_p.cell(ri, ci, val)
            c.fill = PatternFill('solid', fgColor=bk)
            c.font = Font(name='Arial', bold=True, size=10, color=fk)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = BORDER
        elif cn == 'Sistem Öneri' and str(val) in ONERI_R:
            bk,fk = ONERI_R[str(val)]
            c = ws_p.cell(ri, ci, val)
            c.fill = PatternFill('solid', fgColor=bk)
            c.font = Font(name='Arial', bold=True, size=10, color=fk)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = BORDER
        elif cn == 'Giriş Kararı' and str(val) in GIRIS_R:
            bk,fk = GIRIS_R[str(val)]
            c = ws_p.cell(ri, ci, val)
            c.fill = PatternFill('solid', fgColor=bk)
            c.font = Font(name='Arial', bold=True, size=10, color=fk)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = BORDER
        elif cn == 'K/Z (%)':
            c = dc(ws_p, ri, ci, f"%{val:+.1f}" if pd.notna(val) else '-', bg=bg, bold=True, sz=12)
            try:
                c.font = Font(name='Arial', bold=True, size=12,
                    color='00FF88' if float(val) >= 0 else 'FF4757')
            except: pass
        elif cn == 'K/Z (TL)':
            c = dc(ws_p, ri, ci, f"{val:+,.0f}₺" if pd.notna(val) else '-', bg=bg, bold=True, sz=11)
            try:
                c.font = Font(name='Arial', bold=True, size=11,
                    color='00FF88' if float(val) >= 0 else 'FF4757')
            except: pass
        elif cn == 'Stop Mesafe (%)':
            c = dc(ws_p, ri, ci, f"%{val:.1f}" if pd.notna(val) else '-', bg=bg, bold=True)
            try:
                c.font = Font(name='Arial', bold=True, size=10,
                    color='FF4757' if float(val)<=3 else 'FF8C00' if float(val)<=8 else '00FF88')
            except: pass
        elif cn in ['Alış','Son Fiyat','Stop']:
            dc(ws_p, ri, ci, f"{val}₺" if pd.notna(val) else '-', bg=bg, bold=True, fg='FFD700')
        elif cn in ['Maliyet (TL)','Güncel Değer']:
            dc(ws_p, ri, ci, f"{val:,.0f}₺" if pd.notna(val) else '-', bg=bg, fg='E8F4F8')
        elif cn == 'Ne Yapmalıyım':
            val_str = str(val) if val else ''
            if '🛑' in val_str:   dc(ws_p, ri, ci, val, bg='2D0000', fg='FF4757', bold=True, sz=12)
            elif '🎯' in val_str:  dc(ws_p, ri, ci, val, bg='001A10', fg='00FF88', bold=True, sz=12)
            elif '⚡' in val_str:  dc(ws_p, ri, ci, val, bg='1A1400', fg='FFB300', bold=True, sz=12)
            elif '⚠️' in val_str:  dc(ws_p, ri, ci, val, bg='1A0A00', fg='FF8C00', bold=True, sz=12)
            elif '✅' in val_str:  dc(ws_p, ri, ci, val, bg='001A10', fg='00D4FF', bold=True, sz=12)
            else:                  dc(ws_p, ri, ci, val, bg=bg, fg='FFB300', bold=True, sz=12)
        elif cn == 'Çıkış Stratejisi':
            val_str = str(val) if val else ''
            if '🛑' in val_str:   dc(ws_p, ri, ci, val, bg='2D0000', fg='FF4757', bold=True, sz=9, wrap=True)
            elif '🎯' in val_str:  dc(ws_p, ri, ci, val, bg='001A10', fg='00FF88', bold=True, sz=9, wrap=True)
            elif '⚡' in val_str:  dc(ws_p, ri, ci, val, bg='1A1400', fg='FFB300', sz=9, wrap=True)
            elif '⚠️' in val_str:  dc(ws_p, ri, ci, val, bg='1A0A00', fg='FF8C00', sz=9, wrap=True)
            else:                  dc(ws_p, ri, ci, val, bg=bg, fg='E8F4F8', sz=9, wrap=True)
        elif cn in ['Çıkış Detayı','Trailing Stop']:
            dc(ws_p, ri, ci, val, bg='051A10', fg='B2DFDB', sz=8, wrap=True)
        elif cn == 'Süper Final':
            c = dc(ws_p, ri, ci, val, bg=bg, bold=True, sz=11)
            try:
                v = float(val)
                c.font = Font(name='Arial', bold=True, size=11,
                    color='FFD700' if v>=150 else '00FF88' if v>=120 else '00D4FF' if v>=100 else 'FFB300')
            except: pass
        elif cn in ['ATR Hedef 1','ATR Hedef 2']:
            dc(ws_p, ri, ci, f"{val}₺" if pd.notna(val) and val and val != '-' else '-',
               bg=bg, fg='00D4FF', bold=True)
        else:
            dc(ws_p, ri, ci, val, bg=bg)
    ws_p.row_dimensions[ri].height = 45

# Özet satırı
oz_r = len(df_portfoy) + 3
ws_p.merge_cells(f'A{oz_r}:{get_column_letter(len(p_cols))}{oz_r}')
ws_p[f'A{oz_r}'] = f"TOPLAM → Maliyet:{top_maliyet:,.0f}₺ | Değer:{top_deger:,.0f}₺ | K/Z:{top_kz:+,.0f}₺ (%{top_kz_yuz:+.1f})"
ws_p[f'A{oz_r}'].font = Font(name='Arial', bold=True, size=12, color=renk_p)
ws_p[f'A{oz_r}'].fill = PatternFill('solid', fgColor='001A10' if top_kz >= 0 else '2D0000')
ws_p[f'A{oz_r}'].alignment = Alignment(horizontal='center', vertical='center')
ws_p.row_dimensions[oz_r].height = 35

gw_p = {'Kod':12,'Adet':8,'Alış':10,'Son Fiyat':10,'Stop':10,
        'ATR Hedef 1':11,'ATR Hedef 2':11,
        'Stop Mesafe (%)':13,'Stop Durumu':18,'K/Z (%)':10,'K/Z (TL)':14,
        'Maliyet (TL)':14,'Güncel Değer':14,'Durum':14,
        'Ne Yapmalıyım':16,'Çıkış Stratejisi':50,'Çıkış Detayı':30,
        'Trailing Stop':28,'Sistem Öneri':14,'Giriş Kararı':14,'Süper Final':12}
for ci, cn in enumerate(p_cols, 1):
    ws_p.column_dimensions[get_column_letter(ci)].width = gw_p.get(cn, 10)
ws_p.freeze_panes = 'A3'

wb2.save(EXCEL_YOLU)
print(f"\n✅ Portföy sayfası güncellendi!")
print(f"\n🎉 HÜCRE 7 TAMAMLANDI!")
print(f"\n📊 ÖZET:")
print(f"   Analiz edilen: {len(df_sonuc)} hisse")
print(f"   🟢 Güçlü AL / AL: {len(df_sonuc[df_sonuc['Öneri'].isin(['🟢 GÜÇLÜ AL','🟢 AL'])])}")
print(f"   🟡 İZLE: {len(df_sonuc[df_sonuc['Öneri']=='🟡 İZLE'])}")
print(f"   🔴 ALMA: {len(df_sonuc[df_sonuc['Öneri']=='🔴 ALMA'])}")
print(f"   📌 Portföy K/Z: {top_kz:+,.0f}₺ (%{top_kz_yuz:+.1f})")