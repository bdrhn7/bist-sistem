# ═══════════════════════════════════════════════════════════════
#  HÜCRE 6 — YAPAY ZEKÂ HABER ANALİZİ  (v2.3 — 15.10.2026)
#  v2 değişiklikleri:
#   - API anahtarı koddan çıkarıldı: Colab Secrets'tan okunuyor
#   - Güncel model; başlamadan önce bağlantı testi, hata varsa sebebi yazılıyor
#     (eskiden her hisse sessizce "⚪ HATA" ve 0 puan alıyordu)
#   - Yapay zekâya haberlerin tarihi de veriliyor; önemli haber yoksa NÖTR
#   - Cevap okuma sağlamlaştırıldı ("SAT" yazıp içinde "AL" geçen cevaplar
#     artık yanlışlıkla AL sayılmıyor)
#   - "🎯 Kombine AL" sayfasını artık sadece hücre 7 yazıyor
#   - Analiz edilecek hisse yoksa çökmüyor
#   - v2.1: yapay zekâ hukuki riski ayrı bir satırda (RİSK: VAR/YOK) bildiriyor;
#     hücre 7'nin vetosu buna bakıyor
#   - v2.2: geçici bağlantı hatasında 3 kez deneniyor; yine olmazsa hata sebebi
#     'AI Gerekçe' sütununa yazılıyor (DOCO'da sebepsiz "⚪ HATA" çıkmıştı)
#   - v2.3: önbellek — bir hissenin haber başlıkları değişmediyse yapay zekâya tekrar
#     sorulmuyor, önceki cevap kullanılıyor (Drive'da ai_haber_onbellek.json, 7 gün).
#     Aynı gün tekrar çalıştırmak neredeyse hiç API kredisi harcamaz.
# ═══════════════════════════════════════════════════════════════
import anthropic
import pandas as pd
import numpy as np
import re
import time
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import json, hashlib

# ── AYARLAR ──────────────────────────────────────────────
MODEL       = "claude-sonnet-5-5"   # çalışmazsa hata mesajında sebebi görünür
SISTEM_YOLU = '/content/drive/MyDrive/isyatirim/BIST_Sistem.xlsx'
EXCEL_YOLU  = '/content/drive/MyDrive/isyatirim/BIST_Tam_KAP_v9_11mart.xlsx'
ONBELLEK_YOLU = '/content/drive/MyDrive/isyatirim/ai_haber_onbellek.json'
ONBELLEK_GUN  = 7     # bu kadar günden eski cevaplar silinir

# ── ÖNBELLEK: aynı başlıklar için yapay zekâya tekrar sorma ──
try:
    with open(ONBELLEK_YOLU, encoding='utf-8') as f:
        ONBELLEK = json.load(f)
except Exception:
    ONBELLEK = {}
_sinir = (datetime.now() - timedelta(days=ONBELLEK_GUN)).strftime('%Y-%m-%d')
ONBELLEK = {k: v for k, v in ONBELLEK.items() if v.get('tarih', '') >= _sinir}
onbellek_sayisi = 0

# ── API ANAHTARI: Colab Secrets'tan ──────────────────────
# Sol menü → 🔑 Secrets → Name: ANTHROPIC_API_KEY, Value: anahtarın, Notebook access: açık
API_KEY = None
try:
    from google.colab import userdata
    API_KEY = userdata.get('ANTHROPIC_API_KEY')
except Exception as e:
    print(f"⚠️ API anahtarı Secrets'ta bulunamadı: {type(e).__name__}")
    print("   Sol menüdeki 🔑 Secrets bölümüne ANTHROPIC_API_KEY adıyla ekle.")

BORDER = Border(left=Side(style='thin',color='1E3A5F'), right=Side(style='thin',color='1E3A5F'),
                top=Side(style='thin',color='1E3A5F'), bottom=Side(style='thin',color='1E3A5F'))

KARAR_R = {
    '🟢 AL':   ('001A10','00FF88'),
    '🔴 SAT':  ('2D0000','FF4757'),
    '🟡 TUT':  ('1A1400','FFB300'),
    '⚪ NÖTR': ('0D1B2E','AAAAAA'),
    '⚪ HATA': ('0D1B2E','AAAAAA'),
}
KARAR_ESLE = {'AL': '🟢 AL', 'SAT': '🔴 SAT', 'TUT': '🟡 TUT', 'NÖTR': '⚪ NÖTR', 'NOTR': '⚪ NÖTR'}

# ── BAĞLANTI TESTİ ───────────────────────────────────────
client = None
AI_AKTIF = False
if API_KEY:
    try:
        client = anthropic.Anthropic(api_key=API_KEY)
        client.messages.create(model=MODEL, max_tokens=5,
                               messages=[{"role": "user", "content": "Sadece OK yaz."}])
        AI_AKTIF = True
        print(f"✅ Yapay zekâ bağlantısı çalışıyor (model: {MODEL})")
    except Exception as e:
        print(f"❌ Yapay zekâ bağlantı hatası: {type(e).__name__}: {str(e)[:200]}")
        print("   Bu mesajı Claude'a göster. Analiz yapılmadan devam ediliyor (haber puanı 0).")

def yanit_oku(yanit):
    """KARAR / SKOR / GEREKÇE satırlarını sağlam şekilde okur."""
    metin = yanit.replace('*', '')
    karar = '⚪ NÖTR'; skor = 0; gerekce = ''; risk = '✅ YOK'
    m = re.search(r'KARAR\s*:\s*(AL|TUT|SAT|NÖTR|NOTR)\b', metin, re.IGNORECASE)
    if m:
        karar = KARAR_ESLE.get(m.group(1).upper(), '⚪ NÖTR')
    m = re.search(r'SKOR\s*:\s*([+-]?\d+)', metin, re.IGNORECASE)
    if m:
        skor = max(-10, min(10, int(m.group(1))))
    m = re.search(r'GEREK[CÇ]E\s*:\s*(.+)', metin, re.IGNORECASE)
    if m:
        gerekce = m.group(1).strip()[:300]
    m = re.search(r'R[İI]SK\s*:\s*(VAR|YOK)', metin, re.IGNORECASE)
    if m and m.group(1).upper() == 'VAR':
        risk = '🚫 VAR'
    if karar == '⚪ NÖTR':
        skor = 0
    return karar, skor, gerekce, risk

def haber_analiz(kod, haberler):
    """haberler: [(tarih, başlık), ...] en yeni başta"""
    global onbellek_sayisi
    if not haberler:
        return "Haber yok", "⚪ NÖTR", 0, "Son günlerde haber yok", "-"
    haber_metni = "\n".join(f"- {t} | {b}" for t, b in haberler[:10])
    anahtar = kod + '|' + hashlib.sha1(haber_metni.encode('utf-8')).hexdigest()
    if anahtar in ONBELLEK:                       # aynı başlıklar daha önce soruldu
        o = ONBELLEK[anahtar]; onbellek_sayisi += 1
        return o['yanit'], o['karar'], o['skor'], o['gerekce'], o['risk']
    if not AI_AKTIF:
        return "AI kapalı", "⚪ HATA", 0, "Yapay zekâ bağlantısı yok", "-"
    prompt = f"""Aşağıda Borsa İstanbul'da {kod} koduyla işlem gören şirketle ilgili son haber başlıkları var (tarih | başlık):
{haber_metni}

Görev: Sadece bu başlıklara bakarak haberlerin şirketin hissesi için önümüzdeki 1-3 ayda olumlu mu olumsuz mu olduğunu değerlendir.
Kurallar:
- Başlık bu şirketle ilgili değilse yok say.
- "Hisse yükseldi/düştü", fiyat yorumu, teknik analiz gibi piyasa özetleri yeni bilgi değildir; bunlara ağırlık verme.
- Somut gelişmelere bak: bilanço, sözleşme, yatırım, temettü, sermaye artırımı, dava, SPK veya KAP açıklaması, yönetim değişikliği.
- Önemli bir gelişme yoksa KARAR: NÖTR ve SKOR: 0 ver.
- Haberlerde olmayan bilgiyi kullanma, fiyat tahmini yapma.

Cevabı sadece şu dört satırla ver:
KARAR: AL veya TUT veya SAT veya NÖTR
SKOR: -10 ile +10 arası tam sayı
RİSK: VAR veya YOK (VAR sadece haberlerde bu şirketle ilgili SPK soruşturması, dava, manipülasyon, suç duyurusu ya da işlem yasağı açıkça geçiyorsa)
GEREKÇE: tek cümle"""
    son_hata = ''
    for deneme in range(3):                     # geçici hatalarda tekrar dene
        try:
            mesaj = client.messages.create(
                model=MODEL,
                max_tokens=200,
                messages=[{"role": "user", "content": prompt}]
            )
            yanit = mesaj.content[0].text.strip()
            karar, skor, gerekce, risk = yanit_oku(yanit)
            ONBELLEK[anahtar] = {'yanit': yanit, 'karar': karar, 'skor': skor, 'gerekce': gerekce,
                                 'risk': risk, 'tarih': datetime.now().strftime('%Y-%m-%d')}
            return yanit, karar, skor, gerekce, risk
        except Exception as e:
            son_hata = f"{type(e).__name__}: {str(e)[:150]}"
            time.sleep(5 * (deneme + 1))
    return son_hata, "⚪ HATA", 0, f"AI hatası (3 deneme): {son_hata[:120]}", "-"

def uyari_hesapla(row_sis):
    uyarilar = []
    try:
        aylik = float(row_sis.get('Aylık (%)', 0) or 0)
        haftalik = float(row_sis.get('Haftalık (%)', 0) or 0)
        zirveden = float(row_sis.get('Zirveden (%)', 0) or 0)
        dipten = float(row_sis.get('Dipten (%)', 0) or 0)
        if aylik >= 50:
            uyarilar.append(f'⚠️ Aylık +%{aylik:.0f} spekülatif risk!')
        elif aylik >= 30:
            uyarilar.append(f'🟡 Aylık +%{aylik:.0f} dikkatli ol')
        if haftalik >= 20:
            uyarilar.append(f'⚠️ Haftalık +%{haftalik:.0f} ani yükseliş!')
        if zirveden >= -5:
            uyarilar.append('📍 52H zirvesinde')
        if dipten >= 100:
            uyarilar.append(f'📈 Dipten +%{dipten:.0f} uzun yol gelmiş')
    except:
        pass
    return ' | '.join(uyarilar) if uyarilar else '✅ Normal'

# ── VERİLERİ OKU ──
print("\n📰 Haber analizi başlıyor...")
df_haberler = pd.read_excel(EXCEL_YOLU, sheet_name='📰 Haberler', header=1)
df_haberler = df_haberler.iloc[:, :5]
df_haberler.columns = ['Hisse','Kaynak','Tarih','Başlık','Link']
df_haberler['Hisse'] = df_haberler['Hisse'].astype(str).str.replace('📌', '').str.strip()
print(f"✅ {len(df_haberler)} haber yüklendi")

df_teknik = pd.read_excel(EXCEL_YOLU, sheet_name='📊 Teknik Analiz', header=2)
df_sistem_uy = pd.read_excel(SISTEM_YOLU)
df_sistem_uy = df_sistem_uy.rename(columns={'Sembol':'Kod'}) if 'Sembol' in df_sistem_uy.columns else df_sistem_uy

if 'Teknik Karar' in df_teknik.columns:
    df_analiz = df_teknik[df_teknik['Teknik Karar'].isin(['🟢 TEKNİK ONAY','🟡 ZAYIF ONAY'])].copy()
else:
    df_analiz = pd.DataFrame(columns=['Kod'])
kodlar = df_analiz['Kod'].dropna().tolist()
print(f"🎯 {len(kodlar)} hisse analiz edilecek\n")

# Piyasa yönü
try:
    wb_check = load_workbook(EXCEL_YOLU)
    piyasa_str = wb_check['📊 Teknik Analiz']['A2'].value or ""
    piyasa_uyari = '⚠️ PİYASA AŞAĞI — Yeni pozisyon riskli!' if '🔴 AŞAĞI' in str(piyasa_str) else \
                   '🟡 PİYASA YATAY — Dikkatli ol' if '🟡 YATAY' in str(piyasa_str) else \
                   '🟢 PİYASA YUKARI — Uygun ortam'
except:
    piyasa_uyari = ''
    piyasa_str = ''

print(f"📈 Piyasa: {piyasa_uyari}\n")

# ── AI ANALİZ ──
SONUC_KOLONLARI = ['Kod','Son Fiyat','Hedef (Ort.)','Hedefe Uzaklık (%)','Toplam Skor','Teknik Skor',
                   'Kombine Skor','Teknik Karar','Haber Sayısı','Haber Kararı','Haber Skoru',
                   'Hukuki Risk','AI Gerekçe','Final Skor','Uyarı']
sonuclar = []
hata_sayisi = 0
for i, kod in enumerate(kodlar, 1):
    hh = df_haberler[df_haberler['Hisse'] == kod]
    hisse_haberler = list(zip(hh['Tarih'].astype(str), hh['Başlık'].astype(str)))
    print(f"[{i:2d}/{len(kodlar)}] {kod} ({len(hisse_haberler)} haber)...", end=' ')
    yanit, karar, skor, gerekce, risk = haber_analiz(kod, hisse_haberler)
    if karar == '⚪ HATA' and AI_AKTIF:
        hata_sayisi += 1
        if hata_sayisi == 1:
            print(f"\n   ⚠️ İlk hata: {yanit}")

    row_tek = df_analiz[df_analiz['Kod']==kod].iloc[0]
    row_sis_df = df_sistem_uy[df_sistem_uy['Kod']==kod]
    row_sis = row_sis_df.iloc[0] if len(row_sis_df)>0 else pd.Series(dtype=object)

    kombine = row_tek.get('Kombine Skor', 0)
    kombine = 0 if pd.isna(kombine) else kombine
    final = kombine + skor
    uyari = uyari_hesapla(row_sis)

    sonuclar.append({
        'Kod': kod,
        'Son Fiyat': row_sis.get('Son Fiyat','') if len(row_sis_df)>0 else '',
        'Hedef (Ort.)': row_sis.get('Hedef (Ort.)','') if len(row_sis_df)>0 else '',
        'Hedefe Uzaklık (%)': row_sis.get('Hedefe Uzaklık (%)','') if len(row_sis_df)>0 else '',
        'Toplam Skor': row_tek.get('Toplam Skor', 0),
        'Teknik Skor': row_tek.get('Teknik Skor', 0),
        'Kombine Skor': kombine,
        'Teknik Karar': row_tek.get('Teknik Karar',''),
        'Haber Sayısı': len(hisse_haberler),
        'Haber Kararı': karar,
        'Haber Skoru': skor,
        'Hukuki Risk': risk,
        'AI Gerekçe': gerekce,
        'Final Skor': final,
        'Uyarı': uyari,
    })
    print(f"{karar} ({skor:+d}){' ⚖️ HUKUKİ RİSK' if 'VAR' in str(risk) else ''} → Final:{final} | {uyari[:40]}")
    time.sleep(0.5)

df_sonuc = pd.DataFrame(sonuclar, columns=SONUC_KOLONLARI)
df_sonuc = df_sonuc.sort_values('Final Skor', ascending=False).reset_index(drop=True)

if hata_sayisi:
    print(f"\n⚠️ {hata_sayisi} hissede yapay zekâ hatası oldu (haber puanı 0 sayıldı).")
try:
    with open(ONBELLEK_YOLU, 'w', encoding='utf-8') as f:
        json.dump(ONBELLEK, f, ensure_ascii=False)
except Exception as e:
    print(f"⚠️ Önbellek kaydedilemedi: {e}")
if onbellek_sayisi:
    print(f"♻️ {onbellek_sayisi} hissenin başlıkları değişmemiş: önceki cevap kullanıldı (yapay zekâya sorulmadı)")
print(f"\n📊 Kararlar: { {k: int(v) for k, v in df_sonuc['Haber Kararı'].value_counts().items()} }")
print(f"\n🏆 FINAL TOP 10:")
for _, r in df_sonuc.head(10).iterrows():
    print(f"  {r['Kod']:8s} Fiyat:{r['Son Fiyat']} Hedef:{r['Hedef (Ort.)']} Final:{r['Final Skor']:3.0f} | {r['Haber Kararı']} | {str(r['Uyarı'])[:30]}")

# ── EXCEL'E YAZ ──
wb = load_workbook(EXCEL_YOLU)
TARIH = datetime.now().strftime('%d.%m.%Y %H:%M')

def hc(ws,r,c,v,bg='051020',fg='00FF88',sz=10):
    x=ws.cell(r,c,v); x.font=Font(name='Arial',bold=True,color=fg,size=sz)
    x.fill=PatternFill('solid',fgColor=bg); x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    x.border=BORDER; return x

def dc(ws,r,c,v,bg='091422',fg='E8F4F8',bold=False,sz=9,wrap=False,align='center'):
    try: val=None if isinstance(v,float) and np.isnan(v) else v
    except: val=v
    x=ws.cell(r,c,val); x.font=Font(name='Arial',bold=bold,color=fg,size=sz)
    x.fill=PatternFill('solid',fgColor=bg); x.alignment=Alignment(horizontal=align,vertical='center',wrap_text=wrap)
    x.border=BORDER; return x

# ── 🤖 AI ANALİZ SAYFASI ──
if '🤖 AI Analiz' in wb.sheetnames: del wb['🤖 AI Analiz']
ws_ai = wb.create_sheet('🤖 AI Analiz')
ws_ai.sheet_tab_color = 'FF6B35'
ws_ai.sheet_view.showGridLines = False

ai_cols = ['Kod','Son Fiyat','Hedef (Ort.)','Hedefe Uzaklık (%)','Kombine Skor',
           'Haber Sayısı','Haber Kararı','Haber Skoru','Hukuki Risk','Final Skor','Uyarı','AI Gerekçe']

ai_durum = f"model: {MODEL}" if AI_AKTIF else "⚠️ yapay zekâ bağlantısı yok"
ws_ai.merge_cells(f'A1:{get_column_letter(len(ai_cols))}1')
ws_ai['A1'] = f"🤖 AI Haber Analizi — {TARIH} | {len(df_sonuc)} hisse | {ai_durum} | {piyasa_uyari}"
ws_ai['A1'].font = Font(name='Arial',bold=True,size=13,color='FF6B35')
ws_ai['A1'].fill = PatternFill('solid',fgColor='020D1A')
ws_ai['A1'].alignment = Alignment(horizontal='center',vertical='center')
ws_ai.row_dimensions[1].height = 36
for ci,col in enumerate(ai_cols,1): hc(ws_ai,2,ci,col,fg='FF6B35')
ws_ai.row_dimensions[2].height = 28

for ri,(_,row) in enumerate(df_sonuc.iterrows(),3):
    bg='0D1B2E' if ri%2==0 else '091422'
    for ci,cn in enumerate(ai_cols,1):
        val=row.get(cn)
        if cn=='Final Skor':
            c=dc(ws_ai,ri,ci,val,bg=bg,bold=True,sz=11)
            if pd.notna(val): c.font=Font(name='Arial',bold=True,size=11,color='00FF88' if val>=130 else '00D4FF' if val>=100 else 'FFB300')
        elif cn=='Son Fiyat':
            dc(ws_ai,ri,ci,f"{val}₺" if pd.notna(val) and val!='' else '-',bg=bg,bold=True,fg='FFD700')
        elif cn=='Hedef (Ort.)':
            dc(ws_ai,ri,ci,f"{val}₺" if pd.notna(val) and val!='' else '-',bg=bg,fg='00FF88')
        elif cn=='Hedefe Uzaklık (%)':
            c=dc(ws_ai,ri,ci,val,bg=bg,bold=True)
            if pd.notna(val) and val!='':
                try: c.font=Font(name='Arial',bold=True,size=9,color='00FF88' if float(val)>0 else 'FF4757')
                except: pass
        elif cn=='Haber Kararı' and str(val) in KARAR_R:
            bk,fk=KARAR_R[str(val)]; c=ws_ai.cell(ri,ci,val)
            c.fill=PatternFill('solid',fgColor=bk); c.font=Font(name='Arial',bold=True,size=10,color=fk)
            c.alignment=Alignment(horizontal='center',vertical='center'); c.border=BORDER
        elif cn=='Uyarı':
            val_str=str(val) if val else ''
            if '⚠️' in val_str: dc(ws_ai,ri,ci,val,bg='2D1500',fg='FF6B35',sz=8,wrap=True,align='left')
            elif '🟡' in val_str: dc(ws_ai,ri,ci,val,bg='1A1400',fg='FFB300',sz=8,wrap=True,align='left')
            else: dc(ws_ai,ri,ci,val,bg='001A10',fg='00FF88',sz=8,wrap=True,align='left')
        elif cn=='AI Gerekçe':
            dc(ws_ai,ri,ci,val,bg='051A10',fg='B2DFDB',sz=8,wrap=True,align='left')
        else:
            dc(ws_ai,ri,ci,val,bg=bg)
    ws_ai.row_dimensions[ri].height = 45

for ci,cn in enumerate(ai_cols,1):
    ws_ai.column_dimensions[get_column_letter(ci)].width = {
        'Kod':8,'Son Fiyat':10,'Hedef (Ort.)':10,'Hedefe Uzaklık (%)':12,
        'Kombine Skor':11,'Haber Sayısı':10,'Haber Kararı':12,
        'Haber Skoru':10,'Final Skor':10,'Uyarı':35,'AI Gerekçe':55
    }.get(cn,10)
ws_ai.freeze_panes='A3'

wb.save(EXCEL_YOLU)
print(f"\n✅ Tamamlandı!")
print(f"   🤖 AI Analiz — {len(df_sonuc)} hisse ({ai_durum})")
print(f"   🎯 Kombine AL sayfası hücre 7'de oluşturulacak")
print(f"   📈 Piyasa: {piyasa_uyari}")
