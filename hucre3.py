# ═══════════════════════════════════════════════════════════════
#  HÜCRE 3 — HABER TARAYICI  (v2.2 — 16.10.2026)
#  v2 değişiklikleri:
#   - Sadece son HABER_GUN gündeki haberler alınıyor (eskiden 2017'ye
#     kadar eski sayfalar geliyordu)
#   - TradingView gibi "fiyat/grafik" sayfaları haber sayılmıyor
#   - Çalışmayan KAP ve Mynet okuyucuları çıkarıldı; yerine Google
#     Haberler'de "KAP" araması eklendi
#   - Portföy bölümü kaldırıldı: sadece AL listesinin haberleri çekiliyor
#   v2.1:
#   - Başka şirkete ait KAP bülteni satırları atılıyor ("*** KUVVA *** Devre Kesici"
#     gibi, aramada çıkan ama bu hisseyle ilgisi olmayanlar)
#   - Mynet "Canlı Grafik" / "Günlük Teknik Analiz", YouTube teknik bakış ve
#     "en çok yükselen/işlem gören hisseler" gibi liste sayfaları haber sayılmıyor
#   - Başlığında hisse kodu geçen haberler öne alınıyor (BEYAZ → "Beyaz Saray",
#     LOGO → "Apple logo" gibi kelime benzerlikleri sona düşüyor)
#   v2.2:
#   - Google hata verirse (engel / çok istek) artık sessizce "0 haber" demiyor:
#     hata kodunu yazıyor, 3 kez deniyor, olmazsa Bing Haberler'e geçiyor
#   - Bir kaynak art arda 3 hissede hata verirse o kaynak kapatılıyor (boşuna beklemesin)
#   - Hiç haber gelmezse eski '📰 Haberler' sayfası SİLİNMİYOR (hücre 6 boş haberle
#     yapay zekâ kredisi harcamasın) ve sonda ne yapman gerektiği yazıyor
# ═══════════════════════════════════════════════════════════════
import requests
from bs4 import BeautifulSoup
import pandas as pd
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from urllib.parse import quote
import re
import time
import warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

SISTEM_YOLU = '/content/drive/MyDrive/isyatirim/BIST_Sistem.xlsx'
EXCEL_YOLU  = '/content/drive/MyDrive/isyatirim/BIST_Tam_KAP_v9_11mart.xlsx'

HABER_GUN     = 14    # kaç günlük haber alınsın
HISSE_BASINA  = 6     # bir hisse için en fazla kaç haber

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept-Language': 'tr-TR,tr;q=0.9',
}

# Haber olmayan fiyat/grafik sayfaları
HABER_DEGIL = ['tradingview', 'grafiği, fiyatı', 'fiyatı ve grafiği', 'hisse senedi fiyatı',
               'hisse fiyatı', 'canlı borsa', 'anlık fiyat', 'hisse yorumu ve',
               'canlı grafik', 'günlük teknik analiz', 'teknik bakış', '- youtube',
               'en çok işlem gören hisseler', 'en çok yükselen ve düşen', 'günlük hisse önerileri',
               'borsa yorumu /']

def norm(metin):
    """Türkçe büyük/küçük harf farkını kaldırır (İ/I/ı → i) — 'İşlem' ile 'işlem' eşleşsin."""
    return str(metin).replace('İ', 'i').replace('I', 'i').lower().replace('ı', 'i').replace('\u0307', '')

HABER_DEGIL = [norm(k) for k in HABER_DEGIL]

def baska_sirket_kap(baslik, kod):
    """'KAP *** ŞİRKET *** KOD1,KOD2 *** konu' biçimindeki bülten başka bir şirkete aitse True."""
    if '***' not in baslik:
        return False
    parcalar = [p.strip() for p in baslik.split('***')]
    if len(parcalar) >= 4:
        kodlar = [k.strip().upper() for k in parcalar[2].split(',')]
        return kod.upper() not in kodlar
    return not kod_geciyor(baslik, kod)       # kod bölümü yoksa başlıkta kod aranır

def kod_geciyor(baslik, kod):
    """Hisse kodu başlıkta büyük harfle ayrı kelime olarak geçiyor mu? (BEYAZ evet, Beyaz Saray hayır)"""
    return re.search(rf'(?<![A-ZÇĞİÖŞÜ0-9]){re.escape(kod.upper())}(?![A-ZÇĞİÖŞÜ0-9])', baslik) is not None

# ── AL LİSTESİ ──
df_sistem = pd.read_excel(SISTEM_YOLU)
if 'Sembol' in df_sistem.columns:
    df_sistem = df_sistem.rename(columns={'Sembol': 'Kod'})

AL_HISSELER = df_sistem[df_sistem['Toplam Skor'] >= 55]['Kod'].tolist()
AL_HISSELER = list(dict.fromkeys(AL_HISSELER))
print(f"✅ Haber listesi: {len(AL_HISSELER)} hisse (AL listesi)")
print(f"   {AL_HISSELER[:10]}...")
print(f"   Sadece son {HABER_GUN} günün haberleri alınacak")

SINIR = datetime.now(timezone.utc) - timedelta(days=HABER_GUN)

HATALAR       = {}    # kaynak → {hata türü: adet}
ARDISIK_HATA  = {}    # kaynak → art arda hata veren hisse sayısı
KAPALI        = set() # art arda 3 hissede hata verip kapatılan kaynaklar

def rss_cek(url):
    """RSS adresini en fazla 3 kez dener. (soup, hata) döner; başarılıysa hata None."""
    son_hata = None
    for deneme in range(3):
        try:
            r = requests.get(url, headers=HEADERS, timeout=15)
            bas = r.content[:3000].lower()
            if r.status_code == 200 and (b'<rss' in bas or b'<channel' in bas):
                return BeautifulSoup(r.content, 'xml'), None
            son_hata = (f'HTTP {r.status_code}' if r.status_code != 200
                        else 'RSS yerine başka sayfa (engel/onay sayfası)')
        except Exception as e:
            son_hata = f'bağlantı hatası ({type(e).__name__})'
        if deneme < 2:
            time.sleep(3 * (deneme + 1))
    return None, son_hata

def ogeleri_isle(soup, kod, kaynak_adi):
    """RSS içindeki haberleri filtreleyip listeye çevirir (eski filtrelerin aynısı)."""
    haberler = []
    for item in soup.find_all('item')[:20]:
        try:
            baslik = item.find('title').text.strip()
            link   = item.find('link').text.strip() if item.find('link') else ''
            pub    = item.find('pubDate').text.strip() if item.find('pubDate') else ''
            kaynak_etiket = item.find('News:Source') or item.find('Source')
            if kaynak_etiket is not None and kaynak_etiket.text.strip() and ' - ' not in baslik[-40:]:
                baslik = f"{baslik} - {kaynak_etiket.text.strip()}"   # Bing: kaynak adını Google gibi sona ekle
            if len(baslik) <= 10:
                continue
            if any(k in norm(baslik) for k in HABER_DEGIL):
                continue
            if baska_sirket_kap(baslik, kod):
                continue                      # başka şirketin KAP bildirimi
            try:
                dt = parsedate_to_datetime(pub)
                if dt.tzinfo is None:
                    dt = dt.replace(tzinfo=timezone.utc)
            except:
                continue                      # tarihsiz haberi alma
            if dt < SINIR:
                continue                      # eski haber
            yerel = dt.astimezone()
            haberler.append({'Hisse': kod, 'Kaynak': kaynak_adi,
                             'Tarih': yerel.strftime('%d.%m.%Y %H:%M'),
                             'Başlık': baslik[:150], 'Link': link, '_dt': dt})
        except:
            continue
    return haberler

def ara(kaynak, kod, sorgu, kaynak_adi):
    """kaynak: 'Google' veya 'Bing'. (haberler, hata) döner."""
    if kaynak in KAPALI:
        return [], 'kapalı'
    if kaynak == 'Google':
        url = f"https://news.google.com/rss/search?q={quote(f'{sorgu} when:{HABER_GUN}d')}&hl=tr&gl=TR&ceid=TR:tr"
    else:
        url = f"https://www.bing.com/news/search?q={quote(sorgu)}&format=rss&setlang=tr-TR&cc=TR"
    soup, hata = rss_cek(url)
    if hata:
        HATALAR.setdefault(kaynak, {})
        HATALAR[kaynak][hata] = HATALAR[kaynak].get(hata, 0) + 1
        return [], hata
    return ogeleri_isle(soup, kod, kaynak_adi), None

def kaynak_sonucu(kaynak, hata_var):
    """Art arda 3 hissede hata veren kaynağı kapatır."""
    if kaynak in KAPALI:
        return None
    ARDISIK_HATA[kaynak] = ARDISIK_HATA.get(kaynak, 0) + 1 if hata_var else 0
    if ARDISIK_HATA[kaynak] >= 3:
        KAPALI.add(kaynak)
        return f"⚠️ {kaynak} art arda 3 hissede hata verdi, bu çalıştırmada kapatıldı"
    return None

def hisse_haberleri(kod):
    h, notlar = [], []
    g1, e1 = ara('Google', kod, f'"{kod}" hisse', 'Google')
    g2, e2 = ara('Google', kod, f'"{kod}" KAP', 'Google-KAP')
    h += g1 + g2
    google_hata = e1 or e2
    if google_hata:
        notlar.append(f"Google: {google_hata}")
    if google_hata != 'kapalı':
        uyari = kaynak_sonucu('Google', bool(e1 and e2))
        if uyari: notlar.append(uyari + ' → Bing kullanılacak')
    if e1 and e2:                                   # Google bu hisse için hiç cevap vermedi → Bing
        b1, eb = ara('Bing', kod, f'"{kod}" hisse', 'Bing')
        h += b1
        if eb:
            notlar.append(f"Bing: {eb}")
        if eb != 'kapalı':
            uyari = kaynak_sonucu('Bing', bool(eb))
            if uyari: notlar.append(uyari)
    # aynı başlığı bir kez al; önce kodu başlıkta geçenler, sonra diğerleri (her grup en yeniden eskiye)
    gorulen = set(); temiz = []
    for x in sorted(h, key=lambda x: (not kod_geciyor(x['Başlık'], kod), -x['_dt'].timestamp())):
        anahtar = x['Başlık'].lower()
        if anahtar not in gorulen:
            gorulen.add(anahtar); temiz.append(x)
    secilen = temiz[:HISSE_BASINA]
    return sorted(secilen, key=lambda x: x['_dt'], reverse=True), notlar

def tum_haberleri_cek():
    print(f"\n📰 {len(AL_HISSELER)} hisse için haberler çekiliyor...\n")
    tum = []
    for i, kod in enumerate(AL_HISSELER, 1):
        print(f"[{i:2d}/{len(AL_HISSELER)}] {kod}...", end=' ')
        h, notlar = hisse_haberleri(kod)
        tum.extend(h)
        print(f"{'✅' if h else '⚪'} {len(h)} haber" + (f"   ({'; '.join(notlar)})" if notlar else ''))
        time.sleep(1.2)
    print(f"\n✅ Toplam {len(tum)} haber toplandı (son {HABER_GUN} gün).")
    return tum

def excel_yaz(haber_listesi):
    kolonlar = ['Hisse', 'Kaynak', 'Tarih', 'Başlık', 'Link']
    df = pd.DataFrame(haber_listesi, columns=kolonlar + ['_dt'])
    df = df.drop_duplicates(subset=['Hisse', 'Başlık']).reset_index(drop=True)
    try:
        wb = load_workbook(EXCEL_YOLU)
    except:
        from openpyxl import Workbook
        wb = Workbook()
    if '📰 Haberler' in wb.sheetnames:
        del wb['📰 Haberler']
    ws = wb.create_sheet('📰 Haberler')
    kaynak_renk = {'Google': 'FFE8F5E9', 'Google-KAP': 'FFFCE4EC'}
    ws.merge_cells('A1:E1')
    ws['A1'] = (f"📰 HABER AKIŞI — {datetime.now().strftime('%d.%m.%Y %H:%M')} — {len(df)} haber — "
                f"{len(AL_HISSELER)} hisse — son {HABER_GUN} gün")
    ws['A1'].fill = PatternFill('solid', fgColor='FF0F3460')
    ws['A1'].font = Font(bold=True, color='FFFFFFFF', size=12)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28
    for col, ad in enumerate(['Hisse', 'Kaynak', 'Tarih', 'Başlık / Haber', 'Link'], 1):
        c = ws.cell(2, col, ad)
        c.fill = PatternFill('solid', fgColor='FF1A1A2E')
        c.font = Font(bold=True, color='FFFFFFFF', size=11)
        c.alignment = Alignment(horizontal='center', vertical='center')
    for idx, row in df.iterrows():
        r = idx + 3
        renk = kaynak_renk.get(row['Kaynak'], 'FFFFF8E1')
        c1 = ws.cell(r, 1, row['Hisse'])
        c1.font = Font(bold=True, size=10, color='000000')
        ws.cell(r, 2, row['Kaynak'])
        ws.cell(r, 3, row['Tarih'])
        ws.cell(r, 4, row['Başlık']).alignment = Alignment(wrap_text=True, vertical='center')
        ws.cell(r, 5, row['Link'])
        for col in range(1, 6):
            ws.cell(r, col).fill = PatternFill('solid', fgColor=renk)
        ws.row_dimensions[r].height = 30
    for i, g in enumerate([10, 16, 16, 80, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.freeze_panes = 'A3'
    ws.auto_filter.ref = f"A2:E{max(len(df), 1) + 2}"
    wb.save(EXCEL_YOLU)
    print(f"\n✅ '📰 Haberler' sayfası güncellendi! {len(df)} haber | {df['Hisse'].nunique()} hisse")
    if len(df):
        print(f"📊 Kaynak: { {k: int(v) for k, v in df['Kaynak'].value_counts().items()} }")
    habersiz = [k for k in AL_HISSELER if k not in set(df['Hisse'])]
    if habersiz:
        print(f"ℹ️  Son {HABER_GUN} günde haberi olmayan: {', '.join(habersiz)}")

def teshis(haberler):
    """Sorun varsa ne olduğunu ve ne yapılacağını yazar."""
    if HATALAR:
        print("\n🔎 BAĞLANTI DURUMU:")
        for kaynak, d in HATALAR.items():
            print(f"   {kaynak}: " + ', '.join(f'{k} × {v}' for k, v in d.items()))
    if not haberler and HATALAR:
        print("\n❗ Haber kaynakları bu Colab bağlantısına cevap vermedi (geçici engel / çok istek).")
        print("   Çözüm: üst menüden Çalışma zamanı → 'Bağlantıyı kes ve çalışma zamanını sil',")
        print("   sonra hücre 1 → 2 → 3'ü yeniden çalıştır (yeni bağlantı genelde yeni adres alır).")
        print("   Ya da 1–2 saat sonra tekrar dene. Zirve Stratejisi (hücre 10) haberlere bakmaz, onu etkilemez.")
    elif not haberler:
        print("\n❗ Bağlantı hatası yok ama hiç haber de çıkmadı — bu çıktıyı Claude'a gönder.")

# ÇALIŞTIR
if not AL_HISSELER:
    print("❗ AL listesinde hisse yok (Toplam Skor ≥ 55 olan hisse çıkmadı) — haber çekilecek hisse yok.")
    print("   Hücre 2 hatasız bitti mi? Bitmediyse önce onu çalıştır.")
else:
    haberler = tum_haberleri_cek()
    teshis(haberler)
    if haberler:
        excel_yaz(haberler)
    else:
        print("\n📁 Eski '📰 Haberler' sayfası olduğu gibi bırakıldı (boş sayfayla değiştirilmedi).")
        print("   Hücre 6'yı şimdi çalıştırırsan eski haberlerle çalışır; başlıklar aynı olduğu için çoğu hisse önbellekten gelir, kredi harcamaz.")
